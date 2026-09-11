"""问题三：采用细网格有限体积法计算药材全域含水率达标时间。

内部求解采用单元中心网格，输出仍按附件要求为每 60 s、每 0.1 cm。
表面半个单元的内部阻力与外部对流阻力串联，避免把输出节点直接
作为控制体中心所造成的边界离散误差。
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent
INPUT_FILE = PROJECT_ROOT / "附件" / "附件1.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "附件" / "附件3" / "result3.xlsx"
DIAGNOSTICS_FILE = PROJECT_ROOT / "附件" / "附件3" / "result3_diagnostics.json"

PELLET_RADIUS = 0.02
CONVECTIVE_HEAT_COEFF = 25.0
CONVECTIVE_MASS_COEFF = 8.0e-7
INITIAL_TEMPERATURE = 28.0
INITIAL_MOISTURE = 2.55
CRITICAL_MOISTURE = 0.15

INTERNAL_INTERVALS = 320          # 内部步长 0.00625 cm
OUTPUT_SPACING = 0.001            # 输出步长 0.1 cm
TIME_STEP = 30.0                    # 内部推进 30 s，结果仍每 60 s 输出
REPORT_INTERVAL = 60.0
MAX_SIMULATION_TIME = 5 * 24 * 3600.0

CONVERGENCE_TOL = 1.0e-9
MAX_ITERATIONS = 200
RELAXATION_FACTOR = 0.6
USE_LAST_HOUR_MEAN = True


class DryingRoomBoundary(NamedTuple):
    """附件一边界数据和 4 h 后平台值。"""

    times: np.ndarray
    temperatures: np.ndarray
    moisture: np.ndarray
    plateau_temperature: float
    plateau_moisture: float
    plateau_window: float
    plateau_sample_count: int
    plateau_temperature_std: float
    plateau_moisture_std: float


class BoundaryScenario(NamedTuple):
    """用于长期边界情景复算的名称、说明和边界数据。"""

    name: str
    label: str
    boundary: DryingRoomBoundary


class CellCenteredGrid(NamedTuple):
    """圆柱径向单元中心有限体积网格。"""

    faces: np.ndarray
    centers: np.ndarray
    volumes: np.ndarray
    spacing: float


def read_drying_boundary(
    path: Path,
    plateau_window: float = 3600.0,
    temperature_offset: float = 0.0,
    moisture_offset: float = 0.0,
    use_last_sample: bool = False,
) -> DryingRoomBoundary:
    """读取附件一，并按指定末段窗口或末点构造 4 h 后平台。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = [
        row for row in workbook.active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一缺少有效边界数据。")
    if not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError("附件一时间必须严格递增。")

    if plateau_window <= 0:
        raise ValueError("平台均值窗口必须为正数。")
    stable = data[:, 0] >= data[-1, 0] - plateau_window
    if use_last_sample or not USE_LAST_HOUR_MEAN:
        plateau_temperature = float(data[-1, 1])
        plateau_moisture = float(data[-1, 2])
        sample_count = 1
        temperature_std = 0.0
        moisture_std = 0.0
    else:
        stable_values = data[stable, 1:3]
        plateau_temperature = float(np.mean(stable_values[:, 0]))
        plateau_moisture = float(np.mean(stable_values[:, 1]))
        sample_count = int(stable_values.shape[0])
        ddof = 1 if sample_count > 1 else 0
        temperature_std = float(np.std(stable_values[:, 0], ddof=ddof))
        moisture_std = float(np.std(stable_values[:, 1], ddof=ddof))
    return DryingRoomBoundary(
        data[:, 0],
        data[:, 1],
        data[:, 2],
        plateau_temperature + temperature_offset,
        plateau_moisture + moisture_offset,
        plateau_window,
        sample_count,
        temperature_std,
        moisture_std,
    )


def build_boundary_scenarios(path: Path) -> list[BoundaryScenario]:
    """构造末段窗口、末点及温度/水分单因素情景。"""
    baseline = read_drying_boundary(path)
    return [
        BoundaryScenario(
            "last_30_min_mean",
            "末 30 min 均值",
            read_drying_boundary(path, plateau_window=1800.0),
        ),
        BoundaryScenario(
            "last_sample",
            "最后采样点",
            read_drying_boundary(path, use_last_sample=True),
        ),
        BoundaryScenario(
            "temperature_minus_1sigma",
            "温度平台 -1σ",
            read_drying_boundary(
                path, temperature_offset=-baseline.plateau_temperature_std
            ),
        ),
        BoundaryScenario(
            "temperature_plus_1sigma",
            "温度平台 +1σ",
            read_drying_boundary(
                path, temperature_offset=baseline.plateau_temperature_std
            ),
        ),
        BoundaryScenario(
            "moisture_minus_1sigma",
            "环境水分平台 -1σ",
            read_drying_boundary(
                path, moisture_offset=-baseline.plateau_moisture_std
            ),
        ),
        BoundaryScenario(
            "moisture_plus_1sigma",
            "环境水分平台 +1σ",
            read_drying_boundary(
                path, moisture_offset=baseline.plateau_moisture_std
            ),
        ),
    ]


def linear_threshold_crossing(
    previous_time: float,
    previous_value: float,
    current_time: float,
    current_value: float,
    threshold: float,
) -> float:
    """在相邻时间层之间线性估计从上方穿越阈值的时刻。"""
    if current_time <= previous_time:
        raise ValueError("当前时刻必须晚于前一时刻。")
    if not (previous_value >= threshold and current_value < threshold):
        raise ValueError("相邻数值没有从上方严格穿越阈值。")
    return previous_time + (
        (previous_value - threshold) / (previous_value - current_value)
    ) * (current_time - previous_time)


def relative_balance_imbalance(
    initial_inventory: float,
    current_inventory: float,
    cumulative_outflow: float,
) -> float:
    """返回离散 C 方程累计收支相对不平衡，不解释为真实质量误差。"""
    if initial_inventory <= 0:
        raise ValueError("初始积分必须为正数。")
    return abs(
        initial_inventory - current_inventory - cumulative_outflow
    ) / initial_inventory


def boundary_value(
    current_time: float,
    boundary: DryingRoomBoundary,
) -> tuple[float, float]:
    """0–4 h 线性插值，4 h 后直接采用平台值。"""
    if current_time <= boundary.times[-1]:
        return (
            float(np.interp(current_time, boundary.times, boundary.temperatures)),
            float(np.interp(current_time, boundary.times, boundary.moisture)),
        )
    return boundary.plateau_temperature, boundary.plateau_moisture


def density(moisture: np.ndarray) -> np.ndarray:
    return 650.0 + 128.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    """附录三扩散系数，温度转换为 K。"""
    safe_moisture = np.maximum(moisture, 1.0e-12)
    temperature_kelvin = temperature_celsius + 273.15
    return (
        2.4e-3
        * np.exp(-0.45 / safe_moisture)
        * np.exp(-3850.0 / temperature_kelvin)
    )


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    return 2.0 * left * right / np.maximum(left + right, 1.0e-30)


def build_grid(intervals: int) -> CellCenteredGrid:
    if intervals < 4:
        raise ValueError("内部径向区间数至少为 4。")
    faces = np.linspace(0.0, PELLET_RADIUS, intervals + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return CellCenteredGrid(
        faces,
        centers,
        volumes,
        PELLET_RADIUS / intervals,
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    """Thomas 算法。"""
    lower = lower.astype(float, copy=True)
    diagonal = diagonal.astype(float, copy=True)
    upper = upper.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    for index in range(1, len(diagonal)):
        factor = lower[index - 1] / diagonal[index - 1]
        diagonal[index] -= factor * upper[index - 1]
        right_hand_side[index] -= factor * right_hand_side[index - 1]
    solution = np.empty_like(right_hand_side)
    solution[-1] = right_hand_side[-1] / diagonal[-1]
    for index in range(len(diagonal) - 2, -1, -1):
        solution[index] = (
            right_hand_side[index] - upper[index] * solution[index + 1]
        ) / diagonal[index]
    return solution


def assemble_system(
    old_values: np.ndarray,
    storage_capacity: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """组装后向 Euler 单元中心有限体积方程。

    表面有效系数：
    h_eff = 1 / [1/h + (Δr/2)/Γ_N]。
    """
    storage = storage_capacity * grid.volumes / time_step
    face_transport = harmonic_mean(
        transport_coefficients[:-1], transport_coefficients[1:]
    )
    internal_conductance = (
        grid.faces[1:-1] * face_transport / grid.spacing
    )
    diagonal = storage.copy()
    diagonal[:-1] += internal_conductance
    diagonal[1:] += internal_conductance
    lower = -internal_conductance.copy()
    upper = -internal_conductance.copy()
    right_hand_side = storage * old_values

    last_transport = max(float(transport_coefficients[-1]), 1.0e-30)
    effective_transfer = 1.0 / (
        1.0 / external_transfer_coefficient
        + 0.5 * grid.spacing / last_transport
    )
    surface_conductance = PELLET_RADIUS * effective_transfer
    diagonal[-1] += surface_conductance
    right_hand_side[-1] += surface_conductance * external_value
    return lower, diagonal, upper, right_hand_side


def advance_temperature(
    old_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_temperature: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> np.ndarray:
    system = assemble_system(
        old_temperature,
        density(reference_moisture) * heat_capacity(reference_moisture),
        thermal_conductivity(reference_moisture),
        CONVECTIVE_HEAT_COEFF,
        room_temperature,
        grid,
        time_step,
    )
    return solve_tridiagonal(*system)


def advance_moisture(
    old_moisture: np.ndarray,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_moisture: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> np.ndarray:
    system = assemble_system(
        old_moisture,
        np.ones_like(old_moisture),
        moisture_diffusivity(reference_moisture, reference_temperature),
        CONVECTIVE_MASS_COEFF,
        room_moisture,
        grid,
        time_step,
    )
    return solve_tridiagonal(*system)


def reconstruct_output_field(
    cell_values: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    output_nodes: np.ndarray,
    grid: CellCenteredGrid,
) -> np.ndarray:
    """从单元中心解重构中心、内部输出节点和表面值。"""
    output_values = np.interp(output_nodes, grid.centers, cell_values)
    output_values[0] = (9.0 * cell_values[0] - cell_values[1]) / 8.0

    last_transport = max(float(transport_coefficients[-1]), 1.0e-30)
    half_cell_conductance = 2.0 * last_transport / grid.spacing
    output_values[-1] = (
        half_cell_conductance * cell_values[-1]
        + external_transfer_coefficient * external_value
    ) / (half_cell_conductance + external_transfer_coefficient)
    return output_values


def solve_problem_three(
    internal_intervals: int = INTERNAL_INTERVALS,
    time_step: float = TIME_STEP,
    report_interval: float = REPORT_INTERVAL,
    max_time: float = MAX_SIMULATION_TIME,
    boundary: DryingRoomBoundary | None = None,
    critical_moisture: float = CRITICAL_MOISTURE,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, dict]:
    """求解到 60 s 输出网格上全域首次严格达标。"""
    report_steps = int(round(report_interval / time_step))
    total_steps = int(round(max_time / time_step))
    if report_steps < 1 or not np.isclose(
        report_steps * time_step, report_interval
    ):
        raise ValueError("输出间隔必须是内部时间步长的整数倍。")
    if total_steps < 1 or not np.isclose(total_steps * time_step, max_time):
        raise ValueError("最大时间必须是内部时间步长的整数倍。")

    if boundary is None:
        boundary = read_drying_boundary(INPUT_FILE)
    grid = build_grid(internal_intervals)
    output_nodes = np.arange(
        0.0, PELLET_RADIUS + 0.5 * OUTPUT_SPACING, OUTPUT_SPACING
    )
    temperature = np.full(internal_intervals, INITIAL_TEMPERATURE)
    moisture = np.full(internal_intervals, INITIAL_MOISTURE)

    output_times: list[float] = []
    output_moisture: list[np.ndarray] = []
    iteration_counts: list[int] = []
    previous_time = 0.0
    previous_maximum = INITIAL_MOISTURE
    continuous_threshold_time = None
    threshold_before = None
    threshold_after = None
    initial_inventory = float(np.dot(moisture, grid.volumes))
    current_inventory = initial_inventory
    cumulative_outflow = 0.0
    maximum_step_balance_residual = 0.0
    center_controls_every_step = True
    radially_nonincreasing_every_step = True
    checked_steps = 0

    for step in range(1, total_steps + 1):
        current_time = step * time_step
        room_temperature, room_moisture = boundary_value(
            current_time, boundary
        )
        temperature_iterate = temperature.copy()
        moisture_iterate = moisture.copy()

        for iteration in range(1, MAX_ITERATIONS + 1):
            temperature_candidate = advance_temperature(
                temperature,
                moisture_iterate,
                room_temperature,
                grid,
                time_step,
            )
            moisture_candidate = advance_moisture(
                moisture,
                temperature_candidate,
                moisture_iterate,
                room_moisture,
                grid,
                time_step,
            )
            new_temperature = (
                RELAXATION_FACTOR * temperature_candidate
                + (1.0 - RELAXATION_FACTOR) * temperature_iterate
            )
            new_moisture = (
                RELAXATION_FACTOR * moisture_candidate
                + (1.0 - RELAXATION_FACTOR) * moisture_iterate
            )
            temperature_error = np.max(
                np.abs(new_temperature - temperature_iterate)
                / (1.0 + np.abs(new_temperature))
            )
            moisture_error = np.max(
                np.abs(new_moisture - moisture_iterate)
                / (1.0 + np.abs(new_moisture))
            )
            temperature_iterate = new_temperature
            moisture_iterate = new_moisture
            if max(temperature_error, moisture_error) < CONVERGENCE_TOL:
                break
        else:
            raise RuntimeError(
                f"{current_time:.0f} s 未收敛："
                f"{temperature_error:.3e}, {moisture_error:.3e}"
            )

        temperature = temperature_iterate
        moisture = moisture_iterate
        if not np.all(np.isfinite(moisture)) or np.min(moisture) <= 0:
            raise FloatingPointError(f"{current_time:.0f} s 含水率异常。")

        diffusivity = moisture_diffusivity(moisture, temperature)
        last_diffusivity = max(float(diffusivity[-1]), 1.0e-30)
        effective_mass_transfer = 1.0 / (
            1.0 / CONVECTIVE_MASS_COEFF
            + 0.5 * grid.spacing / last_diffusivity
        )
        surface_outflow = (
            PELLET_RADIUS
            * effective_mass_transfer
            * (float(moisture[-1]) - room_moisture)
        )
        next_inventory = float(np.dot(moisture, grid.volumes))
        step_balance_residual = abs(
            next_inventory - current_inventory + time_step * surface_outflow
        ) / initial_inventory
        maximum_step_balance_residual = max(
            maximum_step_balance_residual, step_balance_residual
        )
        cumulative_outflow += time_step * surface_outflow
        current_inventory = next_inventory

        reconstructed = reconstruct_output_field(
            moisture,
            diffusivity,
            CONVECTIVE_MASS_COEFF,
            room_moisture,
            output_nodes,
            grid,
        )
        current_maximum = max(
            float(np.max(moisture)),
            float(np.max(reconstructed)),
        )
        center_value = float(reconstructed[0])
        center_controls_every_step = center_controls_every_step and bool(
            current_maximum <= center_value + 1.0e-12
        )
        radially_nonincreasing_every_step = (
            radially_nonincreasing_every_step
            and bool(np.all(np.diff(moisture) <= 1.0e-12))
            and bool(np.all(np.diff(reconstructed) <= 1.0e-12))
        )
        checked_steps += 1
        if (
            continuous_threshold_time is None
            and previous_maximum >= critical_moisture
            and current_maximum < critical_moisture
        ):
            continuous_threshold_time = linear_threshold_crossing(
                previous_time,
                previous_maximum,
                current_time,
                current_maximum,
                critical_moisture,
            )
            threshold_before = previous_maximum
            threshold_after = current_maximum

        if step % report_steps == 0:
            output_times.append(current_time)
            output_moisture.append(reconstructed.copy())
            if current_maximum < critical_moisture:
                iteration_counts.append(iteration)
                break

        iteration_counts.append(iteration)
        previous_time = current_time
        previous_maximum = current_maximum
    else:
        raise RuntimeError(f"{max_time / 3600:.1f} h 内未达到阈值。")

    moisture_field = np.asarray(output_moisture)
    diagnostics = {
        "continuous_threshold_time_s": float(continuous_threshold_time),
        "discrete_threshold_time_s": float(output_times[-1]),
        "maximum_moisture_before": float(threshold_before),
        "maximum_moisture_after": float(threshold_after),
        "center_controls_threshold": center_controls_every_step,
        "all_domain_checked_every_step": checked_steps == len(iteration_counts),
        "radially_nonincreasing_every_step": radially_nonincreasing_every_step,
        "moisture_balance_relative_imbalance": relative_balance_imbalance(
            initial_inventory, current_inventory, cumulative_outflow
        ),
        "maximum_step_balance_relative_residual": maximum_step_balance_residual,
        "initial_moisture_integral": initial_inventory,
        "final_moisture_integral": current_inventory,
        "cumulative_boundary_outflow": cumulative_outflow,
        "maximum_picard_iterations": int(max(iteration_counts)),
        "mean_picard_iterations": float(np.mean(iteration_counts)),
        "internal_intervals": internal_intervals,
        "internal_spacing_cm": 100.0 * grid.spacing,
        "time_step_s": time_step,
        "report_interval_s": report_interval,
        "critical_moisture": critical_moisture,
        "plateau_temperature_c": boundary.plateau_temperature,
        "plateau_moisture_kgkg": boundary.plateau_moisture,
        "plateau_window_s": boundary.plateau_window,
        "plateau_sample_count": boundary.plateau_sample_count,
    }
    return (
        np.asarray(output_times),
        output_nodes,
        moisture_field,
        diagnostics,
    )


def write_result_workbook(
    times: np.ndarray,
    nodes: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    """写入 result3.xlsx；底层保留全精度，显示四位小数。"""
    workbook = load_workbook(OUTPUT_FILE)
    worksheet = workbook.worksheets[0]
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    worksheet.cell(1, 1).value = "时间\\到药材中心的距离"
    worksheet.cell(1, 1).number_format = "@"
    for column, radius in enumerate(nodes, start=2):
        worksheet.cell(1, column).value = round(float(radius * 100.0), 1)
        worksheet.cell(1, column).number_format = "0.0"
    for row, current_time in enumerate(times, start=2):
        worksheet.cell(row, 1).value = int(round(float(current_time)))
        worksheet.cell(row, 1).number_format = "0"
        for column, value in enumerate(moisture_field[row - 2], start=2):
            worksheet.cell(row, column).value = float(value)
            worksheet.cell(row, column).number_format = "0.0000"
    workbook.save(OUTPUT_FILE)


def main() -> None:
    parser = argparse.ArgumentParser(description="问题三全域含水率达标时间求解")
    parser.add_argument(
        "--verification",
        action="store_true",
        help="额外复算边界情景及空间、时间加密案例并保存诊断 JSON",
    )
    args = parser.parse_args()

    baseline_boundary = read_drying_boundary(INPUT_FILE)
    times, nodes, moisture_field, diagnostics = solve_problem_three(
        boundary=baseline_boundary
    )
    write_result_workbook(times, nodes, moisture_field)
    print(f"问题三重算完成：{OUTPUT_FILE}")
    print(
        f"连续阈值：{diagnostics['continuous_threshold_time_s'] / 3600:.4f} h；"
        f"首次 60 s 离散达标："
        f"{diagnostics['discrete_threshold_time_s'] / 3600:.4f} h"
    )
    print(
        f"内部网格 {diagnostics['internal_intervals']} 个区间，"
        f"Δr={diagnostics['internal_spacing_cm']:.5f} cm；"
        f"输出尺寸 {moisture_field.shape[0]}×{moisture_field.shape[1]}"
    )
    print(
        f"阈值前后最大含水率："
        f"{diagnostics['maximum_moisture_before']:.10f} / "
        f"{diagnostics['maximum_moisture_after']:.10f}"
    )
    print(
        f"离散 C 方程累计相对收支不平衡："
        f"{diagnostics['moisture_balance_relative_imbalance']:.3e}；"
        f"最大 Picard 迭代次数：{diagnostics['maximum_picard_iterations']}"
    )

    if not args.verification:
        return

    def case_record(name: str, label: str, case_diagnostics: dict) -> dict:
        record = dict(case_diagnostics)
        record["name"] = name
        record["label"] = label
        record["continuous_threshold_time_h"] = (
            case_diagnostics["continuous_threshold_time_s"] / 3600.0
        )
        record["discrete_threshold_time_h"] = (
            case_diagnostics["discrete_threshold_time_s"] / 3600.0
        )
        return record

    report = {
        "scope": "离散方程诊断与情景敏感性；不是概率置信区间或真实质量守恒证明。",
        "baseline": case_record(
            "last_1_hour_mean", "末 1 h 均值（基准）", diagnostics
        ),
        "boundary_scenarios": [],
        "numerical_refinement": [],
    }

    for scenario in build_boundary_scenarios(INPUT_FILE):
        print(f"复算边界情景：{scenario.label}")
        _, _, _, case_diagnostics = solve_problem_three(boundary=scenario.boundary)
        report["boundary_scenarios"].append(
            case_record(scenario.name, scenario.label, case_diagnostics)
        )

    refinement_cases = [
        ("space_refined", "空间加密 N=640, Δt=30 s", 640, 30.0),
        ("time_refined", "时间加密 N=320, Δt=15 s", 320, 15.0),
    ]
    for name, label, intervals, time_step in refinement_cases:
        print(f"复算数值加密：{label}")
        _, _, _, case_diagnostics = solve_problem_three(
            internal_intervals=intervals,
            time_step=time_step,
            boundary=baseline_boundary,
        )
        report["numerical_refinement"].append(
            case_record(name, label, case_diagnostics)
        )

    with DIAGNOSTICS_FILE.open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    print(f"验证诊断已保存：{DIAGNOSTICS_FILE}")


if __name__ == "__main__":
    main()
