"""问题二：固定圆柱药材的变物性径向传热—传质计算。

输入：附件/附件1.xlsx
输出：附件/附件3/result2.xlsx

模型：一维轴对称变物性 Fourier 导热 + Fick 水分扩散。
空间离散：有限体积法；时间离散：后向 Euler；非线性：Picard 迭代。

说明：问题二统一从 t=0 开始采用附录3经验公式。附录3未重新给出
对流换热系数和对流传质系数，故基线沿用问题一的 h 和 h_m；这一点
属于待文献或题意进一步确认的建模假设。
"""

from __future__ import annotations

from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook


# 路径以脚本所在目录为基准，换电脑或换工作目录都不用改代码
PROJECT_ROOT = Path(__file__).resolve().parent
INPUT_FILE = PROJECT_ROOT / "附件" / "附件1.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "附件" / "附件3" / "result2.xlsx"

# 药材几何和边界传递参数
PELLET_RADIUS = 0.02             # 药材半径，单位 m
CONVECTIVE_HEAT_COEFF = 25.0     # 对流换热系数，单位 W/(m^2·K)，沿用附录2
CONVECTIVE_MASS_COEFF = 8.0e-7   # 对流传质系数，单位 m/s，沿用附录2

# 初始条件
INITIAL_TEMPERATURE = 28.0       # 初始温度，单位 ℃
INITIAL_MOISTURE = 2.55          # 初始干基含水率，单位 kg/kg

# 时间、空间和非线性迭代设置
END_TIME = 10800                  # 计算终点，单位 s（3 h）
TIME_STEP = 1.0                   # 时间步长，单位 s
RADIAL_INTERVALS = 20             # 径向区间数，20 个区间对应 0.1 cm
CONVERGENCE_TOL = 1.0e-9          # 热湿耦合 Picard 相对误差容差
MAX_ITERATIONS = 60               # 单个时间步最大迭代次数
RELAXATION_FACTOR = 0.8           # 欠松弛因子，改善变物性耦合的收敛性


class DryingRoomBoundary(NamedTuple):
    """附件一给出的烘房侧边界时间序列。"""

    times: np.ndarray              # 时间，单位 s
    temperatures: np.ndarray       # 烘房温度，单位 ℃
    moisture: np.ndarray           # 烘房水分浓度，单位 kg/kg


class RadialGrid(NamedTuple):
    """一维轴对称有限体积网格的几何量。"""

    nodes: np.ndarray              # 节点半径，共 RADIAL_INTERVALS + 1 个
    interfaces: np.ndarray         # 相邻节点中点半径
    area_factors: np.ndarray       # 控制体径向面积因子


def read_drying_boundary(path: Path) -> DryingRoomBoundary:
    """读取附件一，并检查时间序列是否严格递增。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    rows = [
        row for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)

    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一必须包含至少两行有效的时间、温度和水分浓度数据。")

    times = data[:, 0]
    if not np.all(np.diff(times) > 0):
        raise ValueError("附件一中的时间必须严格递增。")

    return DryingRoomBoundary(
        times=times,
        temperatures=data[:, 1],
        moisture=data[:, 2],
    )


def linear_interp(t: float, times: np.ndarray, values: np.ndarray) -> float:
    """在附件一数据范围内进行分段线性插值。

    问题二只计算到 10800 s，而附件一已覆盖到 14400 s，
    因此这里不需要对边界数据进行外推。
    """
    if t < times[0] or t > times[-1]:
        raise ValueError(
            f"计算时刻 {t} s 超出附件一边界范围 "
            f"[{times[0]}, {times[-1]}] s。"
        )
    return float(np.interp(t, times, values))


def density(moisture: np.ndarray) -> np.ndarray:
    """附录3密度关系，单位 kg/m^3。"""
    return 650.0 + 128.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    """附录3比热容关系，单位 J/(kg·K)。"""
    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    """附录3导热系数关系，单位 W/(m·K)。"""
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    """附录3水分扩散系数，单位 m^2/s。

    附录3中的指数温度必须使用绝对温度，因此先将摄氏温度
    转换为 T_K = T_C + 273.15。
    """
    safe_moisture = np.maximum(moisture, 1.0e-12)
    temperature_kelvin = temperature_celsius + 273.15
    return (
        2.4e-3
        * np.exp(-0.45 / safe_moisture)
        * np.exp(-3850.0 / temperature_kelvin)
    )


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """计算相邻控制体之间的调和平均传输系数。"""
    denominator = np.maximum(left + right, 1.0e-30)
    return 2.0 * left * right / denominator


def build_radial_grid(intervals: int) -> RadialGrid:
    """建立等距径向节点和轴对称控制体面积因子。"""
    if intervals < 2:
        raise ValueError("径向区间数至少为2。")

    nodes = np.linspace(0.0, PELLET_RADIUS, intervals + 1)
    interfaces = (nodes[:-1] + nodes[1:]) / 2.0

    area_factors = np.empty(intervals + 1)
    area_factors[0] = 0.5 * interfaces[0] ** 2
    area_factors[1:-1] = 0.5 * (interfaces[1:] ** 2 - interfaces[:-1] ** 2)
    area_factors[-1] = 0.5 * (PELLET_RADIUS**2 - interfaces[-1] ** 2)

    return RadialGrid(
        nodes=nodes,
        interfaces=interfaces,
        area_factors=area_factors,
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    """Thomas 算法：求解三对角线性方程组。"""
    diagonal = diagonal.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    upper = upper.astype(float, copy=True)

    if np.any(np.abs(diagonal) < 1.0e-30):
        raise FloatingPointError("三对角方程组出现近零主对角元。")

    for i in range(1, len(diagonal)):
        factor = lower[i - 1] / diagonal[i - 1]
        diagonal[i] -= factor * upper[i - 1]
        right_hand_side[i] -= factor * right_hand_side[i - 1]

    solution = np.empty_like(right_hand_side)
    solution[-1] = right_hand_side[-1] / diagonal[-1]
    for i in range(len(diagonal) - 2, -1, -1):
        solution[i] = (
            right_hand_side[i] - upper[i] * solution[i + 1]
        ) / diagonal[i]
    return solution


def assemble_control_volume_system(
    storage_coefficients: np.ndarray,
    old_values: np.ndarray,
    grid: RadialGrid,
    interface_coefficients: np.ndarray,
    surface_transfer_coefficient: float,
    boundary_value: float,
    time_step: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """组装变物性后向 Euler 轴对称控制体方程。

    变物性只进入界面通量系数和局部储存系数；
    统一的 2πL 因子在守恒方程两边约去。
    """
    del time_step  # 时间步已包含在 storage_coefficients 中，保留接口便于阅读
    node_count = len(grid.nodes)
    spacing = grid.nodes[1] - grid.nodes[0]

    lower = np.zeros(node_count - 1)
    upper = np.zeros(node_count - 1)
    diagonal = storage_coefficients.copy()
    right_hand_side = storage_coefficients * old_values

    for i in range(node_count):
        inner_coefficient = (
            grid.interfaces[i - 1] * interface_coefficients[i - 1] / spacing
            if i > 0 else 0.0
        )
        outer_coefficient = (
            grid.interfaces[i] * interface_coefficients[i] / spacing
            if i < node_count - 1 else 0.0
        )

        diagonal[i] += inner_coefficient + outer_coefficient
        if i > 0:
            lower[i - 1] = -inner_coefficient
        if i < node_count - 1:
            upper[i] = -outer_coefficient

    # 圆心为对称零通量；表面采用 Robin 边界条件。
    surface_coefficient = PELLET_RADIUS * surface_transfer_coefficient
    diagonal[-1] += surface_coefficient
    right_hand_side[-1] += surface_coefficient * boundary_value
    return lower, diagonal, upper, right_hand_side


def advance_temperature(
    old_temperature: np.ndarray,
    current_time: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    reference_moisture: np.ndarray,
    time_step: float,
) -> np.ndarray:
    """在当前 Picard 迭代的含水率下推进一个温度时间步。"""
    rho = density(reference_moisture)
    cp = heat_capacity(reference_moisture)
    conductivity = thermal_conductivity(reference_moisture)

    storage_coefficients = rho * cp * grid.area_factors / time_step
    interface_conductivity = harmonic_mean(
        conductivity[:-1], conductivity[1:]
    )
    room_temperature = linear_interp(
        current_time, boundary.times, boundary.temperatures
    )

    system = assemble_control_volume_system(
        storage_coefficients=storage_coefficients,
        old_values=old_temperature,
        grid=grid,
        interface_coefficients=interface_conductivity,
        surface_transfer_coefficient=CONVECTIVE_HEAT_COEFF,
        boundary_value=room_temperature,
        time_step=time_step,
    )
    return solve_tridiagonal(*system)


def advance_moisture(
    old_moisture: np.ndarray,
    current_time: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    time_step: float,
) -> np.ndarray:
    """在当前 Picard 迭代的温度和含水率下推进一个水分时间步。"""
    diffusivity = moisture_diffusivity(
        reference_moisture, reference_temperature
    )
    interface_diffusivity = harmonic_mean(
        diffusivity[:-1], diffusivity[1:]
    )
    storage_coefficients = grid.area_factors / time_step
    room_moisture = linear_interp(
        current_time, boundary.times, boundary.moisture
    )

    system = assemble_control_volume_system(
        storage_coefficients=storage_coefficients,
        old_values=old_moisture,
        grid=grid,
        interface_coefficients=interface_diffusivity,
        surface_transfer_coefficient=CONVECTIVE_MASS_COEFF,
        boundary_value=room_moisture,
        time_step=time_step,
    )
    return solve_tridiagonal(*system)


def solve_problem_two(
    end_time: float = END_TIME,
    time_step: float = TIME_STEP,
    radial_intervals: int = RADIAL_INTERVALS,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, dict]:
    """求解问题二，返回时间、节点、温度场、含水率场和诊断信息。"""
    if end_time <= 0 or time_step <= 0:
        raise ValueError("终止时间和时间步长必须为正。")

    boundary = read_drying_boundary(INPUT_FILE)
    grid = build_radial_grid(radial_intervals)
    step_count = int(round(end_time / time_step))
    if not np.isclose(step_count * time_step, end_time):
        raise ValueError("终止时间必须是时间步长的整数倍。")

    temperature = np.full(len(grid.nodes), INITIAL_TEMPERATURE, dtype=float)
    moisture = np.full(len(grid.nodes), INITIAL_MOISTURE, dtype=float)

    times = []
    temperature_field = []
    moisture_field = []
    iteration_counts = []
    max_temperature_error = 0.0
    max_moisture_error = 0.0

    for step in range(1, step_count + 1):
        current_time = step * time_step
        temperature_iterate = temperature.copy()
        moisture_iterate = moisture.copy()
        converged = False

        for iteration in range(1, MAX_ITERATIONS + 1):
            # 先更新温度，再用新温度更新扩散系数，属于同一时间步内的
            # Gauss–Seidel 型 Picard 迭代。
            temperature_candidate = advance_temperature(
                old_temperature=temperature,
                current_time=current_time,
                grid=grid,
                boundary=boundary,
                reference_moisture=moisture_iterate,
                time_step=time_step,
            )
            moisture_candidate = advance_moisture(
                old_moisture=moisture,
                current_time=current_time,
                grid=grid,
                boundary=boundary,
                reference_temperature=temperature_candidate,
                reference_moisture=moisture_iterate,
                time_step=time_step,
            )

            new_temperature = (
                RELAXATION_FACTOR * temperature_candidate
                + (1.0 - RELAXATION_FACTOR) * temperature_iterate
            )
            new_moisture = (
                RELAXATION_FACTOR * moisture_candidate
                + (1.0 - RELAXATION_FACTOR) * moisture_iterate
            )

            temperature_error = np.max(
                np.abs(new_temperature - temperature_iterate)
                / (1.0 + np.abs(new_temperature))
            )
            moisture_error = np.max(
                np.abs(new_moisture - moisture_iterate)
                / (1.0 + np.abs(new_moisture))
            )

            temperature_iterate = new_temperature
            moisture_iterate = new_moisture
            max_temperature_error = max(
                max_temperature_error, float(temperature_error)
            )
            max_moisture_error = max(
                max_moisture_error, float(moisture_error)
            )

            if max(temperature_error, moisture_error) < CONVERGENCE_TOL:
                converged = True
                break

        if not converged:
            raise RuntimeError(
                f"{current_time:.0f} s 的热湿 Picard 迭代未收敛，"
                f"温度误差={temperature_error:.3e}，"
                f"含水率误差={moisture_error:.3e}。"
            )

        temperature = temperature_iterate
        moisture = moisture_iterate

        if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(moisture)):
            raise FloatingPointError(f"{current_time:.0f} s 出现非有限场变量。")
        if np.min(moisture) <= 0.0:
            raise FloatingPointError(f"{current_time:.0f} s 出现非正含水率。")

        times.append(current_time)
        temperature_field.append(temperature.copy())
        moisture_field.append(moisture.copy())
        iteration_counts.append(iteration)

    diagnostics = {
        "max_picard_iterations": int(max(iteration_counts)),
        "mean_picard_iterations": float(np.mean(iteration_counts)),
        "max_temperature_iteration_error": max_temperature_error,
        "max_moisture_iteration_error": max_moisture_error,
    }
    return (
        np.asarray(times),
        grid.nodes,
        np.asarray(temperature_field),
        np.asarray(moisture_field),
        diagnostics,
    )


def write_result_workbook(
    times: np.ndarray,
    nodes: np.ndarray,
    temperature_field: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    """保留附件三模板结构，将问题二结果写入 result2.xlsx。"""
    workbook = load_workbook(OUTPUT_FILE)
    if len(workbook.worksheets) < 2:
        raise ValueError("result2.xlsx 必须包含温度和水分浓度两个工作表。")

    for worksheet, field in zip(
        workbook.worksheets[:2], (temperature_field, moisture_field)
    ):
        # 删除模板示例行，但保留工作表、表头和基础格式。
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)

        worksheet.cell(1, 1).value = "时间\\到药材中心的距离"
        worksheet.cell(1, 1).number_format = "@"

        for column, radius in enumerate(nodes, start=2):
            worksheet.cell(1, column).value = round(float(radius * 100.0), 1)
            worksheet.cell(1, column).number_format = "0.0"

        for row, current_time in enumerate(times, start=2):
            worksheet.cell(row, 1).value = int(round(float(current_time)))
            worksheet.cell(row, 1).number_format = "0"
            for column, value in enumerate(field[row - 2], start=2):
                worksheet.cell(row, column).value = round(float(value), 4)
                worksheet.cell(row, column).number_format = "0.0000"

    workbook.save(OUTPUT_FILE)


def main() -> None:
    times, nodes, temperature_field, moisture_field, diagnostics = solve_problem_two()
    write_result_workbook(times, nodes, temperature_field, moisture_field)

    print(f"问题二计算完成，结果已写入：{OUTPUT_FILE}")
    print(
        f"场尺寸：{temperature_field.shape[0]}×{temperature_field.shape[1]}；"
        f"最大/平均 Picard 迭代次数："
        f"{diagnostics['max_picard_iterations']}/"
        f"{diagnostics['mean_picard_iterations']:.2f}"
    )
    print(
        f"温度范围：{temperature_field.min():.4f}–"
        f"{temperature_field.max():.4f} ℃；"
        f"含水率范围：{moisture_field.min():.4f}–"
        f"{moisture_field.max():.4f} kg/kg"
    )


if __name__ == "__main__":
    main()
