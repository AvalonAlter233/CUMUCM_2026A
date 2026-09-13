from __future__ import annotations

import sys
from functools import lru_cache
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib as mpl
import matplotlib.font_manager as font_manager
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
from openpyxl import load_workbook

if (Path(__file__).parent / "问题1_求解_改.py").is_file():
    import 问题1_求解_改 as problem1_solver
else:
    import 问题1_求解 as problem1_solver


nature_figure_scripts = Path.home() / ".codex" / "skills" / "nature-figure" / "scripts"
if str(nature_figure_scripts) not in sys.path:
    sys.path.insert(0, str(nature_figure_scripts))
from audit_panel_alignment import require_matplotlib_panel_alignment


project_root = Path(__file__).resolve().parent
attachment_one = project_root / "附件" / "附件1.xlsx"
result_one = project_root / "附件" / "附件3" / "result1.xlsx"
figure_dir = project_root / "figures" / "问题一"

cjk_font_candidates = (
    "STSong", "SimSun", "Source Han Serif SC", "Noto Serif CJK SC",
)
installed_fonts = {font.name for font in font_manager.fontManager.ttflist}
available_cjk_fonts = [
    name for name in cjk_font_candidates if name in installed_fonts
]
if not available_cjk_fonts:
    print("提示：未检测到中文字体，图中中文可能显示为方框。")
heading_font = "STZhongsong" if "STZhongsong" in installed_fonts else available_cjk_fonts[0]

mpl.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": [*available_cjk_fonts, "Times New Roman", "Times", "DejaVu Serif"],
        "mathtext.fontset": "stix",
        "axes.unicode_minus": False,
        "font.size": 8,
        "axes.titlesize": 9.5,
        "axes.titleweight": "bold",
        "axes.labelsize": 8,
        "axes.labelweight": "bold",
        "legend.fontsize": 7.5,
        "xtick.labelsize": 7.5,
        "ytick.labelsize": 7.5,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.linewidth": 0.7,
        "legend.frameon": False,
        "figure.dpi": 120,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    }
)

snapshot_times = (100.0, 600.0, 1200.0, 1800.0)
export_dpi = 600
pad_inches = 0.06
figure_formats = ("png",)

reference_palette = ("#44757A", "#452A3D", "#D44C3C", "#EED5B7")
color_teal, color_plum, color_coral, color_sand = reference_palette
color_center = color_teal
color_average = color_plum
color_surface = color_coral
profile_colors = (color_teal, color_plum, "#B88E63", color_coral)

new_figure_bases = (
    "图1_实测边界条件",
    "图2_中心平均表面响应",
    "图3_热湿场时空演化",
    "图4_热湿径向剖面",
    "图5_数值验证",
)


def add_initial_state(
    times: np.ndarray, field: np.ndarray, initial_value: float
) -> tuple[np.ndarray, np.ndarray]:
    if field.ndim != 2 or field.shape[0] != len(times):
        raise ValueError("时间轴与场数组的第一维不一致。")
    initial_row = np.full((1, field.shape[1]), float(initial_value))
    return np.concatenate(([0.0], times), axis=0), np.vstack((initial_row, field))


def radial_error(reference: np.ndarray, refined: np.ndarray) -> np.ndarray:
    # 先把两份场数据的形状对一下，对齐了再按时刻往下比较。
    # 每个时刻留下径向上差得最大的那个数，后面画误差曲线用。
    if reference.shape != refined.shape or reference.ndim != 2:
        raise ValueError("待比较场必须是形状相同的二维数组。")
    return np.max(np.abs(reference - refined), axis=1)


def check_output_axes(
    reference_times: np.ndarray,
    refined_times: np.ndarray,
    reference_radii: np.ndarray,
    refined_radii: np.ndarray,
) -> None:
    # 两份结果先对一下时间和位置，看看是不是用的同一套输出点。
    if not np.allclose(reference_times, refined_times):
        raise ValueError("细化结果的时间轴与基准结果不一致。")
    if not np.allclose(reference_radii, refined_radii):
        raise ValueError("细化结果的径向输出网格不一致。")


def snapshot_index(times: np.ndarray, snapshot_time: float) -> int:
    matches = np.flatnonzero(np.isclose(times, snapshot_time))
    if len(matches) != 1:
        raise ValueError(f"结果中未找到唯一的 {snapshot_time:g} s 快照。")
    return int(matches[0])


def read_boundary_data() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    # 环境温度和水分都在这张表里
    book = load_workbook(attachment_one, data_only=True, read_only=True)
    sheet = book.active
    rows = [
        row for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] < 3:
        raise ValueError("附件一必须包含时间、温度和水分浓度三列。")
    return data[:, 0], data[:, 1], data[:, 2]


def read_results() -> tuple[
    np.ndarray, np.ndarray, np.ndarray, np.ndarray
]:
    workbook = load_workbook(result_one, data_only=True, read_only=True)
    if len(workbook.worksheets) < 2:
        raise ValueError("result1.xlsx 必须包含温度和含水率两个工作表。")
    temperature_sheet, moisture_sheet = workbook.worksheets[:2]

    def read_sheet(sheet) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
        # 时间、半径和数值分开整理好
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        rows = [
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        ]
        times = np.asarray([row[0] for row in rows], dtype=float)
        radii_cm = np.asarray(list(header)[1:], dtype=float)
        field = np.asarray([row[1:] for row in rows], dtype=float)
        return times, radii_cm, field

    times_s, radii_cm, temperature_field = read_sheet(temperature_sheet)
    moisture_times_s, moisture_radii_cm, moisture_field = read_sheet(moisture_sheet)
    check_output_axes(
        times_s,
        moisture_times_s,
        radii_cm,
        moisture_radii_cm,
    )
    return times_s, radii_cm, temperature_field, moisture_field


@lru_cache(maxsize=1)
def load_plot_data() -> dict[str, object]:
    # 后面每张图只取各需，余者仍然留在这份数据里。
    env_seconds, air_temp, air_water = read_boundary_data()
    sheet_times, sheet_radii, sheet_temp, sheet_water = (
        read_results()
    )

    reference = problem1_solver.simulate_question1()
    ref_radii = reference.output_radii * 100.0
    check_output_axes(
        reference.times,
        sheet_times,
        ref_radii,
        sheet_radii,
    )
    round_tol = 5.1e-5
    if not np.allclose(
        reference.temperature_field,
        sheet_temp,
        rtol=0.0,
        atol=round_tol,
    ):
        raise ValueError("求解器温度场与 result1.xlsx 的四位小数结果不一致。")
    if not np.allclose(
        reference.moisture_field,
        sheet_water,
        rtol=0.0,
        atol=round_tol,
    ):
        raise ValueError("求解器含水率场与 result1.xlsx 的四位小数结果不一致。")

    fine_space = problem1_solver.simulate_question1(cell_count=320)
    fine_time = problem1_solver.simulate_question1(time_step=0.5)
    for refined in (fine_space, fine_time):
        check_output_axes(
            reference.times,
            refined.times,
            reference.output_radii,
            refined.output_radii,
        )

    return {
        "boundary_times_s": env_seconds,
        "boundary_temperature": air_temp,
        "boundary_moisture": air_water,
        "times_s": reference.times,
        "radii_cm": ref_radii,
        "temperature_field": reference.temperature_field,
        "moisture_field": reference.moisture_field,
        "average_temperature": reference.average_temperature,
        "average_moisture": reference.average_moisture,
        "diagnostics": reference.diagnostics,
        "fine_space_temperature": fine_space.temperature_field,
        "fine_space_moisture": fine_space.moisture_field,
        "fine_time_temperature": fine_time.temperature_field,
        "fine_time_moisture": fine_time.moisture_field,
    }


def save_figure(
    figure: plt.Figure,
    base_name: str,
    exclude_axes: list[plt.Axes] | None = None,
) -> None:
    # 检查走完并且通过了，才把图片保存下来
    figure_dir.mkdir(parents=True, exist_ok=True)
    figure.canvas.draw()
    stem = figure_dir / base_name
    with TemporaryDirectory(prefix="problem1-figure-qa-") as tmp_dir:
        require_matplotlib_panel_alignment(
            figure,
            json_out=Path(tmp_dir) / f"{base_name}.alignment.json",
            exclude_axes=exclude_axes or [],
            row_groups=getattr(figure, "_alignment_row_groups", None),
            tolerance_pt=1.5,
            gutter_tolerance_pt=1.5,
            require_panel_labels=False,
            strict=True,
        )
    save_opts = {"bbox_inches": "tight", "pad_inches": pad_inches, "facecolor": "white"}
    figure.savefig(stem.with_suffix(".png"), dpi=export_dpi, **save_opts)
    plt.close(figure)


def close_figure(figure: plt.Figure) -> None:
    plt.close(figure)

# 下文还是比较好理解的，不注释了
def style_axis(axis: plt.Axes, *, show_grid: bool = True) -> None:
    axis.set_axisbelow(True)
    if show_grid:
        axis.grid(
            True,
            which="major",
            color="#AAA5A8",
            linestyle="--",
            linewidth=0.55,
            alpha=0.70,
        )
    else:
        axis.grid(False)
    axis.tick_params(length=3, width=0.7)
    for label in (*axis.get_xticklabels(), *axis.get_yticklabels()):
        if not any("\u4e00" <= character <= "\u9fff" for character in label.get_text()):
            label.set_fontfamily("Times New Roman")
    for text in (axis.title, axis.xaxis.label, axis.yaxis.label):
        text.set_fontfamily(heading_font)
        text.set_fontweight("bold")
    for spine in (axis.spines["left"], axis.spines["bottom"]):
        spine.set_color("#51474D")
        spine.set_linewidth(0.75)


def set_title(figure: plt.Figure, title: str, y: float = 0.97) -> None:
    text = figure.suptitle(title, x=0.5, y=y, fontsize=12, fontweight="bold")
    text.set_fontfamily(heading_font)


def plot_boundary(data: dict[str, object]) -> plt.Figure:
    # 环境数据已在前面读好
    # 先把两边的曲线画出来，再补齐标题和坐标等等。
    env_hours = np.asarray(data["boundary_times_s"], dtype=float) / 3600.0
    air_temp = np.asarray(data["boundary_temperature"], dtype=float)
    air_water = np.asarray(data["boundary_moisture"], dtype=float)
    fig, axs = plt.subplots(1, 2, figsize=(7.0, 2.6), sharex=True)
    series = (
        (axs[0], air_temp, "T∞ / °C", "温度边界", color_center, "a"),
        (
            axs[1],
            air_water,
            "C∞ / kg/kg",
            "水分边界",
            color_surface,
            "b",
        ),
    )
    for ax, values, ylabel, title, color, label in series:
        ax.plot(env_hours, values, color=color, linewidth=1.7)
        ax.set_title(title, loc="left", pad=7, fontsize=9.2, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("时间 / h")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0.0, 0.5)
        style_axis(ax)
    set_title(fig, "实测边界条件", y=0.96)
    fig.subplots_adjust(wspace=0.30, bottom=0.21, top=0.73, left=0.09, right=0.98)
    return fig


def plot_positions(
    axis: plt.Axes,
    times_h: np.ndarray,
    field: np.ndarray,
    average: np.ndarray,
    ylabel: str,
    title: str,
) -> None:
    # 先把中心、平均和表面三条线画完。
    # 颜色和线型按各自的标签来
    for values, label, color, linestyle in (
        (field[:, 0], "中心", color_center, "-"),
        (average, "体积平均", color_average, "--"),
        (field[:, -1], "表面", color_surface, "-"),
    ):
        axis.plot(
            times_h,
            values,
            color=color,
            linewidth=1.55 if label != "体积平均" else 1.25,
            linestyle=linestyle,
            label=label,
        )
    axis.set_title(title, loc="left", pad=7, fontsize=9.2, fontweight="bold", fontfamily=heading_font)
    axis.set_xlabel("时间 / h")
    axis.set_ylabel(ylabel)
    axis.set_xlim(times_h[0], times_h[-1] + 0.045)
    style_axis(axis)


def plot_response(data: dict[str, object]) -> plt.Figure:
    times_h = np.asarray(data["times_s"], dtype=float) / 3600.0
    temp_map = np.asarray(data["temperature_field"], dtype=float)
    water_map = np.asarray(data["moisture_field"], dtype=float)
    temp_mean = np.asarray(data["average_temperature"], dtype=float)
    water_mean = np.asarray(data["average_moisture"], dtype=float)
    fig, axs = plt.subplots(1, 2, figsize=(7.0, 2.8), sharex=True)
    plot_positions(
        axs[0], times_h, temp_map, temp_mean, "温度 / °C", "温度"
    )
    plot_positions(
        axs[1],
        times_h,
        water_map,
        water_mean,
        "干基含水率 / kg/kg",
        "含水率",
    )
    handles, labels = axs[0].get_legend_handles_labels()
    set_title(fig, "药材内部温湿响应")
    fig.legend(
        handles,
        labels,
        ncol=3,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.86),
        handlelength=1.8,
        columnspacing=1.5,
    )
    fig.subplots_adjust(wspace=0.31, bottom=0.21, top=0.68, left=0.09, right=0.98)
    return fig


def plot_field(data: dict[str, object]) -> plt.Figure:
    # 把不同时间、不同位置上的数放到同一块图里展开。
    # 数据先按时间和半径排好，后面看颜色的时候，就能从两个方向看变化。
    times_s = np.asarray(data["times_s"], dtype=float)
    radii_cm = np.asarray(data["radii_cm"], dtype=float)
    temp_times, temp_map = add_initial_state(
        times_s,
        np.asarray(data["temperature_field"], dtype=float),
        problem1_solver.initial_temperature,
    )
    water_times, water_map = add_initial_state(
        times_s,
        np.asarray(data["moisture_field"], dtype=float),
        problem1_solver.initial_moisture,
    )
    if not np.allclose(temp_times, water_times):
        raise ValueError("温度场和含水率场的时空网格不一致。")
    times_h = temp_times / 3600.0
    fig, axs = plt.subplots(1, 2, figsize=(7.0, 3.0), sharex=True, sharey=True)
    cbars = []
    panels = (
        (axs[0], temp_map, "magma", "温度场", "°C", "a"),
        (
            axs[1],
            water_map,
            "viridis",
            "含水率场",
            "kg/kg",
            "b",
        ),
    )
    for ax, field, cmap, title, bar_label, label in panels:
        mesh = ax.pcolormesh(
            radii_cm,
            times_h,
            field,
            shading="auto",
            cmap=cmap,
            vmin=float(np.min(field)),
            vmax=float(np.max(field)),
        )
        cbar = fig.colorbar(mesh, ax=ax, fraction=0.046, pad=0.04)
        cbar.set_label(bar_label, rotation=270, labelpad=13)
        cbar.ax.yaxis.label.set_fontfamily(heading_font)
        cbar.ax.yaxis.label.set_fontweight("bold")
        for tick in cbar.ax.get_yticklabels():
            tick.set_fontfamily("Times New Roman")
        cbar.ax.grid(False)
        cbars.append(cbar.ax)
        ax.set_title(title, loc="left", pad=7, fontsize=9.2, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("半径 r / cm")
        ax.set_ylabel("时间 / h")
        ax.set_xlim(radii_cm[0], radii_cm[-1])
        ax.set_ylim(times_h[0], times_h[-1])
        style_axis(ax, show_grid=False)
    set_title(fig, "热湿场时空演化")
    fig._alignment_exclude_axes = cbars
    fig._alignment_row_groups = [["a", "b"]]
    fig.subplots_adjust(wspace=0.28, bottom=0.17, top=0.75, left=0.09, right=0.92)
    return fig


def plot_profiles(data: dict[str, object]) -> plt.Figure:
    times_s = np.asarray(data["times_s"], dtype=float)
    radii_cm = np.asarray(data["radii_cm"], dtype=float)
    temp_map = np.asarray(data["temperature_field"], dtype=float)
    water_map = np.asarray(data["moisture_field"], dtype=float)
    fig, axs = plt.subplots(1, 2, figsize=(7.0, 3.15), sharex=True)
    lines = []
    panels = (
        (axs[0], temp_map, "温度 / °C", "温度", "a"),
        (
            axs[1],
            water_map,
            "干基含水率 / kg/kg",
            "含水率",
            "b",
        ),
    )
    for ax, field, ylabel, title, label in panels:
        for color, snapshot_time in zip(profile_colors, snapshot_times):
            (line,) = ax.plot(
                radii_cm,
                field[snapshot_index(times_s, snapshot_time)],
                color=color,
                linewidth=1.65,
                label=f"{snapshot_time:g} s",
            )
            if ax is axs[0]:
                lines.append(line)
        ax.set_title(title, loc="left", pad=7, fontsize=9.2, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("半径 r / cm")
        ax.set_ylabel(ylabel)
        ax.set_xlim(radii_cm[0], radii_cm[-1])
        style_axis(ax)
    set_title(fig, "关键时刻径向剖面")
    fig.legend(
        handles=lines,
        labels=[f"{value:g} s" for value in snapshot_times],
        ncol=4,
        loc="upper center",
        bbox_to_anchor=(0.52, 0.87),
        handlelength=1.9,
        columnspacing=1.3,
    )
    fig.subplots_adjust(wspace=0.30, bottom=0.17, top=0.70, left=0.10, right=0.98)
    return fig


def collect_errors(data: dict[str, object]) -> dict[str, np.ndarray]:
    # 空间加密的结果和时间加密的结果分别同参考结果比较。
    # 这几串误差存放，后面画到验证图时再取。
    ref_temp = np.asarray(data["temperature_field"], dtype=float)
    ref_water = np.asarray(data["moisture_field"], dtype=float)
    return {
        "grid_temperature": radial_error(
            ref_temp,
            np.asarray(data["fine_space_temperature"], dtype=float),
        ),
        "grid_moisture": radial_error(
            ref_water,
            np.asarray(data["fine_space_moisture"], dtype=float),
        ),
        "time_temperature": radial_error(
            ref_temp,
            np.asarray(data["fine_time_temperature"], dtype=float),
        ),
        "time_moisture": radial_error(
            ref_water,
            np.asarray(data["fine_time_moisture"], dtype=float),
        ),
    }


def plot_validation(data: dict[str, object]) -> plt.Figure:
    # 先把需要的数取齐，再分别放到对应的位置
    times_h = np.asarray(data["times_s"], dtype=float) / 3600.0
    errors = collect_errors(data)
    figure, axes = plt.subplots(2, 2, figsize=(7.0, 4.55), sharex=True)
    panels = (
        (
            axes[0, 0],
            errors["grid_temperature"],
            "空间加密｜温度",
            "最大径向误差 / °C",
            "N = 160 vs N = 320",
            "a",
        ),
        (
            axes[0, 1],
            errors["grid_moisture"],
            "空间加密｜含水率",
            "最大径向误差 / kg/kg",
            "N = 160 vs N = 320",
            "b",
        ),
        (
            axes[1, 0],
            errors["time_temperature"],
            "时间加密｜温度",
            "最大径向误差 / °C",
            "Δt = 1.0 s vs 0.5 s",
            "c",
        ),
        (
            axes[1, 1],
            errors["time_moisture"],
            "时间加密｜含水率",
            "最大径向误差 / kg/kg",
            "Δt = 1.0 s vs 0.5 s",
            "d",
        ),
    )
    for axis, values, title, ylabel, comparison, label in panels:
        axis.semilogy(times_h, values, color=color_teal, linewidth=1.45)
        axis.yaxis.set_major_formatter(mticker.FuncFormatter(lambda value, _: f"{value:.0e}"))
        axis.set_title(title, loc="left", pad=7, fontsize=9.2, fontweight="bold", fontfamily=heading_font)
        axis.set_xlabel("时间 / h")
        axis.set_ylabel(ylabel)
        axis.set_xlim(times_h[0], times_h[-1])
        axis.text(
            0.02,
            0.96,
            comparison.replace(" vs ", " → "),
            transform=axis.transAxes,
            fontsize=6.4,
            ha="left",
            va="top",
            color=color_teal,
        )
        style_axis(axis)
    set_title(figure, "数值收敛验证")
    figure.subplots_adjust(
        wspace=0.34,
        hspace=0.38,
        bottom=0.14,
        top=0.84,
        left=0.10,
        right=0.98,
    )
    return figure


old_figure_names = (
    "图1_附件一烘房温度曲线.png",
    "图2_附件一烘房水分浓度曲线.png",
    "图3_附件一边界采样间隔.png",
    "图4_温度径向剖面.png",
    "图5_水分径向剖面.png",
    "图6_中心与表面轨迹.png",
    "图7_温度场时空分布.png",
    "图8_水分场时空分布.png",
    "图9_空间非均匀性.png",
)


def clear_old_figures() -> None:
    figure_dir.mkdir(parents=True, exist_ok=True)
    for basename in old_figure_names:
        target = figure_dir / basename
        if target.parent != figure_dir:
            raise ValueError(f"拒绝清理图集目录之外的路径：{target}")
        if target.exists():
            target.unlink()


def export_figures(data: dict[str, object]) -> None:
    # 依次生成图片
    builders = (
        (plot_boundary, new_figure_bases[0]),
        (plot_response, new_figure_bases[1]),
        (plot_field, new_figure_bases[2]),
        (plot_profiles, new_figure_bases[3]),
        (plot_validation, new_figure_bases[4]),
    )
    for builder, base_name in builders:
        fig = builder(data)
        save_figure(
            fig,
            base_name,
            exclude_axes=list(getattr(fig, "_alignment_exclude_axes", [])),
        )

    missing_outputs = [
        figure_dir / f"{base_name}.{suffix}"
        for base_name in new_figure_bases
        for suffix in figure_formats
        if not (figure_dir / f"{base_name}.{suffix}").exists()
    ]
    if missing_outputs:
        raise FileNotFoundError(f"新图集导出不完整：{missing_outputs}")
    clear_old_figures()
    print(
        f"问题一绘图完成：生成 {len(new_figure_bases)} 张图，"
        f"每张包含 {len(figure_formats)} 种格式，输出目录：{figure_dir}"
    )


def main() -> None:
    export_figures(load_plot_data())

# 好耶ww
if __name__ == "__main__":
    main()
