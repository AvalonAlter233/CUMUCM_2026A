from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Callable, NamedTuple

import numpy as np
from openpyxl import load_workbook

if (Path(__file__).parent / "问题3_求解_改.py").is_file():
    import 问题3_求解_改 as problem3
else:
    import 问题3_求解 as problem3

# 单击此处以添加标题（）
project_root = Path(__file__).resolve().parent
radius_file = project_root / "附件" / "附件2.xlsx"
output_file = project_root / "附件" / "附件3" / "result4.xlsx"
diagnostics_file = project_root / "附件" / "附件3" / "result4_diagnostics.json"
internal_field_file = project_root / "附件" / "附件3" / "result4_internal_field.npz"

initial_radius = 0.02
convective_heat_coeff = 25.0
convective_mass_coeff = 8.0e-7
initial_temperature = 28.0
initial_moisture = 2.55
default_critical_moisture = 0.15

default_intervals = 320
default_time_step = 30.0
default_report_interval = 60.0
max_simulation_time = 6.0 * 24.0 * 3600.0

convergence_tol = 1.0e-9
max_iterations = 200
relaxation_factor = 0.6


class RadiusHistory(NamedTuple):

    times: np.ndarray
    radii: np.ndarray


class ReferenceGrid(NamedTuple):

    faces: np.ndarray
    centers: np.ndarray
    volumes: np.ndarray
    spacing: float


class MaterialProperties(NamedTuple):

    density: Callable[[np.ndarray], np.ndarray]
    heat_capacity: Callable[[np.ndarray], np.ndarray]
    thermal_conductivity: Callable[[np.ndarray], np.ndarray]
    moisture_diffusivity: Callable[[np.ndarray, np.ndarray], np.ndarray]


class InternalFieldData(NamedTuple):

    times: np.ndarray
    xi_centers: np.ndarray
    moisture: np.ndarray
    surface_radii: np.ndarray


def read_radius_data(path: Path) -> RadiusHistory:
    # 这一问多了一份半径数据，所以先把这张表也读进来。
    # 表里的数取出来以后，把单位换成计算要用的单位，后面就一直用换算好的这份。
    book = load_workbook(path, data_only=True, read_only=True)
    rows = [
        row
        for row in book.active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] < 2 or data.shape[0] < 2:
        raise ValueError("附件二缺少有效半径数据。")
    if not np.all(np.diff(data[:, 0]) > 0.0):
        raise ValueError("附件二时间必须严格递增。")
    radii = data[:, 1] / 100.0
    if np.any(radii <= 0.0):
        raise ValueError("附件二半径必须为正数。")
    if np.any(np.diff(radii) > 1.0e-12):
        raise ValueError("附件二半径必须单调不增。")
    return RadiusHistory(data[:, 0], radii)


def radius_value(current_time: float, history: RadiusHistory) -> float:
    # 算到哪个时刻，就来这里拿那个时刻的半径。
    # 中间的时刻用插值接上，走过最后一个采样时刻以后，就继续用末尾的那个半径值。
    # 这里每次只负责给当前这一刻的半径，外面拿到了再继续做后面的计算。
    return float(
        np.interp(
            current_time,
            history.times,
            history.radii,
            left=history.radii[0],
            right=history.radii[-1],
        )
    )


def density(moisture: np.ndarray) -> np.ndarray:
    return 760.0 + 90.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    return 1850.0 + 2150.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    return 0.12 + 0.20 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    c_safe = np.maximum(moisture, 1.0e-12)
    temp_k = temperature_celsius + 273.15
    return (
        4.2e-4
        * np.exp(-0.30 / c_safe)
        * np.exp(-3850.0 / temp_k)
    )


def material_properties(name: str) -> MaterialProperties:
    # 这一次选哪一组物性，先在这里确定好。
    # 几个物性函数一起交出去，后面的温度和水分计算就照着选中的这一组来调用。
    # 先选定这一趟用哪组，后面调用时就一直顺着这份配置来。
    # 要做别的对照时，再在入口处换成相应的那组就好了。
    if name == "appendix4":
        return MaterialProperties(
            density, heat_capacity, thermal_conductivity, moisture_diffusivity
        )
    if name == "appendix3":
        return MaterialProperties(
            problem3.density,
            problem3.heat_capacity,
            problem3.thermal_conductivity,
            problem3.moisture_diffusivity,
        )
    raise ValueError("物性模型必须为 appendix3 或 appendix4。")


def build_ref_grid(intervals: int) -> ReferenceGrid:
    # 参考位置先放在零到一之间，网格就照此范围来分。
    # 等实际计算时，再把当前半径带进去，所以半径变了以后也还是用这一份参考网格。
    if intervals < 4:
        raise ValueError("内部径向区间数至少为 4。")
    faces = np.linspace(0.0, 1.0, intervals + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    vol = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return ReferenceGrid(faces, centers, vol, 1.0 / intervals)


def assemble_ref_system(
    old_values: np.ndarray,
    storage_capacity: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 本轮的方程系数还是在这里一项一项凑起来。
    # 只是现在还要带上这一时刻的半径，内部的项和表面的项都处理完了，再合到一起去解。
    # 半径也是本轮要用的量，先带进来，再把整套系数接给求解那一步。
    if radius <= 0:
        raise ValueError("药材半径必须为正数。")
    storage = storage_capacity * grid.volumes / time_step
    coef_face = problem3.harmonic_mean(
        transport_coefficients[:-1], transport_coefficients[1:]
    )
    g_inner = (
        grid.faces[1:-1]
        * coef_face
        / (radius**2 * grid.spacing)
    )
    diag = storage.copy()
    diag[:-1] += g_inner
    diag[1:] += g_inner
    lower = -g_inner.copy()
    upper = -g_inner.copy()
    rhs = storage * old_values

    coef_last = max(float(transport_coefficients[-1]), 1.0e-30)
    half_dr = 0.5 * radius * grid.spacing
    h_eff = 1.0 / (
        1.0 / external_transfer_coefficient
        + half_dr / coef_last
    )
    g_surf = h_eff / radius
    diag[-1] += g_surf
    rhs[-1] += g_surf * external_value
    return lower, diag, upper, rhs


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    return problem3.solve_tridiagonal(
        lower, diagonal, upper, right_hand_side
    )


def advance_temperature(
    old_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_temperature: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
    properties: MaterialProperties,
) -> np.ndarray:
    # 这一轮选好的物性从这儿接着用，当前半径也一起带进方程里。
    # 先把温度这一步求出来，再把结果交回时间循环
    lin_sys = assemble_ref_system(
        old_temperature,
        properties.density(reference_moisture)
        * properties.heat_capacity(reference_moisture),
        properties.thermal_conductivity(reference_moisture),
        convective_heat_coeff,
        room_temperature,
        grid,
        radius,
        time_step,
    )
    return solve_tridiagonal(*lin_sys)


def advance_moisture(
    old_moisture: np.ndarray,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_moisture: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
    properties: MaterialProperties,
) -> np.ndarray:
    # 水分的系数也按本轮的温度和含水率来
    # 这一份求完以后先返回，外面还要看它和上一轮相比变了多少
    lin_sys = assemble_ref_system(
        old_moisture,
        np.ones_like(old_moisture),
        properties.moisture_diffusivity(
            reference_moisture, reference_temperature
        ),
        convective_mass_coeff,
        room_moisture,
        grid,
        radius,
        time_step,
    )
    return solve_tridiagonal(*lin_sys)


def sample_field(
    cell_values: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    physical_output_nodes: np.ndarray,
    grid: ReferenceGrid,
    radius: float,
) -> np.ndarray:
    if np.any(physical_output_nodes < 0.0) or np.any(
        physical_output_nodes >= radius
    ):
        raise ValueError("固定输出位置必须位于当前药材内部。")
    xi_out = physical_output_nodes / radius
    inner_values = np.interp(xi_out, grid.centers, cell_values)
    inner_values[0] = (9.0 * cell_values[0] - cell_values[1]) / 8.0

    coef_last = max(float(transport_coefficients[-1]), 1.0e-30)
    half_dr = 0.5 * radius * grid.spacing
    g_half = coef_last / half_dr
    val_surf = (
        g_half * cell_values[-1]
        + external_transfer_coefficient * external_value
    ) / (g_half + external_transfer_coefficient)
    return np.concatenate([inner_values, np.array([val_surf])])


def output_nodes() -> np.ndarray:
    # 固定输出位置先按这里的间隔排出来，后面取值时一直照着这一排找。
    # 表面的位置还会单独处理，先把这些固定位置准备好
    return np.arange(0.0, 0.011 + 0.0005, 0.001)


def output_headers(nodes: np.ndarray) -> list[float | str]:
    # 位置换成表头里要写的单位以后，再把表面这一列接到最后。
    # 整排标题先准备，后面写表时就可以直接顺着放
    return [round(float(node * 100.0), 1) for node in nodes] + ["药材表面"]


def solve_question4(
    internal_intervals: int = default_intervals,
    time_step: float = default_time_step,
    report_interval: float = default_report_interval,
    max_time: float = max_simulation_time,
    critical_moisture: float = default_critical_moisture,
    moving_radius: bool = True,
    property_model: str = "appendix4",
    capture_internal: bool = False,
    boundary: problem3.DryingRoomBoundary | None = None,
    radius_history: RadiusHistory | None = None,
) -> tuple[
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    dict,
    InternalFieldData | None,
]:
    # 主计算开始前，先把这次要用的数据和设置接好。
    # 半径记录、环境条件和网格都准备齐了，下面才开始一小步一小步往后算。
    # 这一步要准备的东西比前几问多一些额
    save_every = int(round(report_interval / time_step))
    n_steps = int(round(max_time / time_step))
    if save_every < 1 or not np.isclose(
        save_every * time_step, report_interval
    ):
        raise ValueError("输出间隔必须是内部时间步长的整数倍。")
    if n_steps < 1 or not np.isclose(n_steps * time_step, max_time):
        raise ValueError("最大时间必须是内部时间步长的整数倍。")

    if boundary is None:
        boundary = problem3.read_room_data(problem3.input_file)
    if radius_history is None:
        radius_history = read_radius_data(radius_file)

    mesh = build_ref_grid(internal_intervals)
    properties = material_properties(property_model)
    r_nodes = output_nodes()
    temp = np.full(internal_intervals, initial_temperature)
    water = np.full(internal_intervals, initial_moisture)

    time_log: list[float] = []
    water_log: list[np.ndarray] = []
    radius_log: list[float] = []
    inner_water_log: list[np.ndarray] = []
    iter_log: list[int] = []
    t_prev = 0.0
    c_max_prev = initial_moisture
    t_cross: float | None = None
    c_before: float | None = None
    c_after: float | None = None
    stock_init = float(np.dot(water, mesh.volumes))
    stock_now = stock_init
    outflow_sum = 0.0
    max_balance_err = 0.0
    center_controls_every_step = True
    radially_nonincreasing_every_step = True
    checked_steps = 0

    for step in range(1, n_steps + 1):
        t_now = step * time_step
        # 先看这一时刻该用多大的半径，再往下更新场变量。
        # 做收缩计算就从数据里取，做固定半径对照时则接着用起始的值。
        # 先把这一刻的大小定下来
        radius = (
            radius_value(t_now, radius_history)
            if moving_radius
            else float(radius_history.radii[0])
        )
        temp_air, water_air = problem3.boundary_value(
            t_now, boundary
        )
        temp_iter = temp.copy()
        water_iter = water.copy()

        # 半径选好以后，就在当前这个时间步里继续迭代温度和含水率。
        # 前后变化还不够小就接着算，等检查满足要求了，再往下面的步骤走。
        # 现在还是在同一个时间步里，先等温度和水分这一轮都更新好了，再看要不要接着迭代。
        # 说点废话，这一步走完以后，外面的时间循环才轮到下一步
        for iteration in range(1, max_iterations + 1):
            temp_trial = advance_temperature(
                temp,
                water_iter,
                temp_air,
                mesh,
                radius,
                time_step,
                properties,
            )
            water_trial = advance_moisture(
                water,
                temp_trial,
                water_iter,
                water_air,
                mesh,
                radius,
                time_step,
                properties,
            )
            temp_new = (
                relaxation_factor * temp_trial
                + (1.0 - relaxation_factor) * temp_iter
            )
            water_new = (
                relaxation_factor * water_trial
                + (1.0 - relaxation_factor) * water_iter
            )
            err_t = float(
                np.max(
                    np.abs(temp_new - temp_iter)
                    / (1.0 + np.abs(temp_new))
                )
            )
            err_c = float(
                np.max(
                    np.abs(water_new - water_iter)
                    / (1.0 + np.abs(water_new))
                )
            )
            temp_iter = temp_new
            water_iter = water_new
            if max(err_t, err_c) < convergence_tol:
                break
        else:
            raise RuntimeError(
                f"{t_now:.0f} s 未收敛："
                f"{err_t:.3e}, {err_c:.3e}"
            )

        temp = temp_iter
        water = water_iter
        iter_log.append(iteration)
        if not np.all(np.isfinite(water)) or np.min(water) <= 0.0:
            raise FloatingPointError(f"{t_now:.0f} s 含水率异常。")

        d_cell = properties.moisture_diffusivity(water, temp)
        d_last = max(float(d_cell[-1]), 1.0e-30)
        hm_eff = 1.0 / (
            1.0 / convective_mass_coeff
            + 0.5 * radius * mesh.spacing / d_last
        )
        flux_surf = (
            hm_eff
            / radius
            * (float(water[-1]) - water_air)
        )
        # 参考网格上的积分和边界流出量，也跟着这儿一起记下来。
        # 后面还会不断加上新的记录，等最后整理诊断时，再把累计的这些数一起带出去。
        stock_next = float(np.dot(water, mesh.volumes))
        balance_err = abs(
            stock_next - stock_now + time_step * flux_surf
        ) / stock_init
        max_balance_err = max(
            max_balance_err, balance_err
        )
        outflow_sum += time_step * flux_surf
        stock_now = stock_next

        c_nodes = sample_field(
            water,
            d_cell,
            convective_mass_coeff,
            water_air,
            r_nodes,
            mesh,
            radius,
        )
        c_max_now = max(
            float(np.max(water)), float(np.max(c_nodes))
        )
        axis_val = float(c_nodes[0])
        center_controls_every_step = center_controls_every_step and bool(
            c_max_now <= axis_val + 1.0e-12
        )
        radially_nonincreasing_every_step = (
            radially_nonincreasing_every_step
            and bool(np.all(np.diff(water) <= 1.0e-12))
            and bool(np.all(np.diff(c_nodes) <= 1.0e-12))
        )
        checked_steps += 1

        if (
            t_cross is None
            and c_max_prev >= critical_moisture
            and c_max_now < critical_moisture
        ):
            # 第一次越过阈值时，先把对应的时间记在这里。
            # 接下来还要保存输出结果，这个时间记录先留着，后面汇总的时候还会用到。
            # 先留住第一次经过时的，后面汇总时再把它拿出来和输出时刻一起看
            t_cross = problem3.threshold_time(
                t_prev,
                c_max_prev,
                t_now,
                c_max_now,
                critical_moisture,
            )
            c_before = c_max_prev
            c_after = c_max_now

        reached_on_report = False
        # 到了保存时刻，把当前含水率和当前半径成对记下来。
        # 含水率和半径要对应同一个时刻，所以存的时候就把它们一起放。
        # 额外的细网格记录也在这一段安排
        if step % save_every == 0:
            time_log.append(t_now)
            water_log.append(c_nodes.copy())
            radius_log.append(radius)
            if capture_internal:
                inner_water_log.append(water.copy())
            reached_on_report = c_max_now < critical_moisture

        t_prev = t_now
        c_max_prev = c_max_now
        if reached_on_report:
            break
    else:
        raise RuntimeError(f"{max_time / 3600:.1f} h 内未达到阈值。")

    water_field = np.asarray(water_log)
    times = np.asarray(time_log)
    radius_series = np.asarray(radius_log)
    temp_final = sample_field(
        temp,
        properties.thermal_conductivity(water),
        convective_heat_coeff,
        problem3.boundary_value(times[-1], boundary)[0],
        r_nodes,
        mesh,
        radius_series[-1],
    )
    # 这一趟算出来的时间、半径和检查量，现在整理一下。
    # 后面查看结果时就从这一份里面找，需要哪个数就取哪个数，相关的记录都放在一起。
    # 这份结果先打包好，再往外返回
    diagnostics = {
        "continuous_threshold_time_s": float(t_cross),
        "discrete_threshold_time_s": float(times[-1]),
        "maximum_moisture_before": float(c_before),
        "maximum_moisture_after": float(c_after),
        "radius_at_continuous_threshold_cm": 100.0
        * (
            radius_value(float(t_cross), radius_history)
            if moving_radius
            else float(radius_history.radii[0])
        ),
        "radius_at_discrete_threshold_cm": 100.0 * radius_series[-1],
        "moving_radius": moving_radius,
        "property_model": property_model,
        "center_controls_threshold": center_controls_every_step,
        "all_domain_checked_every_step": checked_steps == len(iter_log),
        "radially_nonincreasing_every_step": radially_nonincreasing_every_step,
        "moisture_balance_relative_imbalance": (
            problem3.balance_error(
                stock_init, stock_now, outflow_sum
            )
        ),
        "maximum_step_balance_relative_residual": (
            max_balance_err
        ),
        "initial_reference_moisture_integral": stock_init,
        "final_reference_moisture_integral": stock_now,
        "cumulative_reference_boundary_outflow": outflow_sum,
        "maximum_picard_iterations": int(max(iter_log)),
        "mean_picard_iterations": float(np.mean(iter_log)),
        "internal_intervals": internal_intervals,
        "reference_spacing": mesh.spacing,
        "time_step_s": time_step,
        "report_interval_s": report_interval,
        "critical_moisture": critical_moisture,
        "final_center_moisture": float(water_field[-1, 0]),
        "final_surface_moisture": float(water_field[-1, -1]),
        "final_center_temperature_c": float(temp_final[0]),
        "final_surface_temperature_c": float(temp_final[-1]),
        "plateau_temperature_c": boundary.plateau_temperature,
        "plateau_moisture_kgkg": boundary.plateau_moisture,
        "radius_data_final_time_s": float(radius_history.times[-1]),
        "radius_data_final_cm": 100.0 * float(radius_history.radii[-1]),
    }
    # 上文好多（）
    internal_field = None
    if capture_internal:
        internal_field = InternalFieldData(
            times,
            mesh.centers.copy(),
            np.asarray(inner_water_log),
            radius_series,
        )
    return (
        times,
        r_nodes,
        water_field,
        radius_series,
        diagnostics,
        internal_field,
    )


def write_internal_field(
    data: InternalFieldData,
    output_path: Path = internal_field_file,
) -> None:
    # 内部网格上的数据再单独存一份，后面需要细一点的数据时会用到。
    # 时间、位置和半径也跟着一起保存，之后读文件的时候就可以一并取出来。
    # 先把文件存好，后面要画细一点的内部场时，就从这里留下的数据接着来
    np.savez_compressed(
        output_path,
        times_s=data.times,
        xi_centers=data.xi_centers,
        moisture=data.moisture,
        surface_radii_m=data.surface_radii,
    )


def write_results(
    times: np.ndarray,
    nodes: np.ndarray,
    moisture_field: np.ndarray,
    output_path: Path = output_file,
) -> None:
    # 算好的数据现在照着模板往表格里放，顺序还是按输出位置来。
    # 前面的位置列写好以后，表面这一列也一起填上，最后把整个工作簿保存下来。
    # 前面每一列的位置都已经定好了，接下来照着表头往后填，填完这一行再接下一行
    book = load_workbook(output_path)
    sheet = book.worksheets[0]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    if sheet.max_column > 1:
        sheet.delete_cols(2, sheet.max_column - 1)
    sheet.cell(1, 1).value = "时间\\到药材中心的距离"
    sheet.cell(1, 1).number_format = "@"
    for column, header in enumerate(output_headers(nodes), start=2):
        sheet.cell(1, column).value = header
        sheet.cell(1, column).number_format = (
            "0.0" if isinstance(header, float) else "@"
        )
    for row, t_now in enumerate(times, start=2):
        sheet.cell(row, 1).value = int(round(float(t_now)))
        sheet.cell(row, 1).number_format = "0"
        for column, value in enumerate(moisture_field[row - 2], start=2):
            sheet.cell(row, column).value = float(value)
            sheet.cell(row, column).number_format = "0.0000"
    book.save(output_path)


def record_case(name: str, label: str, diagnostics: dict) -> dict:
    # 这一种情况先留一份自己的诊断，名字和换算好的时间也一起补齐
    record = dict(diagnostics)
    record["name"] = name
    record["label"] = label
    record["continuous_threshold_time_h"] = (
        diagnostics["continuous_threshold_time_s"] / 3600.0
    )
    record["discrete_threshold_time_h"] = (
        diagnostics["discrete_threshold_time_s"] / 3600.0
    )
    return record


def make_report(baseline_diagnostics: dict) -> dict:
    # 先搭出这份报告要放东西的几个位置，主结果现在先放进去
    return {
        "scope": (
            "移动参考域离散诊断与机制对照；附录三与附录四同时改变时，"
            "不能把总时间差全部解释为收缩效应。"
        ),
        "verification_complete": False,
        "baseline": record_case(
            "appendix4_moving_radius",
            "附录4物性 + 实测收缩半径",
            baseline_diagnostics,
        ),
        "mechanism_comparison": [],
        "numerical_refinement": [],
        "boundary_sensitivity": [],
    }


def write_report(report: dict) -> None:
    # 报告里的内容已经整理好了，这一步就是把整份记录写到文件里。
    # 写完以后随手关文件！
    with diagnostics_file.open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def main() -> None:
    # 先接好环境和半径两份数据，再开始这一问的主计算。
    # 结果出来以后还要写表、存内部场、整理诊断
    # 先跟着主流程往下看，额外的对照放在后面接着处理。
    parser = argparse.ArgumentParser(description="问题四移动边界干燥求解")
    parser.add_argument(
        "--verification",
        action="store_true",
        help="额外复算固定半径对照及空间、时间加密案例",
    )
    args = parser.parse_args()

    boundary = problem3.read_room_data(problem3.input_file)
    radius_history = read_radius_data(radius_file)
    times, nodes, water, radius_series, diagnostics, internal_field = (
        solve_question4(
            boundary=boundary,
            radius_history=radius_history,
            capture_internal=True,
        )
    )
    write_results(times, nodes, water)
    if internal_field is None:
        raise RuntimeError("未生成内部参考域结果。")
    write_internal_field(internal_field)
    print(f"问题四重算完成：{output_file}")
    print(
        f"连续阈值：{diagnostics['continuous_threshold_time_s'] / 3600:.4f} h；"
        f"首次 60 s 离散达标："
        f"{diagnostics['discrete_threshold_time_s'] / 3600:.4f} h"
    )
    print(
        f"达标时半径：{diagnostics['radius_at_discrete_threshold_cm']:.4f} cm；"
        f"输出尺寸：{water.shape[0]}×{water.shape[1]}"
    )
    print(
        "参考域 C 方程累计相对收支不平衡："
        f"{diagnostics['moisture_balance_relative_imbalance']:.3e}"
    )

    # 懒得写了
    report = make_report(diagnostics)
    if not args.verification:
        write_report(report)
        print(f"基准诊断已保存（未运行完整验证）：{diagnostics_file}")
        return

    print("复算机制对照：附录4物性 + 固定半径")
    _, _, _, _, fixed_stats, _ = solve_question4(
        moving_radius=False,
        boundary=boundary,
        radius_history=radius_history,
    )
    report["mechanism_comparison"].append(
        record_case(
            "appendix4_fixed_radius", "附录4物性 + 固定半径", fixed_stats
        )
    )

    print("复算机制对照：附录3物性 + 固定半径（问题三）")
    _, _, _, q3_stats = problem3.solve_question3(
        boundary=boundary
    )
    report["mechanism_comparison"].append(
        record_case(
            "appendix3_fixed_radius",
            "附录3物性 + 固定半径（问题三）",
            q3_stats,
        )
    )

    print("复算机制对照：附录3物性 + 实测收缩半径")
    _, _, _, _, appendix3_moving_diagnostics, _ = solve_question4(
        moving_radius=True,
        property_model="appendix3",
        boundary=boundary,
        radius_history=radius_history,
    )
    report["mechanism_comparison"].append(
        record_case(
            "appendix3_moving_radius",
            "附录3物性 + 实测收缩半径",
            appendix3_moving_diagnostics,
        )
    )

    mesh_cases = [
        ("space_refined", "空间加密 N=640, Δt=30 s", 640, 30.0),
        ("time_refined", "时间加密 N=320, Δt=15 s", 320, 15.0),
    ]
    for name, label, intervals, dt_fine in mesh_cases:
        print(f"复算数值加密：{label}")
        _, _, _, _, refined_stats, _ = solve_question4(
            internal_intervals=intervals,
            time_step=dt_fine,
            boundary=boundary,
            radius_history=radius_history,
        )
        report["numerical_refinement"].append(
            record_case(name, label, refined_stats)
        )

    for scenario in problem3.boundary_cases(problem3.input_file):
        if scenario.name not in {
            "temperature_minus_1sigma",
            "temperature_plus_1sigma",
        }:
            continue
        print(f"复算长期边界情景：{scenario.label}")
        _, _, _, _, case_stats, _ = solve_question4(
            boundary=scenario.boundary,
            radius_history=radius_history,
        )
        report["boundary_sensitivity"].append(
            record_case(scenario.name, scenario.label, case_stats)
        )

    moving_h = report["baseline"]["continuous_threshold_time_h"]
    fixed_h = report["mechanism_comparison"][0][
        "continuous_threshold_time_h"
    ]
    appendix3_h = report["mechanism_comparison"][1][
        "continuous_threshold_time_h"
    ]
    appendix3_moving_h = report["mechanism_comparison"][2][
        "continuous_threshold_time_h"
    ]
    report["effects"] = {
        "appendix3_shrinkage_effect_hours": appendix3_moving_h - appendix3_h,
        "appendix4_shrinkage_effect_hours": moving_h - fixed_h,
        "fixed_radius_property_effect_hours": fixed_h - appendix3_h,
        "moving_radius_property_effect_hours": moving_h - appendix3_moving_h,
        "total_problem3_to_problem4_hours": moving_h - appendix3_h,
        "interpretation": (
            "两种物性下的收缩效应不同，说明物性与几何存在交互；"
            "不作唯一的加性贡献分解。"
        ),
    }
    report["verification_complete"] = True
    write_report(report)
    print(f"验证诊断已保存：{diagnostics_file}")


if __name__ == "__main__":
    main()
# 终于！噫嘘唏！