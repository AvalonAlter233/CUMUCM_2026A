import json
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

import 问题3_求解 as problem3
import 问题4_求解 as problem4


def test_radius_history_is_monotone_and_plateaus_after_measurements():
    history = problem4.read_radius_history(problem4.RADIUS_FILE)
    assert np.all(np.diff(history.times) > 0.0)
    assert np.all(np.diff(history.radii) <= 1.0e-12)
    assert history.radii[0] == pytest.approx(0.02000)
    assert history.radii[-1] == pytest.approx(0.01198)
    assert problem4.radius_value(history.times[-1] + 3600.0, history) == pytest.approx(
        history.radii[-1]
    )


def test_appendix_four_property_formulas_match_specification():
    moisture = np.array([0.15, 1.0, 2.55])
    temperature = np.array([28.0, 40.0, 50.0])
    assert np.allclose(problem4.density(moisture), 760.0 + 90.0 * moisture)
    assert np.allclose(
        problem4.heat_capacity(moisture),
        1850.0 + 2150.0 * moisture / (moisture + 1.0),
    )
    assert np.allclose(
        problem4.thermal_conductivity(moisture),
        0.12 + 0.20 * moisture / (moisture + 1.0),
    )
    expected_d = (
        4.2e-4
        * np.exp(-0.30 / moisture)
        * np.exp(-3850.0 / (temperature + 273.15))
    )
    np.testing.assert_allclose(
        problem4.moisture_diffusivity(moisture, temperature),
        expected_d,
        rtol=1.0e-12,
        atol=0.0,
    )


def test_appendix_three_property_model_reuses_problem_three_formulas():
    moisture = np.array([0.15, 1.0, 2.55])
    temperature = np.array([28.0, 40.0, 50.0])
    properties = problem4.material_properties("appendix3")
    np.testing.assert_allclose(
        properties.density(moisture), problem3.density(moisture)
    )
    np.testing.assert_allclose(
        properties.heat_capacity(moisture), problem3.heat_capacity(moisture)
    )
    np.testing.assert_allclose(
        properties.thermal_conductivity(moisture),
        problem3.thermal_conductivity(moisture),
    )
    np.testing.assert_allclose(
        properties.moisture_diffusivity(moisture, temperature),
        problem3.moisture_diffusivity(moisture, temperature),
        rtol=1.0e-12,
        atol=0.0,
    )


def test_reference_grid_has_correct_axisymmetric_volume():
    grid = problem4.build_reference_grid(320)
    assert len(grid.centers) == 320
    assert grid.spacing == pytest.approx(1.0 / 320.0)
    assert np.all(np.diff(grid.faces) > 0.0)
    assert np.sum(grid.volumes) == pytest.approx(0.5)


def test_reference_grid_rejects_too_few_cells():
    with pytest.raises(ValueError, match="至少为 4"):
        problem4.build_reference_grid(3)


def test_reference_system_has_positive_diagonal_and_expected_shape():
    grid = problem4.build_reference_grid(8)
    old = np.full(8, 1.0)
    storage = np.ones(8)
    transport = np.full(8, 2.0e-9)
    lower, diagonal, upper, rhs = problem4.assemble_reference_system(
        old, storage, transport, 8.0e-7, 0.05, grid, 0.02, 30.0
    )
    assert lower.shape == upper.shape == (7,)
    assert diagonal.shape == rhs.shape == (8,)
    assert np.all(diagonal > 0.0)
    assert np.all(lower < 0.0)
    assert np.all(upper < 0.0)


def test_fixed_radius_reference_system_matches_problem_three_system():
    intervals = 8
    reference_grid = problem4.build_reference_grid(intervals)
    physical_grid = problem3.build_grid(intervals)
    old = np.linspace(1.0, 0.5, intervals)
    storage = np.linspace(1.0, 2.0, intervals)
    transport = np.linspace(1.0e-9, 2.0e-9, intervals)
    reference_system = problem4.assemble_reference_system(
        old,
        storage,
        transport,
        problem4.CONVECTIVE_MASS_COEFF,
        0.05,
        reference_grid,
        problem4.INITIAL_RADIUS,
        30.0,
    )
    physical_system = problem3.assemble_system(
        old,
        storage,
        transport,
        problem3.CONVECTIVE_MASS_COEFF,
        0.05,
        physical_grid,
        30.0,
    )
    reference_solution = problem4.solve_tridiagonal(*reference_system)
    physical_solution = problem3.solve_tridiagonal(*physical_system)
    np.testing.assert_allclose(reference_solution, physical_solution, rtol=1.0e-12)


def test_tridiagonal_solver_matches_dense_solution():
    lower = np.array([-1.0, -1.0])
    diagonal = np.array([4.0, 4.0, 4.0])
    upper = np.array([-1.0, -1.0])
    rhs = np.array([2.0, 4.0, 10.0])
    dense = np.diag(diagonal) + np.diag(lower, -1) + np.diag(upper, 1)
    assert np.allclose(
        problem4.solve_tridiagonal(lower, diagonal, upper, rhs),
        np.linalg.solve(dense, rhs),
    )


def test_reconstruction_returns_fixed_nodes_and_moving_surface():
    grid = problem4.build_reference_grid(8)
    values = 1.0 - 0.5 * grid.centers
    transport = np.full(8, 2.0e-9)
    nodes = np.array([0.0, 0.005, 0.010])
    result = problem4.reconstruct_output_field(
        values, transport, 8.0e-7, 0.05, nodes, grid, 0.02
    )
    assert result.shape == (4,)
    assert np.all(np.diff(result) <= 0.0)
    assert 0.05 < result[-1] < values[-1]


def test_reconstruction_rejects_nodes_outside_current_radius():
    grid = problem4.build_reference_grid(8)
    with pytest.raises(ValueError, match="药材内部"):
        problem4.reconstruct_output_field(
            np.ones(8),
            np.ones(8),
            8.0e-7,
            0.05,
            np.array([0.0, 0.021]),
            grid,
            0.02,
        )


def test_template_nodes_and_headers_follow_result4_layout():
    nodes = problem4.template_output_nodes()
    assert np.allclose(nodes, np.arange(0.0, 0.012, 0.001))
    assert problem4.output_headers(nodes) == [
        0.0,
        0.1,
        0.2,
        0.3,
        0.4,
        0.5,
        0.6,
        0.7,
        0.8,
        0.9,
        1.0,
        1.1,
        "药材表面",
    ]


def test_solver_can_capture_full_internal_reference_field(tmp_path):
    result = problem4.solve_problem_four(
        internal_intervals=8,
        time_step=30.0,
        report_interval=3600.0,
        max_time=3600.0,
        critical_moisture=2.55,
        capture_internal=True,
    )
    times, _, _, surface_radii, _, internal = result
    assert internal is not None
    assert internal.moisture.shape == (1, 8)
    np.testing.assert_allclose(internal.times, times)
    np.testing.assert_allclose(internal.surface_radii, surface_radii)
    assert np.all(np.diff(internal.xi_centers) > 0.0)

    output = tmp_path / "internal_field.npz"
    problem4.write_internal_field_data(internal, output)
    with np.load(output) as saved:
        np.testing.assert_allclose(saved["times_s"], times)
        np.testing.assert_allclose(saved["xi_centers"], internal.xi_centers)
        np.testing.assert_allclose(saved["moisture"], internal.moisture)
        np.testing.assert_allclose(saved["surface_radii_m"], surface_radii)


def test_baseline_report_cannot_reuse_stale_verification_cases():
    baseline = {
        "continuous_threshold_time_s": 10.0,
        "discrete_threshold_time_s": 60.0,
    }
    report = problem4.new_diagnostics_report(baseline)
    assert not report["verification_complete"]
    assert report["mechanism_comparison"] == []
    assert report["numerical_refinement"] == []
    assert report["boundary_sensitivity"] == []


def test_result_workbook_and_diagnostics_confirm_official_end_time():
    workbook = load_workbook(problem4.OUTPUT_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    assert sheet.max_column == 14
    assert [sheet.cell(1, c).value for c in range(1, 15)] == [
        "时间\\到药材中心的距离",
        *problem4.output_headers(problem4.template_output_nodes()),
    ]
    previous = np.array(
        [sheet.cell(sheet.max_row - 1, c).value for c in range(2, sheet.max_column + 1)]
    )
    final = np.array(
        [sheet.cell(sheet.max_row, c).value for c in range(2, sheet.max_column + 1)]
    )
    assert sheet.cell(sheet.max_row, 1).value == 184020
    assert np.max(previous) >= problem4.CRITICAL_MOISTURE
    assert np.max(final) < problem4.CRITICAL_MOISTURE

    diagnostics_path = Path(problem4.DIAGNOSTICS_FILE)
    report = json.loads(diagnostics_path.read_text(encoding="utf-8"))
    baseline = report["baseline"]
    assert baseline["discrete_threshold_time_s"] == pytest.approx(184020.0)
    assert baseline["continuous_threshold_time_h"] == pytest.approx(51.1151668549404)
    assert baseline["center_controls_threshold"]
    assert baseline["all_domain_checked_every_step"]
    assert baseline["radially_nonincreasing_every_step"]
    assert baseline["moisture_balance_relative_imbalance"] < 1.0e-5
