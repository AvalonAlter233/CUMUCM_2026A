from __future__ import annotations
import argparse
from pathlib import Path
from typing import NamedTuple
import numpy as np
from openpyxl import load_workbook

# 依旧很多参数
project_root = Path(__file__).resolve().parent
input_file = project_root / "附件" / "附件1.xlsx"
output_file = project_root / "附件" / "附件3" / "result2.xlsx"

pellet_radius = 0.02
convective_heat_coeff = 25.0
convective_mass_coeff = 8.0e-7
initial_temperature = 28.0
initial_moisture = 2.55

default_end_time = 10800.0
default_time_step = 1.0
internal_cell_count = 160
output_time_interval = 1.0
report_radii = np.linspace(0.0, pellet_radius, 21)

convergence_tol = 1.0e-9
max_iterations = 60
relaxation_factor = 0.8


class DryingRoomBoundary(NamedTuple):

    times: np.ndarray
    temperatures: np.ndarray
    moisture: np.ndarray


class RadialGrid(NamedTuple):

    centers: np.ndarray
    faces: np.ndarray
    volume_factors: np.ndarray
    spacing: float


class SimulationDetails(NamedTuple):

    times: np.ndarray
    output_radii: np.ndarray
    temperature_field: np.ndarray
    moisture_field: np.ndarray
    average_temperature: np.ndarray
    average_moisture: np.ndarray
    diagnostics: dict[str, float]


def read_room_data(path: Path) -> DryingRoomBoundary:
    # 这里还是从那张环境数据表开始，先把要用的几列拿出来。
    # 温度和水分一起读好，后面哪个计算步骤要用，就把对应的一列传过去。
    # 附件先读顺了再去算别的，主循环后面还要一次次来取这里的数据。
    book = load_workbook(path, data_only=True, read_only=True)
    sheet = book.active
    rows = [
        row for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)

    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一必须包含时间、温度、水分浓度三列有效数据。")
    if not np.all(np.isfinite(data[:, :3])):
        raise ValueError("附件一的前三列存在空值或非有限数值。")
    if not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError("附件一中的时间必须严格递增。")

    return DryingRoomBoundary(
        times=data[:, 0],
        temperatures=data[:, 1],
        moisture=data[:, 2],
    )


def linear_interp(t: float, times: np.ndarray, values: np.ndarray) -> float:
    if t < times[0] or t > times[-1]:
        raise ValueError(
            f"计算时刻 {t:g} s 超出附件一范围 "
            f"[{times[0]:g}, {times[-1]:g}] s。"
        )
    return float(np.interp(t, times, values))


def density(moisture: np.ndarray) -> np.ndarray:
    # 这几个物性式子先分别放好，后面会反复叫到它们。
    # 用到哪一个就让对应的函数算一下，主循环里也就不用再把整条式子写一遍。
    # 一眼看过去都是几行小式子，先放在一起，后面用起来反倒比较好找。
    return 650.0 + 128.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    # 当前的含水率传过来以后，比热就照着这一条式子算。
    # 这里算完直接交回去，外面需要下一轮的值时还会再来。
    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    # 导热系数也从含水率这边接着算，放在这里用起来比较顺手。
    # 先返回当前这一份，后面的温度计算会把它带进去。
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    c_safe = np.maximum(moisture, 1.0e-12)
    # 扩散系数的式子里用的是开尔文温度，所以在这里加上换算值。
    # 换算完的这个数接着给下面的式子用，外面传进来的温度还是原来的摄氏度。
    # 这里记一下温度换算发生在哪就够了，后面接着用换算好的数往下算。
    temp_k = temperature_celsius + 273.15
    return (
        2.4e-3
        * np.exp(-0.45 / c_safe)
        * np.exp(-3850.0 / temp_k)
    )


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    denom = np.maximum(left + right, 1.0e-30)
    return 2.0 * left * right / denom


def build_radial_grid(cell_count: int) -> RadialGrid:
    # 位置先分好，中心和分界都存在这份网格里。
    # 后面每到一个新时刻，还用这一套位置，只更新这些位置上对应的数值。
    # 位置这一步准备好以后就先搁着，后面一轮轮变化的是上面的数。
    # 网格里的这几组数组先认一下，下面会经常看到它们。
    if cell_count < 2:
        raise ValueError("径向有限体积单元数至少为 2。")
    faces = np.linspace(0.0, pellet_radius, cell_count + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    vol = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return RadialGrid(
        centers=centers,
        faces=faces,
        volume_factors=vol,
        spacing=float(faces[1] - faces[0]),
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    # 这一段先拿副本往下做消元，随后再倒回来把解求出来。
    # 先看前面的循环怎么推过去，再接着看后面怎么推回来，就能把顺序连上。
    # 外面把系数传到这里，等返回时拿走的就是这一轮的解。
    diagonal = diagonal.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    upper = upper.astype(float, copy=True)

    for i in range(1, len(diagonal)):
        if abs(diagonal[i - 1]) < 1.0e-30:
            raise FloatingPointError("三对角方程组出现近零主对角元。")
        factor = lower[i - 1] / diagonal[i - 1]
        diagonal[i] -= factor * upper[i - 1]
        right_hand_side[i] -= factor * right_hand_side[i - 1]

    sol = np.empty_like(right_hand_side)
    sol[-1] = right_hand_side[-1] / diagonal[-1]
    for i in range(len(diagonal) - 2, -1, -1):
        sol[i] = (
            right_hand_side[i] - upper[i] * sol[i + 1]
        ) / diagonal[i]
    return sol


def surface_transfer(
    internal_coefficient: float,
    boundary_coefficient: float,
    half_cell_width: float,
) -> float:
    # 内部走到表面还有一小段，表面这边又有自己的传递系数。
    # 这两部分在这里接起来，后面组方程时就直接用合好的这一项。
    if internal_coefficient <= 0.0 or boundary_coefficient <= 0.0:
        raise ValueError("内部传输系数和表面传递系数必须为正。")
    return 1.0 / (
        half_cell_width / internal_coefficient + 1.0 / boundary_coefficient
    )


def assemble_system(
    storage_coefficients: np.ndarray,
    old_values: np.ndarray,
    grid: RadialGrid,
    interface_coefficients: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 本轮要解的那些系数先在这一段里凑起来。
    # 里面各个小单元先接好，最外面和环境接触的那一项再添进去，最后一起交给求解函数。
    # 先把方程这一边收齐，求解的事交给后面那一步，顺着传过去就行。
    cell_count = len(grid.centers)
    lower = np.zeros(cell_count - 1)
    upper = np.zeros(cell_count - 1)
    diag = storage_coefficients.copy()
    rhs = storage_coefficients * old_values

    g_face = grid.faces[1:-1] * interface_coefficients / grid.spacing
    diag[:-1] += g_face
    diag[1:] += g_face
    upper[:] = -g_face
    lower[:] = -g_face

    h_eff = surface_transfer(
        internal_coefficient=surface_internal_coefficient,
        boundary_coefficient=surface_transfer_coefficient,
        half_cell_width=0.5 * grid.spacing,
    )
    g_surf = pellet_radius * h_eff
    diag[-1] += g_surf
    rhs[-1] += g_surf * boundary_value
    return lower, diag, upper, rhs


def surface_value(
    last_cell_value: float,
    internal_coefficient: float,
    boundary_coefficient: float,
    boundary_value: float,
    grid_spacing: float,
) -> float:
    # 表面这一点先另外处理，最外面那个单元的值也一起带进来。
    # 算好后先交出去，后面再接着整理整条径向分布。
    g_inner = 2.0 * internal_coefficient / grid_spacing
    return float(
        (g_inner * last_cell_value + boundary_coefficient * boundary_value)
        / (g_inner + boundary_coefficient)
    )


def sample_profile(
    cell_values: np.ndarray,
    grid: RadialGrid,
    output_radii: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> np.ndarray:
    # 要输出的位置和网格中心稍微有些区别，这里先把中心和表面补齐。
    # 两头都接上了，再在中间找需要的输出位置，最后返回表格要的那一排。
    r1_sq = grid.centers[0] ** 2
    r2_sq = grid.centers[1] ** 2
    axis_val = (
        r2_sq * cell_values[0]
        - r1_sq * cell_values[1]
    ) / (r2_sq - r1_sq)
    val_surf = surface_value(
        float(cell_values[-1]),
        surface_internal_coefficient,
        surface_transfer_coefficient,
        boundary_value,
        grid.spacing,
    )
    r_fit = np.concatenate(([0.0], grid.centers, [pellet_radius]))
    val_fit = np.concatenate(([axis_val], cell_values, [val_surf]))
    return np.interp(output_radii, r_fit, val_fit)


def radial_average(cell_values: np.ndarray, grid: RadialGrid) -> float:
    return float(2.0 * np.sum(grid.volume_factors * cell_values) / pellet_radius**2)


def balance_residual(storage_rate: float, outward_flux: float) -> float:
    scale = abs(storage_rate) + abs(outward_flux) + 1.0e-30
    return abs(storage_rate + outward_flux) / scale


def temperature_step(
    old_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    heat_transfer_coefficient: float,
) -> np.ndarray:
    # 传进来的含水率先用来算这一轮的物性，再接着求温度。
    # 这里拿到的是一个候选值，返回外面以后还得继续迭代和检查，现在先把它算出来。
    # 这一轮先拿到一份温度结果，外面的循环还会接着安排后面的事。
    # 先算好这一份，再回到主循环里继续看含水率。
    rho = density(reference_moisture)
    cp = heat_capacity(reference_moisture)
    k_cell = thermal_conductivity(reference_moisture)
    storage = rho * cp * grid.volume_factors / time_step
    k_face = harmonic_mean(k_cell[:-1], k_cell[1:])
    temp_air = linear_interp(current_time, boundary.times, boundary.temperatures)

    lin_sys = assemble_system(
        storage,
        old_temperature,
        grid,
        k_face,
        float(k_cell[-1]),
        heat_transfer_coefficient,
        temp_air,
    )
    return solve_tridiagonal(*lin_sys)


def moisture_step(
    old_moisture: np.ndarray,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    mass_transfer_coefficient: float,
) -> np.ndarray:
    # 这一轮水分先用传进来的温度和含水率准备系数，准备好再解方程。
    # 这里先走完水分这一步，外面的迭代循环还会继续接着安排。
    d_cell = moisture_diffusivity(reference_moisture, reference_temperature)
    d_face = harmonic_mean(d_cell[:-1], d_cell[1:])
    storage = grid.volume_factors / time_step
    water_air = linear_interp(current_time, boundary.times, boundary.moisture)

    lin_sys = assemble_system(
        storage,
        old_moisture,
        grid,
        d_face,
        float(d_cell[-1]),
        mass_transfer_coefficient,
        water_air,
    )
    return solve_tridiagonal(*lin_sys)


def simulate_question2(
    end_time: float = default_end_time,
    time_step: float = default_time_step,
    cell_count: int = internal_cell_count,
    output_interval: float = output_time_interval,
    heat_transfer_coefficient: float = convective_heat_coeff,
    mass_transfer_coefficient: float = convective_mass_coeff,
) -> SimulationDetails:
    # 这一问的主计算从这里往下走，设置的数也都从这里接进来。
    # 先把输入和网格准备好，再把初始值放进去，准备完了以后才进入下面的循环。
    # 入口这段东西稍多一些，先顺着把准备工作看完，再去看真正推进时间的部分。
    if min(end_time, time_step, output_interval) <= 0.0:
        raise ValueError("终止时间、时间步长和输出间隔必须为正。")
    n_steps = int(round(end_time / time_step))
    save_every = int(round(output_interval / time_step))
    if not np.isclose(n_steps * time_step, end_time):
        raise ValueError("终止时间必须是时间步长的整数倍。")
    if save_every < 1 or not np.isclose(save_every * time_step, output_interval):
        raise ValueError("输出间隔必须是时间步长的整数倍。")

    boundary = read_room_data(input_file)
    if end_time > boundary.times[-1]:
        raise ValueError("附件一边界数据没有覆盖要求的计算时段。")
    mesh = build_radial_grid(cell_count)

    temp = np.full(cell_count, initial_temperature, dtype=float)
    water = np.full(cell_count, initial_moisture, dtype=float)

    time_log: list[float] = []
    temp_log: list[np.ndarray] = []
    water_log: list[np.ndarray] = []
    temp_mean_log: list[float] = []
    water_mean_log: list[float] = []
    iter_log: list[int] = []
    max_temp_err = 0.0
    max_moisture_err = 0.0
    max_heat_err = 0.0
    max_water_err = 0.0
    rho_min = float("inf")
    rho_max = -float("inf")
    cp_min = float("inf")
    cp_max = -float("inf")
    k_min = float("inf")
    k_max = -float("inf")
    d_min = float("inf")
    d_max = -float("inf")

    # 新的一步仍然接在旧的一步后面，前面的结果要继续用。
    # 先把旧值留出来，再单独放本轮的迭代值，后面更新的时候就各用各的。
    # 上一时刻和本轮迭代是两套记录，读到这里时稍微分清一下就好。
    # 前一套先留着，本轮有新的变化再记到这一轮对应的变量里。
    for step in range(1, n_steps + 1):
        t_now = step * time_step
        temp_old = temp.copy()
        water_old = water.copy()
        temp_iter = temp.copy()
        water_iter = water.copy()

        # 温度更新一遍，含水率也跟着更新一遍。
        # 算完看一下变化量，还没满足要求的话，就拿这一轮的结果再接着算下一轮。
        # 也就是在当前这一步里来回更新几次，先把这一步算妥当，再接下一步。
        for iteration in range(1, max_iterations + 1):
            temp_trial = temperature_step(
                temp_old,
                water_iter,
                t_now,
                time_step,
                mesh,
                boundary,
                heat_transfer_coefficient,
            )
            water_trial = moisture_step(
                water_old,
                temp_trial,
                water_iter,
                t_now,
                time_step,
                mesh,
                boundary,
                mass_transfer_coefficient,
            )
            # 候选值出来后，先和上一轮的值按比例放到一起。
            # 温度这样做，含水率也这样做，得到新的值以后再去检查前后变化了多少。
            # 这一步先把候选结果接过来，接好以后再看变化量，顺序就照下面这样走。
            temp_new = (
                relaxation_factor * temp_trial
                + (1.0 - relaxation_factor) * temp_iter
            )
            water_new = (
                relaxation_factor * water_trial
                + (1.0 - relaxation_factor) * water_iter
            )
            err_t = float(
                np.max(
                    np.abs(temp_new - temp_iter)
                    / (1.0 + np.abs(temp_new))
                )
            )
            err_c = float(
                np.max(
                    np.abs(water_new - water_iter)
                    / (1.0 + np.abs(water_new))
                )
            )
            temp_iter = temp_new
            water_iter = water_new
            if max(err_t, err_c) < convergence_tol:
                break
        else:
            raise RuntimeError(
                f"{t_now:g} s 的热湿 Picard 迭代未收敛，"
                f"温度变化量={err_t:.3e}，"
                f"含水率变化量={err_c:.3e}。"
            )

        temp = temp_iter
        water = water_iter
        if not np.all(np.isfinite(temp)) or not np.all(np.isfinite(water)):
            raise FloatingPointError(f"{t_now:g} s 出现非有限场变量。")
        if np.min(water) <= 0.0:
            raise FloatingPointError(f"{t_now:g} s 出现非正含水率。")

        rho = density(water)
        cp = heat_capacity(water)
        k_cell = thermal_conductivity(water)
        d_cell = moisture_diffusivity(water, temp)
        temp_air = linear_interp(t_now, boundary.times, boundary.temperatures)
        water_air = linear_interp(t_now, boundary.times, boundary.moisture)
        temp_surf = surface_value(
            float(temp[-1]),
            float(k_cell[-1]),
            heat_transfer_coefficient,
            temp_air,
            mesh.spacing,
        )
        water_surf = surface_value(
            float(water[-1]),
            float(d_cell[-1]),
            mass_transfer_coefficient,
            water_air,
            mesh.spacing,
        )

        heat_rate = float(
            np.sum(
                rho * cp * mesh.volume_factors
                * (temp - temp_old) / time_step
            )
        )
        water_rate = float(
            np.sum(mesh.volume_factors * (water - water_old) / time_step)
        )
        heat_out = (
            pellet_radius * heat_transfer_coefficient
            * (temp_surf - temp_air)
        )
        water_out = (
            pellet_radius * mass_transfer_coefficient
            * (water_surf - water_air)
        )
        max_heat_err = max(
            max_heat_err,
            balance_residual(heat_rate, heat_out),
        )
        max_water_err = max(
            max_water_err,
            balance_residual(water_rate, water_out),
        )
        iter_log.append(iteration)
        max_temp_err = max(max_temp_err, err_t)
        max_moisture_err = max(max_moisture_err, err_c)
        # 这一轮的物性也顺手看一下，碰到更大或者更小的数就记下来。
        # 循环往后走的时候这些记录也跟着更新，最后汇总时就直接用留下来的数。
        # 这些值不急着打印，先一路记着，等主计算结束以后再一起拿出来看。
        rho_min = min(rho_min, float(np.min(rho)))
        rho_max = max(rho_max, float(np.max(rho)))
        cp_min = min(cp_min, float(np.min(cp)))
        cp_max = max(cp_max, float(np.max(cp)))
        k_min = min(k_min, float(np.min(k_cell)))
        k_max = max(k_max, float(np.max(k_cell)))
        d_min = min(d_min, float(np.min(d_cell)))
        d_max = max(d_max, float(np.max(d_cell)))

        # 需要写进表格的那份结果，在这里往列表里添。
        # 前面迭代了几轮是一回事，到输出时刻保存下来又是一回事，这里负责把要输出的收好。
        # 要存的先放好，其余时刻照样往后推进，到下一次保存时再添新的。
        # 等所有时间走完，这些列表里的顺序就是最后输出的顺序。
        if step % save_every == 0:
            time_log.append(t_now)
            temp_log.append(
                sample_profile(
                    temp,
                    mesh,
                    report_radii,
                    float(k_cell[-1]),
                    heat_transfer_coefficient,
                    temp_air,
                )
            )
            water_log.append(
                sample_profile(
                    water,
                    mesh,
                    report_radii,
                    float(d_cell[-1]),
                    mass_transfer_coefficient,
                    water_air,
                )
            )
            temp_mean_log.append(radial_average(temp, mesh))
            water_mean_log.append(radial_average(water, mesh))

    diagnostics = {
        "internal_cell_count": float(cell_count),
        "time_step": float(time_step),
        "max_picard_iterations": float(max(iter_log)),
        "mean_picard_iterations": float(np.mean(iter_log)),
        "max_final_temperature_error": float(max_temp_err),
        "max_final_moisture_error": float(max_moisture_err),
        "max_heat_balance_residual": float(max_heat_err),
        "max_moisture_balance_residual": float(max_water_err),
        "minimum_density": rho_min,
        "maximum_density": rho_max,
        "minimum_heat_capacity": cp_min,
        "maximum_heat_capacity": cp_max,
        "minimum_conductivity": k_min,
        "maximum_conductivity": k_max,
        "minimum_diffusivity": d_min,
        "maximum_diffusivity": d_max,
    }
    return SimulationDetails(
        times=np.asarray(time_log),
        output_radii=report_radii.copy(),
        temperature_field=np.asarray(temp_log),
        moisture_field=np.asarray(water_log),
        average_temperature=np.asarray(temp_mean_log),
        average_moisture=np.asarray(water_mean_log),
        diagnostics=diagnostics,
    )


def solve_question2(
    end_time: float = default_end_time,
    time_step: float = default_time_step,
    radial_intervals: int = internal_cell_count,
    heat_transfer_coefficient: float = convective_heat_coeff,
    mass_transfer_coefficient: float = convective_mass_coeff,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, dict[str, float]]:
    result = simulate_question2(
        end_time=end_time,
        time_step=time_step,
        cell_count=radial_intervals,
        output_interval=output_time_interval,
        heat_transfer_coefficient=heat_transfer_coefficient,
        mass_transfer_coefficient=mass_transfer_coefficient,
    )
    return (
        result.times,
        result.output_radii,
        result.temperature_field,
        result.moisture_field,
        result.diagnostics,
    )


def write_results(
    times: np.ndarray,
    output_radii: np.ndarray,
    temperature_field: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    # 温度先放到它的工作表里，水分再放到另一个工作表里。
    # 两个表都照着时间往下写，表头和格式也一并处理好，全部填完再保存文件。
    # 把一行里的东西排齐了，再去填下一行，后面翻表时看起来就顺一些。
    book = load_workbook(output_file)
    if len(book.worksheets) < 2:
        raise ValueError("result2.xlsx 必须包含温度和水分浓度两个工作表。")

    for sheet, field in zip(
        book.worksheets[:2], (temperature_field, moisture_field)
    ):
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)

        sheet.cell(1, 1).value = "时间\\到药材中心的距离"
        sheet.cell(1, 1).number_format = "@"
        for column, radius in enumerate(output_radii, start=2):
            sheet.cell(1, column).value = round(float(radius * 100.0), 1)
            sheet.cell(1, column).number_format = "0.0"

        for row, t_now in enumerate(times, start=2):
            sheet.cell(row, 1).value = int(round(float(t_now)))
            sheet.cell(row, 1).number_format = "0"
            for column, value in enumerate(field[row - 2], start=2):
                sheet.cell(row, column).value = round(float(value), 4)
                sheet.cell(row, column).number_format = "0.0000"

    book.save(output_file)


def report_difference(
    label: str,
    times: np.ndarray,
    radii: np.ndarray,
    first_field: np.ndarray,
    second_field: np.ndarray,
) -> None:
    diff = np.abs(first_field - second_field)
    it, ir = np.unravel_index(
        int(np.argmax(diff)), diff.shape
    )
    print(
        f"{label}：最大绝对差={diff[it, ir]:.6e}，"
        f"位置为 t={times[it]:g} s、"
        f"r={radii[ir] * 100:g} cm；"
        f"终点最大差={np.max(diff[-1]):.6e}。"
    )


def check_convergence(reference: SimulationDetails) -> None:
    # 这里换一下网格，再换一下时间步长，各自重新算一组。
    # 结果出来以后和主计算放到一起比，终端里会把对应的差值列出来，顺着看就可以。
    # 这里会比只跑主结果多花一点时间，先让几组计算各自跑完。
    # 最后看终端给出的那几行，把每一组的名字和差值对着看就行。
    space_coarse = simulate_question2(cell_count=80)
    time_fine = simulate_question2(cell_count=160, time_step=0.5)

    print("问题二数值检验：")
    report_difference(
        "N=80 与 N=160 的温度场",
        reference.times,
        reference.output_radii,
        space_coarse.temperature_field,
        reference.temperature_field,
    )
    report_difference(
        "N=80 与 N=160 的含水率场",
        reference.times,
        reference.output_radii,
        space_coarse.moisture_field,
        reference.moisture_field,
    )
    report_difference(
        "时间步 1 s 与 0.5 s 的温度场",
        reference.times,
        reference.output_radii,
        reference.temperature_field,
        time_fine.temperature_field,
    )
    report_difference(
        "时间步 1 s 与 0.5 s 的含水率场",
        reference.times,
        reference.output_radii,
        reference.moisture_field,
        time_fine.moisture_field,
    )
    print(f"主计算诊断：{reference.diagnostics}")


def main() -> None:
    # 先跑这一问的主计算，等拿到结果后就按原来的表格安排写出去。
    # 如果还开着验证选项，下面再接着多跑几组，把整段流程做完。
    parser = argparse.ArgumentParser(description="求解问题二的变物性径向传热传质模型")
    parser.add_argument(
        "--validate",
        action="store_true",
        help="额外运行全 3 h 空间和时间步长检验，只在终端输出诊断",
    )
    args = parser.parse_args()

    result = simulate_question2()
    write_results(
        result.times,
        result.output_radii,
        result.temperature_field,
        result.moisture_field,
    )
    print(f"问题二计算完成，结果已写入：{output_file}")
    if args.validate:
        check_convergence(result)


if __name__ == "__main__":
    main()
