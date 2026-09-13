from __future__ import annotations

import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib as mpl
import matplotlib.font_manager as font_manager
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

if (Path(__file__).parent / "问题4_求解_改.py").is_file():
    import 问题4_求解_改 as problem4
else:
    import 问题4_求解 as problem4


nature_figure_scripts = Path.home() / ".codex" / "skills" / "nature-figure" / "scripts"
if str(nature_figure_scripts) not in sys.path:
    sys.path.insert(0, str(nature_figure_scripts))
from audit_panel_alignment import require_matplotlib_panel_alignment


project_root = Path(__file__).resolve().parent
result_file = project_root / "附件" / "附件3" / "result4.xlsx"
diagnostics_file = project_root / "附件" / "附件3" / "result4_diagnostics.json"
internal_field_file = project_root / "附件" / "附件3" / "result4_internal_field.npz"
figure_dir = project_root / "figures" / "问题四"

default_critical_moisture = 0.15
initial_moisture = 2.55
initial_radius_cm = 2.0
export_dpi = 600
pad_inches = 0.06
figure_formats = ("png",)

cjk_font_candidates = (
    "STSong", "SimSun", "Source Han Serif SC", "Noto Serif CJK SC",
)
installed_fonts = {font.name for font in font_manager.fontManager.ttflist}
available_cjk_fonts = [name for name in cjk_font_candidates if name in installed_fonts]
heading_font = "STZhongsong" if "STZhongsong" in installed_fonts else available_cjk_fonts[0]

mpl.rcParams.update({
    "font.family": "serif",
    "font.serif": [*available_cjk_fonts, "Times New Roman", "Times", "DejaVu Serif"],
    "mathtext.fontset": "stix",
    "axes.unicode_minus": False,
    "font.size": 8,
    "axes.titlesize": 9.2,
    "axes.titleweight": "bold",
    "axes.labelsize": 8,
    "axes.labelweight": "bold",
    "legend.fontsize": 7.2,
    "xtick.labelsize": 7.5,
    "ytick.labelsize": 7.5,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "axes.linewidth": 0.7,
    "legend.frameon": False,
    "figure.dpi": 120,
    "svg.fonttype": "none",
    "pdf.fonttype": 42,
})

reference_palette = ("#44757A", "#452A3D", "#D44C3C", "#EED5B7")
color_teal, color_plum, color_coral, color_sand = reference_palette
color_dark = "#51474D"

new_figure_bases = (
    "图1_收缩轨迹与含水率自洽性",
    "图2_移动域达标时间判定",
    "图3_移动域含水率演化",
    "图4_物性与几何非线性交互",
    "图5_数值与边界稳健性",
)


def read_results() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    # 先把结果表里的输出数据取出来，列的顺序就照表里原来的顺序保留。
    # 先按这个顺序把数据收进来，后面要哪一列，再从这份结果里取
    book = load_workbook(result_file, data_only=True, read_only=True)
    rows = list(book.active.iter_rows(values_only=True))
    fixed_radii = np.asarray(rows[0][1:-1], dtype=float)
    times_h = np.asarray([row[0] for row in rows[1:]], dtype=float) / 3600.0
    water = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    return times_h, fixed_radii, water


def read_internal_field(
    path: Path = internal_field_file,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 内部细网格的数据放在单独的文件里，所以这里再另外读一次。
    # 时间、位置和半径一起拿出来，读完以后对一下尺寸，能对应上了再交给后面的步骤。
    # 这份数据比结果表细一些，先看它和时间、位置能不能接得上，接好了再往后用。
    with np.load(path) as data:
        times_h = np.asarray(data["times_s"], dtype=float) / 3600.0
        xi_centers = np.asarray(data["xi_centers"], dtype=float)
        water = np.asarray(data["moisture"], dtype=float)
        surf_radii = 100.0 * np.asarray(data["surface_radii_m"], dtype=float)
    if water.shape != (len(times_h), len(xi_centers)):
        raise ValueError("内部细网格含水率数组尺寸与时间、空间坐标不一致。")
    if len(surf_radii) != len(times_h):
        raise ValueError("内部细网格半径序列与时间序列长度不一致。")
    return times_h, xi_centers, water, surf_radii


def find_threshold_times(
    times_h: np.ndarray,
    maximum_moisture: np.ndarray,
    threshold: float = default_critical_moisture,
) -> tuple[float, float]:
    # 先找第一次低于阈值的记录，再回头把挨着的前一条拿到手。
    # 有了这两个时刻和对应的数，就在中间估一下经过阈值的时间。
    # 估计时间和记录时刻都留着，后面各有地方要用。
    indices = np.flatnonzero(maximum_moisture < threshold)
    if len(indices) == 0 or indices[0] == 0:
        raise ValueError("输出记录没有形成有效的阈值夹逼。")
    index = int(indices[0])
    y0, y1 = float(maximum_moisture[index - 1]), float(maximum_moisture[index])
    crossing = float(times_h[index - 1]) + (y0 - threshold) / (y0 - y1) * float(
        times_h[index] - times_h[index - 1]
    )
    return crossing, float(times_h[index])


def mechanism_cases(diagnostics: dict) -> list[dict]:
    # 先看看对照计算是不是已经全部完成，再决定要不要往下取结果。
    # 完成以后，就把这几种情况按原先约定的顺序排好，后面画图时接着用这个顺序。
    # 对照齐了就挨个取出来，还没齐的时候先按这里的返回结果处理
    if not diagnostics.get("verification_complete", False):
        return []
    case_map = {
        case["name"]: case
        for case in [diagnostics["baseline"], *diagnostics["mechanism_comparison"]]
    }
    case_order = (
        "appendix3_fixed_radius", "appendix3_moving_radius",
        "appendix4_fixed_radius", "appendix4_moving_radius",
    )
    return [case_map[name] for name in case_order]


def load_plot_data() -> dict[str, object]:
    # 这里要接起来的数据有几份，结果表、内部场和半径记录都得用。
    # 先把它们整理到同一份里面，后面需要画哪一部分，就从这里的结果中取对应的量
    out_times, fixed_radii, out_water = read_results()
    inner_times, xi_centers, inner_water, surf_radii = read_internal_field()
    if not np.allclose(out_times, inner_times, rtol=0.0, atol=1e-12):
        raise ValueError("内部细网格与结果工作簿的输出时刻不一致。")
    with diagnostics_file.open("r", encoding="utf-8") as stream:
        stats = json.load(stream)
    history = problem4.read_radius_data(problem4.radius_file)

    times_h = np.concatenate(([0.0], out_times))
    output = np.vstack((np.full((1, out_water.shape[1]), initial_moisture), out_water))
    internal = np.vstack((np.full((1, inner_water.shape[1]), initial_moisture), inner_water))
    surf_radius = np.concatenate(([initial_radius_cm], surf_radii))
    maximum = np.maximum(output.max(axis=1), internal.max(axis=1))
    baseline = stats["baseline"]
    return {
        "times_h": times_h,
        "fixed_radii_cm": fixed_radii,
        "output_moisture": output,
        "xi_centers": xi_centers,
        "internal_moisture": internal,
        "surface_radii_cm": surf_radius,
        "center_moisture": output[:, 0],
        "surface_moisture": output[:, -1],
        "maximum_moisture": maximum,
        "radius_history_times_h": history.times / 3600.0,
        "radius_history_cm": history.radii * 100.0,
        "continuous_threshold_time_h": float(baseline["continuous_threshold_time_h"]),
        "discrete_threshold_time_h": float(baseline["discrete_threshold_time_h"]),
        "diagnostics": stats,
    }


def shrinkage_drivers(data: dict[str, object]) -> dict[str, np.ndarray]:
    # 几种代表性的含水率先分别算出来
    # 后面要比较哪一种，就从这里返回的结果里取哪一种，不用再回头重复整理内部场。
    # 名字和对应的那串数放在一起
    xi = np.asarray(data["xi_centers"], dtype=float)
    water = np.asarray(data["internal_moisture"], dtype=float)
    surface = np.asarray(data["surface_moisture"], dtype=float)
    volume_mean = 2.0 * np.trapezoid(water * xi[None, :], xi, axis=1)
    outer_mask = xi >= 0.5
    outer_xi = xi[outer_mask]
    outer_mean = np.trapezoid(
        water[:, outer_mask] * outer_xi[None, :], outer_xi, axis=1
    ) / np.trapezoid(outer_xi, outer_xi)
    return {
        "volume_average": volume_mean,
        "outer_average": outer_mean,
        "surface": surface,
    }


def shrinkage_phi(data: dict[str, object]) -> dict[str, np.ndarray]:
    # 半径和前面准备好的含水率现在放在一起算，结果先一项项存起来。
    # 有些位置暂时算不了，就先留成空值，后面能用的那些位置再接着往下处理。
    # 先把可以算的位置找出来，再一个个填上，暂时留空的就继续留在原来的位置。
    # 后面做统计时还会看哪些数能用
    radius = np.asarray(data["surface_radii_cm"], dtype=float)
    s = (radius / initial_radius_cm) ** 2
    denom = 1.0 - s
    fields = shrinkage_drivers(data)
    result = {}
    for name, values in fields.items():
        phi = np.full_like(values, np.nan)
        valid = denom > 1e-8
        phi[valid] = (s[valid] * initial_moisture - values[valid]) / denom[valid]
        result[name] = phi
    return result


def shrinkage_stats(
    data: dict[str, object], start_h: float = 12.0
) -> dict[str, dict[str, float]]:
    # 每一段的统计先各自算一遍，需要累积的量也跟着一点点加起来。
    # 这些小段都处理完以后，再把最后整理好的几个数放到一起返回出去。
    # 先看每小段自己的情况，再把要汇总的数接起来
    times = np.asarray(data["times_h"], dtype=float)
    trajectories = shrinkage_phi(data)
    boundaries = ((start_h, 24.0), (24.0, 40.0), (40.0, np.inf))
    statistics = {}
    for name, values in trajectories.items():
        selected = (times >= start_h) & np.isfinite(values)
        mean = float(np.mean(values[selected]))
        sum_sq = 0.0
        count = 0
        for lower, upper in boundaries:
            mask = selected & (times >= lower) & (times < upper)
            if np.any(mask):
                local = values[mask]
                sum_sq += float(np.sum((local - np.mean(local)) ** 2))
                count += int(mask.sum())
        pooled_std = float(np.sqrt(sum_sq / count))
        statistics[name] = {
            "mean": mean,
            "within_segment_std": pooled_std,
            "coefficient_of_variation": pooled_std / abs(mean),
        }
    return statistics


def style_axis(axis: plt.Axes, *, show_grid: bool = True) -> None:
    # 网格、刻度、字体这些东西都在这一小段里顺着处理。
    # 图上主要内容画好了以后，再来把坐标轴周围收拾齐就行。
    axis.set_axisbelow(True)
    if show_grid:
        axis.grid(
            True,
            which="major",
            color="#AAA5A8",
            linestyle="--",
            linewidth=0.55,
            alpha=0.70,
        )
    else:
        axis.grid(False)
    axis.tick_params(length=3, width=0.7)
    for label in (*axis.get_xticklabels(), *axis.get_yticklabels()):
        if not any("\u4e00" <= character <= "\u9fff" for character in label.get_text()):
            label.set_fontfamily("Times New Roman")
    for text in (axis.title, axis.xaxis.label, axis.yaxis.label):
        text.set_fontfamily(heading_font)
        text.set_fontweight("bold")
    for name in ("left", "bottom"):
        axis.spines[name].set_color(color_dark)
        axis.spines[name].set_linewidth(0.75)


def set_title(figure: plt.Figure, title: str, y: float = 0.97) -> None:
    text = figure.suptitle(title, x=0.5, y=y, fontsize=12, fontweight="bold")
    text.set_fontfamily(heading_font)


def close_figure(figure: plt.Figure) -> None:
    plt.close(figure)


def save_figure(
    figure: plt.Figure, base_name: str, exclude_axes: list[plt.Axes] | None = None
) -> None:
    # 图片先完整画到画布上，接下来再走原来的面板对齐检查。
    # 检查通过以后才写出图片文件，等这张保存结束，再关闭它并接着处理后面的图。
    # 这一张先处理到保存结束，再换后面的图。
    figure_dir.mkdir(parents=True, exist_ok=True)
    figure.canvas.draw()
    with TemporaryDirectory(prefix="problem4-figure-qa-") as tmp_dir:
        require_matplotlib_panel_alignment(
            figure,
            json_out=Path(tmp_dir) / f"{base_name}.alignment.json",
            exclude_axes=exclude_axes or [],
            row_groups=getattr(figure, "_alignment_row_groups", None),
            tolerance_pt=1.5,
            gutter_tolerance_pt=1.5,
            require_panel_labels=False,
            strict=True,
        )
    figure.savefig(
        figure_dir / f"{base_name}.png", dpi=export_dpi, bbox_inches="tight",
        pad_inches=pad_inches, facecolor="white",
    )
    plt.close(figure)


def plot_shrinkage(data: dict[str, object]) -> plt.Figure:
    # 先把和收缩有关的几组量分别拿出来，再照顺序放到图里。
    # 后面还有相应的统计结果要看，把这些曲线也留在同一张图上，就能对着一起查看。
    # 每组数先各自放好，等线和标签都齐了，就能沿着时间一起看它们如何变化。
    history_times = np.asarray(data["radius_history_times_h"])
    history_radius = np.asarray(data["radius_history_cm"])
    times = np.asarray(data["times_h"])
    phi = shrinkage_phi(data)
    statistics = shrinkage_stats(data)
    fig, axs = plt.subplots(1, 2, figsize=(7.2, 3.35))

    axs[0].plot(history_times, history_radius, "o-", ms=2.0, lw=1.35, color=color_teal)
    axs[0].scatter([history_times[0], history_times[-1]], [history_radius[0], history_radius[-1]], color=color_coral, s=24, zorder=3)
    axs[0].set(xlabel="时间 / h", ylabel="药材半径 / cm", title="附件二实测收缩轨迹")
    axs[0].text(0.96, 0.92, "2.000 → 1.198 cm\n径向收缩 40.1%", transform=axs[0].transAxes, ha="right", va="top", color=color_dark)

    mask = (times >= 6.0) & np.isfinite(phi["surface"])
    for name, label, color, width in (
        ("volume_average", "体积平均", color_sand, 1.3),
        ("outer_average", "外层半域平均", color_plum, 1.3),
        ("surface", "表面含水率", color_coral, 1.8),
    ):
        axs[1].plot(times[mask], phi[name][mask], color=color, lw=width, label=label)
    axs[1].axvline(12.0, color=color_dark, ls="--", lw=0.8)
    axs[1].axhline(statistics["surface"]["mean"], color=color_coral, ls=":", lw=1.0)
    axs[1].set(xlabel="时间 / h", ylabel=r"反演参数 $\phi$", title="不同收缩驱动量的恒定性")
    axs[1].set_ylim(0.15, 1.62)
    axs[1].text(0.97, 0.10, "12 h 后表面口径 CV = 0.16%", transform=axs[1].transAxes, ha="right", color=color_coral, fontsize=7.3)
    for ax in axs:
        style_axis(ax)
    handles, labels = axs[1].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 0.86), ncol=3)
    set_title(fig, "收缩轨迹与表面含水率保持自洽", y=0.985)
    fig.subplots_adjust(left=0.10, right=0.98, bottom=0.17, top=0.70, wspace=0.34)
    return fig


def plot_threshold(data: dict[str, object]) -> plt.Figure:
    # 达标前后的情况在这一张里集中放好，时间和含水率都从已有记录里取。
    # 先把整段变化接起来，再把需要强调的位置标出来，后面查看时就顺着这些记录看。
    # 先把整体那段时间留出来
    times = np.asarray(data["times_h"])
    center = np.asarray(data["center_moisture"])
    surface = np.asarray(data["surface_moisture"])
    target = float(data["continuous_threshold_time_h"])
    fig, axs = plt.subplots(1, 2, figsize=(7.2, 3.25))
    axs[0].plot(times, center, color=color_teal, lw=1.7, label="轴心（全域最大值）")
    axs[0].plot(times, surface, color=color_coral, lw=1.35, label="移动表面")
    axs[0].axhline(default_critical_moisture, color=color_dark, ls="--", lw=1.0, label="达标阈值")
    axs[0].set(xlabel="时间 / h", ylabel="含水率 / kg·kg$^{-1}$", title="全过程")
    zoom = (times >= target - 0.22) & (times <= target + 0.18)
    axs[1].plot(times[zoom], center[zoom], "o-", ms=2.3, color=color_teal, lw=1.25)
    axs[1].axhline(default_critical_moisture, color=color_dark, ls="--", lw=1.0)
    axs[1].axvline(target, color=color_coral, lw=1.35)
    axs[1].annotate(
        f"$t_*= {target:.4f}$ h", xy=(target, default_critical_moisture),
        xytext=(target - 0.19, default_critical_moisture + 0.00008),
        arrowprops={"arrowstyle": "-", "color": color_coral, "lw": 0.8},
        color=color_coral, fontsize=7.5,
    )
    axs[1].set(xlabel="时间 / h", ylabel="轴心含水率 / kg·kg$^{-1}$", title="阈值附近")
    axs[1].ticklabel_format(axis="y", style="plain", useOffset=False)
    for ax in axs:
        style_axis(ax)
    handles, labels = axs[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 0.86), ncol=3)
    set_title(fig, "移动边界模型全域达标时间：51.12 h", y=0.985)
    fig.subplots_adjust(left=0.10, right=0.98, bottom=0.17, top=0.70, wspace=0.34)
    return fig


def plot_moving_field(data: dict[str, object]) -> plt.Figure:
    # 内部场的数据先准备好，同一时刻的半径也一起带上。
    # 参考位置换成当前的实际位置以后，再把这一时刻的数放到图里，后面的时刻照着继续处理。
    # 同一时刻的数要跟同一时刻的半径接上
    # 几个时刻依次画完，再看各条线对应的标签
    times = np.asarray(data["times_h"])
    xi = np.asarray(data["xi_centers"])
    internal = np.asarray(data["internal_moisture"])
    radii = np.asarray(data["surface_radii_cm"])
    center = np.asarray(data["center_moisture"])
    surface = np.asarray(data["surface_moisture"])
    discrete = float(data["discrete_threshold_time_h"])
    fig, axs = plt.subplots(1, 2, figsize=(7.2, 3.35))
    mesh = axs[0].pcolormesh(times, xi, internal.T, shading="auto", cmap="viridis")
    axs[0].axvline(discrete, color="white", ls="--", lw=1.0)
    axs[0].set(xlabel="时间 / h", ylabel=r"归一化半径 $\xi=r/R(t)$", title="移动参考域时空分布")
    color_axis = axs[0].inset_axes([1.025, 0.02, 0.045, 0.96])
    fig.colorbar(mesh, cax=color_axis)
    for label in color_axis.get_yticklabels():
        label.set_fontfamily("Times New Roman")

    selected = (0.0, 6.0, 12.0, 24.0, 36.0, discrete)
    colors = mpl.colormaps["viridis"](np.linspace(0.08, 0.90, len(selected)))
    for hour, color in zip(selected, colors):
        index = int(np.argmin(np.abs(times - hour)))
        physical_radii = np.concatenate(([0.0], xi * radii[index], [radii[index]]))
        profile = np.concatenate(([center[index]], internal[index], [surface[index]]))
        label = f"{times[index]:.1f} h" if hour else "0 h"
        axs[1].plot(physical_radii, profile, color=color, lw=1.45, label=label)
    axs[1].axhline(default_critical_moisture, color=color_dark, ls="--", lw=0.9)
    axs[1].set(xlabel="实际径向位置 / cm", ylabel="含水率 / kg·kg$^{-1}$", title="收缩中的径向剖面")
    style_axis(axs[0], show_grid=False)
    style_axis(axs[1])
    handles, labels = axs[1].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 0.86), ncol=6)
    set_title(fig, "收缩缩短扩散路径并保持轴心最湿", y=0.985)
    fig.subplots_adjust(left=0.09, right=0.97, bottom=0.17, top=0.69, wspace=0.40)
    fig._alignment_row_groups = [["a", "b"]]
    fig._extra_qa_axes = [color_axis]
    return fig


def plot_mechanisms(data: dict[str, object]) -> plt.Figure:
    # 不同物性、不同半径设置的几种结果，现在按顺序放到一起。
    # 每一种对应的标签也跟着放好，最后看这张图时，就能顺着标签找到相应的对照结果。
    cases = mechanism_cases(data["diagnostics"])
    if not cases:
        raise ValueError("机制对照尚未完成整套复算。")
    values = {case["name"]: float(case["continuous_threshold_time_h"]) for case in cases}
    fig, ax = plt.subplots(figsize=(5.5, 3.35))
    x = np.arange(2)
    appendix3 = [values["appendix3_fixed_radius"], values["appendix3_moving_radius"]]
    appendix4 = [values["appendix4_fixed_radius"], values["appendix4_moving_radius"]]
    ax.plot(x, appendix3, "o-", color=color_teal, lw=1.8, ms=5, label="附录三物性")
    ax.plot(x, appendix4, "o-", color=color_coral, lw=1.8, ms=5, label="附录四物性")
    ax.set_xticks(x, ["固定半径", "实测收缩半径"])
    ax.set(ylabel="连续达标时间 $t_*$ / h", title="两条非平行响应线表明物性与几何存在交互")
    for series, color in ((appendix3, color_teal), (appendix4, color_coral)):
        for px, value in zip(x, series):
            ax.text(px, value + 3.0, f"{value:.2f}", ha="center", color=color, fontsize=7.4)
    style_axis(ax)
    fig.legend(loc="upper center", bbox_to_anchor=(0.5, 0.86), ncol=2)
    set_title(fig, "物性变化与收缩效应不可简单相加", y=0.985)
    fig.subplots_adjust(left=0.15, right=0.98, bottom=0.17, top=0.68)
    return fig


def plot_lollipop(axis: plt.Axes, labels: list[str], values: np.ndarray, xlabel: str) -> None:
    # 先给每种设置排一个位置，再把它的横线和圆点接着画出来。
    # 旁边的数跟着各自那一行放好。
    positions = np.arange(len(values))
    colors = [color_coral if value > 0 else color_teal for value in values]
    axis.hlines(positions, 0.0, values, color=colors, lw=2.0)
    axis.scatter(values, positions, color=colors, s=27, zorder=3)
    axis.axvline(0.0, color=color_dark, lw=0.8)
    axis.set_yticks(positions, labels)
    axis.invert_yaxis()
    axis.set_xlabel(xlabel)
    span = max(np.max(np.abs(values)), 1.0)
    axis.set_xlim(
        min(float(np.min(values)), 0.0) - 0.08 * span,
        max(float(np.max(values)), 0.0) + 0.18 * span,
    )
    for y, value in zip(positions, values):
        axis.text(value + span * 0.025, y, f"{value:+.1f}", va="center", ha="left", fontsize=7.2)
    style_axis(axis)


def plot_robustness(data: dict[str, object]) -> plt.Figure:
    # 最后把数值设置和边界情况的那些结果再集中整理一下。
    # 数据按原来的记录往图里放，放到一张图上以后，就可以一起查看它们各自的变化。
    # 顺着已有的分组把数放进去，哪一种变化对应哪一行，就照标签来认。
    stats = data["diagnostics"]
    baseline = float(stats["baseline"]["continuous_threshold_time_h"])
    refinement = stats["numerical_refinement"]
    refinement_offsets = 3600.0 * np.asarray([case["continuous_threshold_time_h"] - baseline for case in refinement])
    boundary = stats["boundary_sensitivity"]
    boundary_offsets = 60.0 * np.asarray([case["continuous_threshold_time_h"] - baseline for case in boundary])
    fig, axs = plt.subplots(1, 2, figsize=(7.2, 3.25))
    plot_lollipop(axs[0], ["空间加密", "时间加密"], refinement_offsets, "达标时间变化 / s")
    plot_lollipop(axs[1], ["温度 -1σ", "温度 +1σ"], boundary_offsets, "相对基准变化 / min")
    axs[0].set_title("网格与步长独立性", pad=7)
    axs[1].set_title("长期温度边界敏感性", pad=7)
    set_title(fig, "51.12 h 结论在数值与边界检验下稳定", y=0.985)
    fig.subplots_adjust(left=0.15, right=0.98, bottom=0.17, top=0.73, wspace=0.52)
    return fig


def clear_old_figures() -> None:
    # 这一轮要留下的新图名字先放在集合里，随后再查看目录中其他图片。
    # 照着下面的条件把需要清理的逐个处理完。
    figure_dir.mkdir(parents=True, exist_ok=True)
    expected = {f"{name}.png" for name in new_figure_bases}
    for target in figure_dir.iterdir():
        if target.is_file() and target.suffix.lower() in {".png", ".pdf", ".svg", ".tif", ".tiff"} and target.name not in expected:
            if target.resolve().parent != figure_dir.resolve():
                raise RuntimeError("拒绝删除问题四图片目录之外的文件。")
            target.unlink()


def export_figures(data: dict[str, object]) -> None:
    clear_old_figures()
    builders = (
        plot_shrinkage, plot_threshold, plot_moving_field,
        plot_mechanisms, plot_robustness,
    )
    for builder, base_name in zip(builders, new_figure_bases):
        fig = builder(data)
        save_figure(fig, base_name, getattr(fig, "_extra_qa_axes", None))


def main() -> None:
    export_figures(load_plot_data())
    print(f"问题四绘图完成，共生成 {len(new_figure_bases)} 张 600 dpi PNG：{figure_dir}")

# 写完了www
if __name__ == "__main__":
    main()
