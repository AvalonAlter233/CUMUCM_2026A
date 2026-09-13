from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook

# 默认参数呃呃
project_root = Path(__file__).resolve().parent
input_file = project_root / "附件" / "附件1.xlsx"
output_file = project_root / "附件" / "附件3" / "result3.xlsx"
diagnostics_file = project_root / "附件" / "附件3" / "result3_diagnostics.json"

pellet_radius = 0.02
convective_heat_coeff = 25.0
convective_mass_coeff = 8.0e-7
initial_temperature = 28.0
initial_moisture = 2.55
default_critical_moisture = 0.15

default_intervals = 320
output_spacing = 0.001
default_time_step = 30.0
default_report_interval = 60.0
max_simulation_time = 5 * 24 * 3600.0

convergence_tol = 1.0e-9
max_iterations = 200
relaxation_factor = 0.6
use_plateau_mean = True


class DryingRoomBoundary(NamedTuple):

    times: np.ndarray
    temperatures: np.ndarray
    moisture: np.ndarray
    plateau_temperature: float
    plateau_moisture: float
    plateau_window: float
    plateau_sample_count: int
    plateau_temperature_std: float
    plateau_moisture_std: float


class BoundaryScenario(NamedTuple):

    name: str
    label: str
    boundary: DryingRoomBoundary


class CellCenteredGrid(NamedTuple):

    faces: np.ndarray
    centers: np.ndarray
    volumes: np.ndarray
    spacing: float


def read_room_data(
    path: Path,
    plateau_window: float = 3600.0,
    temperature_offset: float = 0.0,
    moisture_offset: float = 0.0,
    use_last_sample: bool = False,
) -> DryingRoomBoundary:
    # 这一次还要往后算很长一段时间，光把表读出来还不够
    # 表里的原始数据先留好，等超过表格时间范围以后要用的那几个值，也在这里一起准备。
    # 前面有实测数据的那段先照常拿着，后面延长出去的部分也得提前备好。
    # 这里先把两头都安排上，主循环走到哪里再取哪里
    book = load_workbook(path, data_only=True, read_only=True)
    rows = [
        row for row in book.active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一缺少有效边界数据。")
    if not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError("附件一时间必须严格递增。")

    if plateau_window <= 0:
        raise ValueError("平台均值窗口必须为正数。")
    # 先从数据末尾挑出这一段，后面几个统计都用它。
    # 要取均值的时候从这里取，要看波动的时候也看这一段，先把范围选好再往下算。
    # 后面说到末尾这段数据时，指的就是这里选出来的这些行
    stable = data[:, 0] >= data[-1, 0] - plateau_window
    if use_last_sample or not use_plateau_mean:
        temp_level = float(data[-1, 1])
        water_level = float(data[-1, 2])
        n_samples = 1
        temp_std = 0.0
        water_std = 0.0
    else:
        tail_values = data[stable, 1:3]
        temp_level = float(np.mean(tail_values[:, 0]))
        water_level = float(np.mean(tail_values[:, 1]))
        n_samples = int(tail_values.shape[0])
        ddof = 1 if n_samples > 1 else 0
        temp_std = float(np.std(tail_values[:, 0], ddof=ddof))
        water_std = float(np.std(tail_values[:, 1], ddof=ddof))
    return DryingRoomBoundary(
        data[:, 0],
        data[:, 1],
        data[:, 2],
        temp_level + temperature_offset,
        water_level + moisture_offset,
        plateau_window,
        n_samples,
        temp_std,
        water_std,
    )


def boundary_cases(path: Path) -> list[BoundaryScenario]:
    # 准备比较的几种情况先逐个列出来，后面会照这个顺序去算。
    # 名称也和数据一起带上，等结果放到一起以后，看名字就知道对应的是哪一种设置。
    # 先把要比较的几份设置放成一排，后面算的时候就能挨个接着来
    baseline = read_room_data(path)
    return [
        BoundaryScenario(
            "last_30_min_mean",
            "末 30 min 均值",
            read_room_data(path, plateau_window=1800.0),
        ),
        BoundaryScenario(
            "last_sample",
            "最后采样点",
            read_room_data(path, use_last_sample=True),
        ),
        BoundaryScenario(
            "temperature_minus_1sigma",
            "温度平台 -1σ",
            read_room_data(
                path, temperature_offset=-baseline.plateau_temperature_std
            ),
        ),
        BoundaryScenario(
            "temperature_plus_1sigma",
            "温度平台 +1σ",
            read_room_data(
                path, temperature_offset=baseline.plateau_temperature_std
            ),
        ),
        BoundaryScenario(
            "moisture_minus_1sigma",
            "环境水分平台 -1σ",
            read_room_data(
                path, moisture_offset=-baseline.plateau_moisture_std
            ),
        ),
        BoundaryScenario(
            "moisture_plus_1sigma",
            "环境水分平台 +1σ",
            read_room_data(
                path, moisture_offset=baseline.plateau_moisture_std
            ),
        ),
    ]


def threshold_time(
    previous_time: float,
    previous_value: float,
    current_time: float,
    current_value: float,
    threshold: float,
) -> float:
    # 前一个时刻还在阈值上面，到了后一个时刻已经到下面了。
    # 就在这两个记录之间再插一下，先给出一个连续时间的估计，表格照旧按输出间隔来记。
    # 夹在两个时刻中间的这一点先估出来，等写结果时再和记录时刻一起看
    if current_time <= previous_time:
        raise ValueError("当前时刻必须晚于前一时刻。")
    if not (previous_value >= threshold and current_value < threshold):
        raise ValueError("相邻数值没有从上方严格穿越阈值。")
    return previous_time + (
        (previous_value - threshold) / (previous_value - current_value)
    ) * (current_time - previous_time)


def balance_error(
    initial_inventory: float,
    current_inventory: float,
    cumulative_outflow: float,
) -> float:
    # 先把起始的量、现在的量和累计流出的量放到同一处对一下。
    # 算出来的检查数交回外面，后面记诊断的时候会接着用
    if initial_inventory <= 0:
        raise ValueError("初始积分必须为正数。")
    return abs(
        initial_inventory - current_inventory - cumulative_outflow
    ) / initial_inventory


def boundary_value(
    current_time: float,
    boundary: DryingRoomBoundary,
) -> tuple[float, float]:
    # 先看看当前已经算到什么时刻，再决定从哪里拿边界值。
    # 还在表格范围里就用表里的数据插值，走过末尾以后就接着用前面备好的平台值。
    # 前一段走完就接后一段，取值的位置在这里换一下，外面的循环继续往下走
    if current_time <= boundary.times[-1]:
        return (
            float(np.interp(current_time, boundary.times, boundary.temperatures)),
            float(np.interp(current_time, boundary.times, boundary.moisture)),
        )
    return boundary.plateau_temperature, boundary.plateau_moisture


def density(moisture: np.ndarray) -> np.ndarray:
    return 650.0 + 128.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    return 1450.0 + 2736.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    return 0.21 + 0.38 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    c_safe = np.maximum(moisture, 1.0e-12)
    temp_k = temperature_celsius + 273.15
    return (
        2.4e-3
        * np.exp(-0.45 / c_safe)
        * np.exp(-3850.0 / temp_k)
    )


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    return 2.0 * left * right / np.maximum(left + right, 1.0e-30)


def build_grid(intervals: int) -> CellCenteredGrid:
    # 后面要用的网格位置在这里一次分好。
    # 这些位置留在返回的对象里，主循环拿到以后一直用这份，接下来主要更新位置上的温度和水分。
    # 这一份网格后面要反复用，先在这里准备一次，接下来一直带着它算
    if intervals < 4:
        raise ValueError("内部径向区间数至少为 4。")
    faces = np.linspace(0.0, pellet_radius, intervals + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    vol = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return CellCenteredGrid(
        faces,
        centers,
        vol,
        pellet_radius / intervals,
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    lower = lower.astype(float, copy=True)
    diagonal = diagonal.astype(float, copy=True)
    upper = upper.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    for index in range(1, len(diagonal)):
        factor = lower[index - 1] / diagonal[index - 1]
        diagonal[index] -= factor * upper[index - 1]
        right_hand_side[index] -= factor * right_hand_side[index - 1]
    sol = np.empty_like(right_hand_side)
    sol[-1] = right_hand_side[-1] / diagonal[-1]
    for index in range(len(diagonal) - 2, -1, -1):
        sol[index] = (
            right_hand_side[index] - upper[index] * sol[index + 1]
        ) / diagonal[index]
    return sol


def assemble_system(
    old_values: np.ndarray,
    storage_capacity: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 一轮方程需要的东西在这里依次接起来，先内部，再到外面的边界。
    # 几条对角线和右边的数组都准备好以后，接着交给下面的求解步骤
    storage = storage_capacity * grid.volumes / time_step
    coef_face = harmonic_mean(
        transport_coefficients[:-1], transport_coefficients[1:]
    )
    g_inner = (
        grid.faces[1:-1] * coef_face / grid.spacing
    )
    diag = storage.copy()
    diag[:-1] += g_inner
    diag[1:] += g_inner
    lower = -g_inner.copy()
    upper = -g_inner.copy()
    rhs = storage * old_values

    coef_last = max(float(transport_coefficients[-1]), 1.0e-30)
    h_eff = 1.0 / (
        1.0 / external_transfer_coefficient
        + 0.5 * grid.spacing / coef_last
    )
    g_surf = pellet_radius * h_eff
    diag[-1] += g_surf
    rhs[-1] += g_surf * external_value
    return lower, diag, upper, rhs


def advance_temperature(
    old_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_temperature: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> np.ndarray:
    # 温度这一步要用的物性先从当前参考含水率里算出来。
    # 把这一轮的系数都带齐，再求出温度，拿到结果后回外面继续下一段
    lin_sys = assemble_system(
        old_temperature,
        density(reference_moisture) * heat_capacity(reference_moisture),
        thermal_conductivity(reference_moisture),
        convective_heat_coeff,
        room_temperature,
        grid,
        time_step,
    )
    return solve_tridiagonal(*lin_sys)


def advance_moisture(
    old_moisture: np.ndarray,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_moisture: float,
    grid: CellCenteredGrid,
    time_step: float,
) -> np.ndarray:
    # 水分这一部分接着准备自己的系数，环境里的水分值也一并带上。
    # 先把当前这一步求完，结果返回以后还会和温度那边一起继续更新。
    lin_sys = assemble_system(
        old_moisture,
        np.ones_like(old_moisture),
        moisture_diffusivity(reference_moisture, reference_temperature),
        convective_mass_coeff,
        room_moisture,
        grid,
        time_step,
    )
    return solve_tridiagonal(*lin_sys)


def sample_field(
    cell_values: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    output_nodes: np.ndarray,
    grid: CellCenteredGrid,
) -> np.ndarray:
    # 输出的几个位置先从内部数据里取，中心和表面再照下面的做法分别补一下
    # 这一排数拼完整了再返回，后面写表时就照它的顺序来。
    node_values = np.interp(output_nodes, grid.centers, cell_values)
    node_values[0] = (9.0 * cell_values[0] - cell_values[1]) / 8.0

    coef_last = max(float(transport_coefficients[-1]), 1.0e-30)
    g_half = 2.0 * coef_last / grid.spacing
    node_values[-1] = (
        g_half * cell_values[-1]
        + external_transfer_coefficient * external_value
    ) / (g_half + external_transfer_coefficient)
    return node_values


def solve_question3(
    internal_intervals: int = default_intervals,
    time_step: float = default_time_step,
    report_interval: float = default_report_interval,
    max_time: float = max_simulation_time,
    boundary: DryingRoomBoundary | None = None,
    critical_moisture: float = default_critical_moisture,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, dict]:
    save_every = int(round(report_interval / time_step))
    n_steps = int(round(max_time / time_step))
    if save_every < 1 or not np.isclose(
        save_every * time_step, report_interval
    ):
        raise ValueError("输出间隔必须是内部时间步长的整数倍。")
    if n_steps < 1 or not np.isclose(n_steps * time_step, max_time):
        raise ValueError("最大时间必须是内部时间步长的整数倍。")

    if boundary is None:
        boundary = read_room_data(input_file)
    mesh = build_grid(internal_intervals)
    r_nodes = np.arange(
        0.0, pellet_radius + 0.5 * output_spacing, output_spacing
    )
    temp = np.full(internal_intervals, initial_temperature)
    water = np.full(internal_intervals, initial_moisture)

    # 先放几个空列表，等有结果需要保存时再往里面加。
    # 还有一些用来判断和记账的变量，也先给个起始值，等进入循环以后再跟着每一步更新。
    # 这些量开始时先各自占个位置
    time_log: list[float] = []
    water_log: list[np.ndarray] = []
    iter_log: list[int] = []
    t_prev = 0.0
    c_max_prev = initial_moisture
    t_cross = None
    c_before = None
    c_after = None
    stock_init = float(np.dot(water, mesh.volumes))
    stock_now = stock_init
    outflow_sum = 0.0
    max_balance_err = 0.0
    center_controls_every_step = True
    radially_nonincreasing_every_step = True
    checked_steps = 0

    # 下面就开始往后推进时间了，一步算完再接下一步。
    # 温度和含水率都要在这一步里更新好，之后再看看全域的含水率有没有达到要求。
    for step in range(1, n_steps + 1):
        t_now = step * time_step
        temp_air, water_air = boundary_value(
            t_now, boundary
        )
        temp_iter = temp.copy()
        water_iter = water.copy()

        for iteration in range(1, max_iterations + 1):
            temp_trial = advance_temperature(
                temp,
                water_iter,
                temp_air,
                mesh,
                time_step,
            )
            water_trial = advance_moisture(
                water,
                temp_trial,
                water_iter,
                water_air,
                mesh,
                time_step,
            )
            temp_new = (
                relaxation_factor * temp_trial
                + (1.0 - relaxation_factor) * temp_iter
            )
            water_new = (
                relaxation_factor * water_trial
                + (1.0 - relaxation_factor) * water_iter
            )
            err_t = np.max(
                np.abs(temp_new - temp_iter)
                / (1.0 + np.abs(temp_new))
            )
            err_c = np.max(
                np.abs(water_new - water_iter)
                / (1.0 + np.abs(water_new))
            )
            temp_iter = temp_new
            water_iter = water_new
            if max(err_t, err_c) < convergence_tol:
                break
        else:
            raise RuntimeError(
                f"{t_now:.0f} s 未收敛："
                f"{err_t:.3e}, {err_c:.3e}"
            )

        temp = temp_iter
        water = water_iter
        if not np.all(np.isfinite(water)) or np.min(water) <= 0:
            raise FloatingPointError(f"{t_now:.0f} s 含水率异常。")

        d_cell = moisture_diffusivity(water, temp)
        d_last = max(float(d_cell[-1]), 1.0e-30)
        hm_eff = 1.0 / (
            1.0 / convective_mass_coeff
            + 0.5 * mesh.spacing / d_last
        )
        flux_surf = (
            pellet_radius
            * hm_eff
            * (float(water[-1]) - water_air)
        )
        # 此处懒得写了
        stock_next = float(np.dot(water, mesh.volumes))
        balance_err = abs(
            stock_next - stock_now + time_step * flux_surf
        ) / stock_init
        max_balance_err = max(
            max_balance_err, balance_err
        )
        outflow_sum += time_step * flux_surf
        stock_now = stock_next

        c_nodes = sample_field(
            water,
            d_cell,
            convective_mass_coeff,
            water_air,
            r_nodes,
            mesh,
        )
        c_max_now = max(
            float(np.max(water)),
            float(np.max(c_nodes)),
        )
        axis_val = float(c_nodes[0])
        center_controls_every_step = center_controls_every_step and bool(
            c_max_now <= axis_val + 1.0e-12
        )
        radially_nonincreasing_every_step = (
            radially_nonincreasing_every_step
            and bool(np.all(np.diff(water) <= 1.0e-12))
            and bool(np.all(np.diff(c_nodes) <= 1.0e-12))
        )
        checked_steps += 1
        if (
            t_cross is None
            and c_max_prev >= critical_moisture
            and c_max_now < critical_moisture
        ):
            # 第一次走过阈值的时候，先把这个时间留下来。
            # 后面还要继续等到合适的输出时刻，这个第一次跨过的记录就先放着，之后整理结果时再取
            t_cross = threshold_time(
                t_prev,
                c_max_prev,
                t_now,
                c_max_now,
                critical_moisture,
            )
            c_before = c_max_prev
            c_after = c_max_now

        # 此处再看一下是不是到了该保存的时刻，到了就往下添一行。
        # 这一行保存完以后，如果此时也已经达标，那这一趟便可停于此。
        if step % save_every == 0:
            time_log.append(t_now)
            water_log.append(c_nodes.copy())
            if c_max_now < critical_moisture:
                iter_log.append(iteration)
                break

        iter_log.append(iteration)
        t_prev = t_now
        c_max_prev = c_max_now
    else:
        raise RuntimeError(f"{max_time / 3600:.1f} h 内未达到阈值。")

    water_field = np.asarray(water_log)
    # 最后把达标时间、阈值前后的数和那些检查量放到一起。
    # 这样返回的是一整份记录，后面要打印或者画图，就从里面挑需要的那一项。
    diagnostics = {
        "continuous_threshold_time_s": float(t_cross),
        "discrete_threshold_time_s": float(time_log[-1]),
        "maximum_moisture_before": float(c_before),
        "maximum_moisture_after": float(c_after),
        "center_controls_threshold": center_controls_every_step,
        "all_domain_checked_every_step": checked_steps == len(iter_log),
        "radially_nonincreasing_every_step": radially_nonincreasing_every_step,
        "moisture_balance_relative_imbalance": balance_error(
            stock_init, stock_now, outflow_sum
        ),
        # 好多东西，注释就不写了这里
        "maximum_step_balance_relative_residual": max_balance_err,
        "initial_moisture_integral": stock_init,
        "final_moisture_integral": stock_now,
        "cumulative_boundary_outflow": outflow_sum,
        "maximum_picard_iterations": int(max(iter_log)),
        "mean_picard_iterations": float(np.mean(iter_log)),
        "internal_intervals": internal_intervals,
        "internal_spacing_cm": 100.0 * mesh.spacing,
        "time_step_s": time_step,
        "report_interval_s": report_interval,
        "critical_moisture": critical_moisture,
        "plateau_temperature_c": boundary.plateau_temperature,
        "plateau_moisture_kgkg": boundary.plateau_moisture,
        "plateau_window_s": boundary.plateau_window,
        "plateau_sample_count": boundary.plateau_sample_count,
    }
    return (
        np.asarray(time_log),
        r_nodes,
        water_field,
        diagnostics,
    )


def write_results(
    times: np.ndarray,
    nodes: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    # 先照着模板把表头放好，时间放在每一行的最前面。
    # 这一时刻各个位置上的含水率再接到后面，整张表就有了。
    # 时间和位置都已经有顺序，写的时候顺着接就好了，填完再统一保存。
    book = load_workbook(output_file)
    sheet = book.worksheets[0]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    sheet.cell(1, 1).value = "时间\\到药材中心的距离"
    sheet.cell(1, 1).number_format = "@"
    for column, radius in enumerate(nodes, start=2):
        sheet.cell(1, column).value = round(float(radius * 100.0), 1)
        sheet.cell(1, column).number_format = "0.0"
    for row, t_now in enumerate(times, start=2):
        sheet.cell(row, 1).value = int(round(float(t_now)))
        sheet.cell(row, 1).number_format = "0"
        for column, value in enumerate(moisture_field[row - 2], start=2):
            sheet.cell(row, column).value = float(value)
            sheet.cell(row, column).number_format = "0.0000"
    book.save(output_file)


def main() -> None:
    # 先把主结果算出来并写好，接着把几项时间和检查记录打印一下。
    parser = argparse.ArgumentParser(description="问题三全域含水率达标时间求解")
    parser.add_argument(
        "--verification",
        action="store_true",
        help="额外复算边界情景及空间、时间加密案例并保存诊断 JSON",
    )
    args = parser.parse_args()

    baseline_boundary = read_room_data(input_file)
    times, nodes, moisture_field, diagnostics = solve_question3(
        boundary=baseline_boundary
    )
    write_results(times, nodes, moisture_field)
    print(f"问题三重算完成：{output_file}")
    print(
        f"连续阈值：{diagnostics['continuous_threshold_time_s'] / 3600:.4f} h；"
        f"首次 60 s 离散达标："
        f"{diagnostics['discrete_threshold_time_s'] / 3600:.4f} h"
    )
    print(
        f"内部网格 {diagnostics['internal_intervals']} 个区间，"
        f"Δr={diagnostics['internal_spacing_cm']:.5f} cm；"
        f"输出尺寸 {moisture_field.shape[0]}×{moisture_field.shape[1]}"
    )
    print(
        f"阈值前后最大含水率："
        f"{diagnostics['maximum_moisture_before']:.10f} / "
        f"{diagnostics['maximum_moisture_after']:.10f}"
    )
    print(
        f"离散 C 方程累计相对收支不平衡："
        f"{diagnostics['moisture_balance_relative_imbalance']:.3e}；"
        f"最大 Picard 迭代次数：{diagnostics['maximum_picard_iterations']}"
    )

    # 如果这次只算主结果，到这里就可以结束。
    # 开了验证选项的话，下面还有几种情况要继续算，就让程序接着往下走。
    # 如果走到了下面，那就让额外的几组也接着算完。
    # 每跑完一种记下一种，最后再把这些记录放到一起看。
    if not args.verification:
        return

    def record_case(name: str, label: str, case_diagnostics: dict) -> dict:
        # 每种情况的诊断先复制一份出来，再把名字和小时单位的时间补进去。
        # 补齐以后这一条记录就比较完整了，后面可以直接和别的情况放到一起。
        record = dict(case_diagnostics)
        record["name"] = name
        record["label"] = label
        record["continuous_threshold_time_h"] = (
            case_diagnostics["continuous_threshold_time_s"] / 3600.0
        )
        record["discrete_threshold_time_h"] = (
            case_diagnostics["discrete_threshold_time_s"] / 3600.0
        )
        return record

    report = {
        "scope": "离散方程诊断与情景敏感性；不是概率置信区间或真实质量守恒证明。",
        "baseline": record_case(
            "last_1_hour_mean", "末 1 h 均值（基准）", diagnostics
        ),
        "boundary_scenarios": [],
        "numerical_refinement": [],
    }

    for scenario in boundary_cases(input_file):
        print(f"复算边界情景：{scenario.label}")
        _, _, _, case_diagnostics = solve_question3(boundary=scenario.boundary)
        report["boundary_scenarios"].append(
            record_case(scenario.name, scenario.label, case_diagnostics)
        )

    refinement_cases = [
        ("space_refined", "空间加密 N=640, Δt=30 s", 640, 30.0),
        ("time_refined", "时间加密 N=320, Δt=15 s", 320, 15.0),
    ]
    for name, label, intervals, time_step in refinement_cases:
        print(f"复算数值加密：{label}")
        _, _, _, case_diagnostics = solve_question3(
            internal_intervals=intervals,
            time_step=time_step,
            boundary=baseline_boundary,
        )
        report["numerical_refinement"].append(
            record_case(name, label, case_diagnostics)
        )

    with diagnostics_file.open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    print(f"验证诊断已保存：{diagnostics_file}")


if __name__ == "__main__":
    main()
# 呃啊，写完了嘻嘻（）