"""问题一：固定圆柱药材的一维径向传热传质计算。

只读取：附件/附件1.xlsx、附件/附件3/result1.xlsx
只写入：附件/附件3/result1.xlsx

模型约定：
1. 计算对象为圆柱中部截面，忽略端面影响；该假设应另用二维模型检验。
2. C 表示药材干基含水率，题给 D(C) 视为与该变量配套的有效扩散系数。
3. 附件一的水分浓度直接作为题设等效外界水分变量，不额外引入未标定的吸附等温线。
4. 题目没有给出潜热和热扩散参数，基准模型不加入潜热项与 Soret 项。

数值方法：单元中心有限体积法、后向 Euler、Picard 迭代。
内部计算网格独立于题目要求的 0.1 cm 输出网格；外表面同时计入
最后半个控制体的内部传递阻力和 Robin 边界阻力。
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook


# 路径以脚本所在目录为基准，换电脑或换工作目录时无需修改
PROJECT_ROOT = Path(__file__).resolve().parent
INPUT_FILE = PROJECT_ROOT / "附件" / "附件1.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "附件" / "附件3" / "result1.xlsx"

# 药材物性与表面传递参数
PELLET_RADIUS = 0.02             # 药材半径，单位 m
PELLET_DENSITY = 820.0           # 药材表观密度，单位 kg/m^3
PELLET_HEAT_CAPACITY = 2600.0    # 药材表观比热容，单位 J/(kg·K)
THERMAL_CONDUCTIVITY = 0.36      # 导热系数，单位 W/(m·K)
CONVECTIVE_HEAT_COEFF = 25.0     # 对流换热系数，单位 W/(m^2·K)
CONVECTIVE_MASS_COEFF = 8.0e-7   # 对流传质系数，单位 m/s

# 初始条件
INITIAL_TEMPERATURE = 28.0       # 初始温度，单位 ℃
INITIAL_MOISTURE = 2.55          # 初始干基含水率，单位 kg/kg

# 时间、内部空间网格和输出设置
END_TIME = 1800.0                # 计算终点，单位 s
TIME_STEP = 1.0                  # 内部时间步长，单位 s
INTERNAL_CELL_COUNT = 160        # 内部有限体积单元数，不等于输出点数
OUTPUT_TIME_INTERVAL = 1.0       # 结果输出间隔，单位 s
OUTPUT_RADII = np.linspace(0.0, PELLET_RADIUS, 21)  # 每隔 0.1 cm 输出

# 非线性迭代设置
CONVERGENCE_TOL = 1.0e-10
MAX_ITERATIONS = 30


class DryingRoomBoundary(NamedTuple):
    """附件一给出的烘房边界时间序列。"""

    times: np.ndarray
    temperatures: np.ndarray
    moisture: np.ndarray


class RadialGrid(NamedTuple):
    """单元中心径向网格及每个柱壳的几何积分因子。"""

    centers: np.ndarray           # 控制体中心半径
    faces: np.ndarray             # 控制体界面半径，包含 0 和 R
    volume_factors: np.ndarray    # 积分 \int r dr 的离散权重
    spacing: float


class SimulationDetails(NamedTuple):
    """内部计算结果及数值诊断。"""

    times: np.ndarray
    output_radii: np.ndarray
    temperature_field: np.ndarray
    moisture_field: np.ndarray
    average_temperature: np.ndarray
    average_moisture: np.ndarray
    diagnostics: dict[str, float]


def read_drying_boundary(path: Path) -> DryingRoomBoundary:
    """读取附件一，并检查边界数据是否覆盖计算时段。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    rows = [
        row for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)

    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一必须包含时间、温度、水分浓度三列有效数据。")
    if not np.all(np.isfinite(data[:, :3])):
        raise ValueError("附件一的前三列存在空值或非有限数值。")

    times = data[:, 0]
    if not np.all(np.diff(times) > 0):
        raise ValueError("附件一中的时间必须严格递增。")

    return DryingRoomBoundary(
        times=times,
        temperatures=data[:, 1],
        moisture=data[:, 2],
    )


def linear_interp(t: float, times: np.ndarray, values: np.ndarray) -> float:
    """对附件一边界数据进行分段线性插值，不允许外推。"""
    if t < times[0] or t > times[-1]:
        raise ValueError(
            f"计算时刻 {t:g} s 超出附件一范围 "
            f"[{times[0]:g}, {times[-1]:g}] s。"
        )
    return float(np.interp(t, times, values))


def moisture_diffusivity(moisture: np.ndarray) -> np.ndarray:
    """题给有效水分扩散系数 D(C)，单位 m^2/s。"""
    safe_moisture = np.maximum(moisture, 1.0e-12)
    return 7.0e-9 * np.exp(-0.89 / safe_moisture)


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """用调和平均计算相邻控制体界面处的传输系数。"""
    denominator = np.maximum(left + right, 1.0e-30)
    return 2.0 * left * right / denominator


def build_radial_grid(cell_count: int) -> RadialGrid:
    """建立等距单元中心有限体积网格。"""
    if cell_count < 2:
        raise ValueError("径向有限体积单元数至少为 2。")

    faces = np.linspace(0.0, PELLET_RADIUS, cell_count + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    volume_factors = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return RadialGrid(
        centers=centers,
        faces=faces,
        volume_factors=volume_factors,
        spacing=float(faces[1] - faces[0]),
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    """Thomas 算法求解三对角线性方程组。"""
    diagonal = diagonal.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    upper = upper.astype(float, copy=True)

    for i in range(1, len(diagonal)):
        if abs(diagonal[i - 1]) < 1.0e-30:
            raise FloatingPointError("三对角方程组出现近零主对角元。")
        factor = lower[i - 1] / diagonal[i - 1]
        diagonal[i] -= factor * upper[i - 1]
        right_hand_side[i] -= factor * right_hand_side[i - 1]

    solution = np.empty_like(right_hand_side)
    solution[-1] = right_hand_side[-1] / diagonal[-1]
    for i in range(len(diagonal) - 2, -1, -1):
        solution[i] = (
            right_hand_side[i] - upper[i] * solution[i + 1]
        ) / diagonal[i]
    return solution


def effective_surface_transfer(
    internal_coefficient: float,
    boundary_coefficient: float,
    half_cell_width: float,
) -> float:
    """合并表面半单元内部阻力与 Robin 边界阻力。"""
    if internal_coefficient <= 0.0 or boundary_coefficient <= 0.0:
        raise ValueError("内部传输系数和表面传递系数必须为正。")
    return 1.0 / (
        half_cell_width / internal_coefficient + 1.0 / boundary_coefficient
    )


def assemble_control_volume_system(
    storage_coefficients: np.ndarray,
    old_values: np.ndarray,
    grid: RadialGrid,
    interface_coefficients: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """组装单元中心有限体积的后向 Euler 三对角方程组。"""
    cell_count = len(grid.centers)
    lower = np.zeros(cell_count - 1)
    upper = np.zeros(cell_count - 1)
    diagonal = storage_coefficients.copy()
    right_hand_side = storage_coefficients * old_values

    # 内部界面通量；中心面 r=0 的面积为零，自动满足对称零通量条件。
    conductance = grid.faces[1:-1] * interface_coefficients / grid.spacing
    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    upper[:] = -conductance
    lower[:] = -conductance

    # 表面通量以最后一个单元中心到环境的总阻力计算。
    effective_coefficient = effective_surface_transfer(
        internal_coefficient=surface_internal_coefficient,
        boundary_coefficient=surface_transfer_coefficient,
        half_cell_width=0.5 * grid.spacing,
    )
    surface_conductance = PELLET_RADIUS * effective_coefficient
    diagonal[-1] += surface_conductance
    right_hand_side[-1] += surface_conductance * boundary_value
    return lower, diagonal, upper, right_hand_side


def surface_value_from_last_cell(
    last_cell_value: float,
    internal_coefficient: float,
    boundary_coefficient: float,
    boundary_value: float,
    grid_spacing: float,
) -> float:
    """根据半单元内部阻力和对流阻力重构真实表面值。"""
    internal_conductance = 2.0 * internal_coefficient / grid_spacing
    return float(
        (internal_conductance * last_cell_value + boundary_coefficient * boundary_value)
        / (internal_conductance + boundary_coefficient)
    )


def reconstruct_output_profile(
    cell_values: np.ndarray,
    grid: RadialGrid,
    output_radii: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> np.ndarray:
    """将单元中心解重构并插值到题目指定的径向位置。"""
    first_radius_squared = grid.centers[0] ** 2
    second_radius_squared = grid.centers[1] ** 2
    center_value = (
        second_radius_squared * cell_values[0]
        - first_radius_squared * cell_values[1]
    ) / (second_radius_squared - first_radius_squared)

    surface_value = surface_value_from_last_cell(
        last_cell_value=float(cell_values[-1]),
        internal_coefficient=surface_internal_coefficient,
        boundary_coefficient=surface_transfer_coefficient,
        boundary_value=boundary_value,
        grid_spacing=grid.spacing,
    )

    reconstruction_radii = np.concatenate(([0.0], grid.centers, [PELLET_RADIUS]))
    reconstruction_values = np.concatenate(([center_value], cell_values, [surface_value]))
    return np.interp(output_radii, reconstruction_radii, reconstruction_values)


def radial_volume_average(cell_values: np.ndarray, grid: RadialGrid) -> float:
    """计算圆柱截面的体积加权平均值。"""
    return float(2.0 * np.sum(grid.volume_factors * cell_values) / PELLET_RADIUS**2)


def relative_balance_residual(storage_rate: float, outward_flux: float) -> float:
    """计算积分守恒式的无量纲残差。"""
    scale = abs(storage_rate) + abs(outward_flux) + 1.0e-30
    return abs(storage_rate + outward_flux) / scale


def advance_temperature(
    old_temperature: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    heat_transfer_coefficient: float,
) -> np.ndarray:
    """后向 Euler 推进一个温度时间步。"""
    storage_coefficients = (
        PELLET_DENSITY * PELLET_HEAT_CAPACITY * grid.volume_factors / time_step
    )
    interface_conductivity = np.full(len(grid.centers) - 1, THERMAL_CONDUCTIVITY)
    room_temperature = linear_interp(current_time, boundary.times, boundary.temperatures)

    system = assemble_control_volume_system(
        storage_coefficients=storage_coefficients,
        old_values=old_temperature,
        grid=grid,
        interface_coefficients=interface_conductivity,
        surface_internal_coefficient=THERMAL_CONDUCTIVITY,
        surface_transfer_coefficient=heat_transfer_coefficient,
        boundary_value=room_temperature,
    )
    return solve_tridiagonal(*system)


def advance_moisture(
    old_moisture: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    mass_transfer_coefficient: float,
) -> tuple[np.ndarray, int, float]:
    """用 Picard 迭代推进一个含水率时间步。"""
    storage_coefficients = grid.volume_factors / time_step
    iterate_moisture = old_moisture.copy()
    room_moisture = linear_interp(current_time, boundary.times, boundary.moisture)

    for iteration in range(1, MAX_ITERATIONS + 1):
        diffusivity = moisture_diffusivity(iterate_moisture)
        interface_diffusivity = harmonic_mean(diffusivity[:-1], diffusivity[1:])
        system = assemble_control_volume_system(
            storage_coefficients=storage_coefficients,
            old_values=old_moisture,
            grid=grid,
            interface_coefficients=interface_diffusivity,
            surface_internal_coefficient=float(diffusivity[-1]),
            surface_transfer_coefficient=mass_transfer_coefficient,
            boundary_value=room_moisture,
        )
        new_moisture = solve_tridiagonal(*system)
        final_error = float(
            np.max(
                np.abs(new_moisture - iterate_moisture)
                / (1.0 + np.abs(new_moisture))
            )
        )
        iterate_moisture = new_moisture
        if final_error < CONVERGENCE_TOL:
            return new_moisture, iteration, final_error

    raise RuntimeError(
        f"{current_time:g} s 的含水率 Picard 迭代未收敛，"
        f"最终相对变化量为 {final_error:.3e}。"
    )


def simulate_problem_one(
    end_time: float = END_TIME,
    time_step: float = TIME_STEP,
    cell_count: int = INTERNAL_CELL_COUNT,
    output_interval: float = OUTPUT_TIME_INTERVAL,
    heat_transfer_coefficient: float = CONVECTIVE_HEAT_COEFF,
    mass_transfer_coefficient: float = CONVECTIVE_MASS_COEFF,
) -> SimulationDetails:
    """完成问题一计算，并返回题目网格上的场与数值诊断。"""
    if min(end_time, time_step, output_interval) <= 0.0:
        raise ValueError("终止时间、时间步长和输出间隔必须为正。")
    step_count = int(round(end_time / time_step))
    output_stride = int(round(output_interval / time_step))
    if not np.isclose(step_count * time_step, end_time):
        raise ValueError("终止时间必须是时间步长的整数倍。")
    if output_stride < 1 or not np.isclose(output_stride * time_step, output_interval):
        raise ValueError("输出间隔必须是时间步长的整数倍。")

    boundary = read_drying_boundary(INPUT_FILE)
    if end_time > boundary.times[-1]:
        raise ValueError("附件一边界数据没有覆盖要求的计算时段。")
    grid = build_radial_grid(cell_count)

    temperature = np.full(cell_count, INITIAL_TEMPERATURE, dtype=float)
    moisture = np.full(cell_count, INITIAL_MOISTURE, dtype=float)

    output_times: list[float] = []
    output_temperatures: list[np.ndarray] = []
    output_moistures: list[np.ndarray] = []
    average_temperatures: list[float] = []
    average_moistures: list[float] = []
    iteration_counts: list[int] = []
    max_final_iteration_error = 0.0
    max_heat_balance_residual = 0.0
    max_moisture_balance_residual = 0.0

    for step in range(1, step_count + 1):
        current_time = step * time_step
        old_temperature = temperature.copy()
        old_moisture = moisture.copy()

        temperature = advance_temperature(
            old_temperature,
            current_time,
            time_step,
            grid,
            boundary,
            heat_transfer_coefficient,
        )
        moisture, iteration_count, final_iteration_error = advance_moisture(
            old_moisture,
            current_time,
            time_step,
            grid,
            boundary,
            mass_transfer_coefficient,
        )

        if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(moisture)):
            raise FloatingPointError(f"{current_time:g} s 出现非有限场变量。")
        if np.min(moisture) <= 0.0:
            raise FloatingPointError(f"{current_time:g} s 出现非正含水率。")

        room_temperature = linear_interp(current_time, boundary.times, boundary.temperatures)
        room_moisture = linear_interp(current_time, boundary.times, boundary.moisture)
        surface_temperature = surface_value_from_last_cell(
            float(temperature[-1]),
            THERMAL_CONDUCTIVITY,
            heat_transfer_coefficient,
            room_temperature,
            grid.spacing,
        )
        surface_diffusivity = float(moisture_diffusivity(moisture[-1:])[0])
        surface_moisture = surface_value_from_last_cell(
            float(moisture[-1]),
            surface_diffusivity,
            mass_transfer_coefficient,
            room_moisture,
            grid.spacing,
        )

        heat_storage_rate = float(
            np.sum(
                PELLET_DENSITY * PELLET_HEAT_CAPACITY * grid.volume_factors
                * (temperature - old_temperature) / time_step
            )
        )
        moisture_storage_rate = float(
            np.sum(grid.volume_factors * (moisture - old_moisture) / time_step)
        )
        heat_flux_outward = (
            PELLET_RADIUS * heat_transfer_coefficient
            * (surface_temperature - room_temperature)
        )
        moisture_flux_outward = (
            PELLET_RADIUS * mass_transfer_coefficient
            * (surface_moisture - room_moisture)
        )
        max_heat_balance_residual = max(
            max_heat_balance_residual,
            relative_balance_residual(heat_storage_rate, heat_flux_outward),
        )
        max_moisture_balance_residual = max(
            max_moisture_balance_residual,
            relative_balance_residual(moisture_storage_rate, moisture_flux_outward),
        )
        iteration_counts.append(iteration_count)
        max_final_iteration_error = max(max_final_iteration_error, final_iteration_error)

        if step % output_stride == 0:
            output_times.append(current_time)
            output_temperatures.append(
                reconstruct_output_profile(
                    temperature,
                    grid,
                    OUTPUT_RADII,
                    THERMAL_CONDUCTIVITY,
                    heat_transfer_coefficient,
                    room_temperature,
                )
            )
            output_moistures.append(
                reconstruct_output_profile(
                    moisture,
                    grid,
                    OUTPUT_RADII,
                    surface_diffusivity,
                    mass_transfer_coefficient,
                    room_moisture,
                )
            )
            average_temperatures.append(radial_volume_average(temperature, grid))
            average_moistures.append(radial_volume_average(moisture, grid))

    diagnostics = {
        "internal_cell_count": float(cell_count),
        "time_step": float(time_step),
        "max_picard_iterations": float(max(iteration_counts)),
        "mean_picard_iterations": float(np.mean(iteration_counts)),
        "max_final_picard_error": float(max_final_iteration_error),
        "max_heat_balance_residual": float(max_heat_balance_residual),
        "max_moisture_balance_residual": float(max_moisture_balance_residual),
    }
    return SimulationDetails(
        times=np.asarray(output_times),
        output_radii=OUTPUT_RADII.copy(),
        temperature_field=np.asarray(output_temperatures),
        moisture_field=np.asarray(output_moistures),
        average_temperature=np.asarray(average_temperatures),
        average_moisture=np.asarray(average_moistures),
        diagnostics=diagnostics,
    )


def solve_problem_one() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """保留原调用接口，返回时间、输出半径、温度场和含水率场。"""
    result = simulate_problem_one()
    return result.times, result.output_radii, result.temperature_field, result.moisture_field


def write_result_workbook(
    times: np.ndarray,
    output_radii: np.ndarray,
    temperature_field: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    """按附件三模板写入 result1.xlsx，不生成其他结果文件。"""
    workbook = load_workbook(OUTPUT_FILE)
    if len(workbook.worksheets) < 2:
        raise ValueError("result1.xlsx 必须包含温度和水分浓度两个工作表。")

    for worksheet, field in zip(
        workbook.worksheets[:2], (temperature_field, moisture_field)
    ):
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)

        worksheet.cell(1, 1).value = "时间\\到药材中心的距离"
        worksheet.cell(1, 1).number_format = "@"
        for column, radius in enumerate(output_radii, start=2):
            worksheet.cell(1, column).value = round(float(radius * 100.0), 1)
            worksheet.cell(1, column).number_format = "0.0"

        for row, current_time in enumerate(times, start=2):
            worksheet.cell(row, 1).value = int(round(float(current_time)))
            worksheet.cell(row, 1).number_format = "0"
            for column, value in enumerate(field[row - 2], start=2):
                worksheet.cell(row, column).value = round(float(value), 4)
                worksheet.cell(row, column).number_format = "0.0000"

    workbook.save(OUTPUT_FILE)


def report_largest_difference(
    label: str,
    times: np.ndarray,
    radii: np.ndarray,
    first_field: np.ndarray,
    second_field: np.ndarray,
) -> None:
    """输出两个同网格结果的最大绝对差及其发生位置。"""
    difference = np.abs(first_field - second_field)
    time_index, radius_index = np.unravel_index(
        int(np.argmax(difference)), difference.shape
    )
    print(
        f"{label}：最大绝对差={difference[time_index, radius_index]:.6e}，"
        f"位置为 t={times[time_index]:g} s、"
        f"r={radii[radius_index] * 100:g} cm；"
        f"终点最大差={np.max(difference[-1]):.6e}。"
    )


def run_numerical_validation(reference: SimulationDetails) -> None:
    """复现论文中的空间、时间离散检验，不生成额外文件。"""
    fine_space = simulate_problem_one(cell_count=320)
    fine_time = simulate_problem_one(cell_count=160, time_step=0.5)

    print("问题一数值检验：")
    report_largest_difference(
        "N=160 与 N=320 的温度场",
        reference.times,
        reference.output_radii,
        reference.temperature_field,
        fine_space.temperature_field,
    )
    report_largest_difference(
        "N=160 与 N=320 的含水率场",
        reference.times,
        reference.output_radii,
        reference.moisture_field,
        fine_space.moisture_field,
    )
    report_largest_difference(
        "时间步 1 s 与 0.5 s 的温度场",
        reference.times,
        reference.output_radii,
        reference.temperature_field,
        fine_time.temperature_field,
    )
    report_largest_difference(
        "时间步 1 s 与 0.5 s 的含水率场",
        reference.times,
        reference.output_radii,
        reference.moisture_field,
        fine_time.moisture_field,
    )
    print(f"主计算诊断：{reference.diagnostics}")


def main() -> None:
    parser = argparse.ArgumentParser(description="求解问题一的一维径向传热传质模型")
    parser.add_argument(
        "--validate",
        action="store_true",
        help="额外运行空间和时间步长检验，只在终端输出诊断",
    )
    arguments = parser.parse_args()

    result = simulate_problem_one()
    write_result_workbook(
        result.times,
        result.output_radii,
        result.temperature_field,
        result.moisture_field,
    )
    print(f"问题一计算完成，结果已写入：{OUTPUT_FILE}")
    if arguments.validate:
        run_numerical_validation(result)


if __name__ == "__main__":
    main()
