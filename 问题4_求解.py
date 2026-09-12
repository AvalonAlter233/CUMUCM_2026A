"""问题四：考虑实测径向收缩的移动边界热湿耦合模型。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Callable, NamedTuple

import numpy as np
from openpyxl import load_workbook

import 问题3_求解 as problem3


PROJECT_ROOT = Path(__file__).resolve().parent
RADIUS_FILE = PROJECT_ROOT / "附件" / "附件2.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "附件" / "附件3" / "result4.xlsx"
DIAGNOSTICS_FILE = PROJECT_ROOT / "附件" / "附件3" / "result4_diagnostics.json"
INTERNAL_FIELD_FILE = PROJECT_ROOT / "附件" / "附件3" / "result4_internal_field.npz"

INITIAL_RADIUS = 0.02
CONVECTIVE_HEAT_COEFF = 25.0
CONVECTIVE_MASS_COEFF = 8.0e-7
INITIAL_TEMPERATURE = 28.0
INITIAL_MOISTURE = 2.55
CRITICAL_MOISTURE = 0.15

INTERNAL_INTERVALS = 320
TIME_STEP = 30.0
REPORT_INTERVAL = 60.0
MAX_SIMULATION_TIME = 6.0 * 24.0 * 3600.0

CONVERGENCE_TOL = 1.0e-9
MAX_ITERATIONS = 200
RELAXATION_FACTOR = 0.6


class RadiusHistory(NamedTuple):
    """附件二给出的时间与药材半径，单位分别为 s 和 m。"""

    times: np.ndarray
    radii: np.ndarray


class ReferenceGrid(NamedTuple):
    """归一化径向坐标 xi 上的单元中心有限体积网格。"""

    faces: np.ndarray
    centers: np.ndarray
    volumes: np.ndarray
    spacing: float


class MaterialProperties(NamedTuple):
    """一组与含水率、温度配套的热湿物性函数。"""

    density: Callable[[np.ndarray], np.ndarray]
    heat_capacity: Callable[[np.ndarray], np.ndarray]
    thermal_conductivity: Callable[[np.ndarray], np.ndarray]
    moisture_diffusivity: Callable[[np.ndarray, np.ndarray], np.ndarray]


class InternalFieldData(NamedTuple):
    """供绘图使用的内部参考域单元中心结果。"""

    times: np.ndarray
    xi_centers: np.ndarray
    moisture: np.ndarray
    surface_radii: np.ndarray


def read_radius_history(path: Path) -> RadiusHistory:
    """读取附件二，校验时间递增和半径单调不增并转换为 SI 单位。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = [
        row
        for row in workbook.active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)
    if data.ndim != 2 or data.shape[1] < 2 or data.shape[0] < 2:
        raise ValueError("附件二缺少有效半径数据。")
    if not np.all(np.diff(data[:, 0]) > 0.0):
        raise ValueError("附件二时间必须严格递增。")
    radii = data[:, 1] / 100.0
    if np.any(radii <= 0.0):
        raise ValueError("附件二半径必须为正数。")
    if np.any(np.diff(radii) > 1.0e-12):
        raise ValueError("附件二半径必须单调不增。")
    return RadiusHistory(data[:, 0], radii)


def radius_value(current_time: float, history: RadiusHistory) -> float:
    """附件范围内分段线性插值，超出末时刻后保持末值。"""
    return float(
        np.interp(
            current_time,
            history.times,
            history.radii,
            left=history.radii[0],
            right=history.radii[-1],
        )
    )


def density(moisture: np.ndarray) -> np.ndarray:
    return 760.0 + 90.0 * moisture


def heat_capacity(moisture: np.ndarray) -> np.ndarray:
    return 1850.0 + 2150.0 * moisture / (moisture + 1.0)


def thermal_conductivity(moisture: np.ndarray) -> np.ndarray:
    return 0.12 + 0.20 * moisture / (moisture + 1.0)


def moisture_diffusivity(
    moisture: np.ndarray,
    temperature_celsius: np.ndarray,
) -> np.ndarray:
    safe_moisture = np.maximum(moisture, 1.0e-12)
    temperature_kelvin = temperature_celsius + 273.15
    return (
        4.2e-4
        * np.exp(-0.30 / safe_moisture)
        * np.exp(-3850.0 / temperature_kelvin)
    )


def material_properties(name: str) -> MaterialProperties:
    """返回附录三或附录四物性，供同一移动边界框架作机制对照。"""
    if name == "appendix4":
        return MaterialProperties(
            density, heat_capacity, thermal_conductivity, moisture_diffusivity
        )
    if name == "appendix3":
        return MaterialProperties(
            problem3.density,
            problem3.heat_capacity,
            problem3.thermal_conductivity,
            problem3.moisture_diffusivity,
        )
    raise ValueError("物性模型必须为 appendix3 或 appendix4。")


def build_reference_grid(intervals: int) -> ReferenceGrid:
    if intervals < 4:
        raise ValueError("内部径向区间数至少为 4。")
    faces = np.linspace(0.0, 1.0, intervals + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return ReferenceGrid(faces, centers, volumes, 1.0 / intervals)


def assemble_reference_system(
    old_values: np.ndarray,
    storage_capacity: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """在 xi=r/R(t) 上组装后向 Euler 有限体积方程。"""
    if radius <= 0:
        raise ValueError("药材半径必须为正数。")
    storage = storage_capacity * grid.volumes / time_step
    face_transport = problem3.harmonic_mean(
        transport_coefficients[:-1], transport_coefficients[1:]
    )
    internal_conductance = (
        grid.faces[1:-1]
        * face_transport
        / (radius**2 * grid.spacing)
    )
    diagonal = storage.copy()
    diagonal[:-1] += internal_conductance
    diagonal[1:] += internal_conductance
    lower = -internal_conductance.copy()
    upper = -internal_conductance.copy()
    right_hand_side = storage * old_values

    last_transport = max(float(transport_coefficients[-1]), 1.0e-30)
    physical_half_cell = 0.5 * radius * grid.spacing
    effective_transfer = 1.0 / (
        1.0 / external_transfer_coefficient
        + physical_half_cell / last_transport
    )
    surface_conductance = effective_transfer / radius
    diagonal[-1] += surface_conductance
    right_hand_side[-1] += surface_conductance * external_value
    return lower, diagonal, upper, right_hand_side


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    """使用稳定、无外部二进制依赖的 Thomas 算法。"""
    return problem3.solve_tridiagonal(
        lower, diagonal, upper, right_hand_side
    )


def advance_temperature(
    old_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_temperature: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
    properties: MaterialProperties,
) -> np.ndarray:
    system = assemble_reference_system(
        old_temperature,
        properties.density(reference_moisture)
        * properties.heat_capacity(reference_moisture),
        properties.thermal_conductivity(reference_moisture),
        CONVECTIVE_HEAT_COEFF,
        room_temperature,
        grid,
        radius,
        time_step,
    )
    return solve_tridiagonal(*system)


def advance_moisture(
    old_moisture: np.ndarray,
    reference_temperature: np.ndarray,
    reference_moisture: np.ndarray,
    room_moisture: float,
    grid: ReferenceGrid,
    radius: float,
    time_step: float,
    properties: MaterialProperties,
) -> np.ndarray:
    system = assemble_reference_system(
        old_moisture,
        np.ones_like(old_moisture),
        properties.moisture_diffusivity(
            reference_moisture, reference_temperature
        ),
        CONVECTIVE_MASS_COEFF,
        room_moisture,
        grid,
        radius,
        time_step,
    )
    return solve_tridiagonal(*system)


def reconstruct_output_field(
    cell_values: np.ndarray,
    transport_coefficients: np.ndarray,
    external_transfer_coefficient: float,
    external_value: float,
    physical_output_nodes: np.ndarray,
    grid: ReferenceGrid,
    radius: float,
) -> np.ndarray:
    """重构固定物理位置及当前移动表面的场值。"""
    if np.any(physical_output_nodes < 0.0) or np.any(
        physical_output_nodes >= radius
    ):
        raise ValueError("固定输出位置必须位于当前药材内部。")
    xi_nodes = physical_output_nodes / radius
    interior_values = np.interp(xi_nodes, grid.centers, cell_values)
    interior_values[0] = (9.0 * cell_values[0] - cell_values[1]) / 8.0

    last_transport = max(float(transport_coefficients[-1]), 1.0e-30)
    physical_half_cell = 0.5 * radius * grid.spacing
    half_cell_conductance = last_transport / physical_half_cell
    surface_value = (
        half_cell_conductance * cell_values[-1]
        + external_transfer_coefficient * external_value
    ) / (half_cell_conductance + external_transfer_coefficient)
    return np.concatenate([interior_values, np.array([surface_value])])


def template_output_nodes() -> np.ndarray:
    """题目模板中始终位于药材内部的固定物理位置，单位 m。"""
    return np.arange(0.0, 0.011 + 0.0005, 0.001)


def output_headers(nodes: np.ndarray) -> list[float | str]:
    return [round(float(node * 100.0), 1) for node in nodes] + ["药材表面"]


def solve_problem_four(
    internal_intervals: int = INTERNAL_INTERVALS,
    time_step: float = TIME_STEP,
    report_interval: float = REPORT_INTERVAL,
    max_time: float = MAX_SIMULATION_TIME,
    critical_moisture: float = CRITICAL_MOISTURE,
    moving_radius: bool = True,
    property_model: str = "appendix4",
    capture_internal: bool = False,
    boundary: problem3.DryingRoomBoundary | None = None,
    radius_history: RadiusHistory | None = None,
) -> tuple[
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    dict,
    InternalFieldData | None,
]:
    """用指定物性求解移动或固定半径干燥过程。"""
    report_steps = int(round(report_interval / time_step))
    total_steps = int(round(max_time / time_step))
    if report_steps < 1 or not np.isclose(
        report_steps * time_step, report_interval
    ):
        raise ValueError("输出间隔必须是内部时间步长的整数倍。")
    if total_steps < 1 or not np.isclose(total_steps * time_step, max_time):
        raise ValueError("最大时间必须是内部时间步长的整数倍。")

    if boundary is None:
        boundary = problem3.read_drying_boundary(problem3.INPUT_FILE)
    if radius_history is None:
        radius_history = read_radius_history(RADIUS_FILE)

    grid = build_reference_grid(internal_intervals)
    properties = material_properties(property_model)
    output_nodes = template_output_nodes()
    temperature = np.full(internal_intervals, INITIAL_TEMPERATURE)
    moisture = np.full(internal_intervals, INITIAL_MOISTURE)

    output_times: list[float] = []
    output_moisture: list[np.ndarray] = []
    output_surface_radius: list[float] = []
    internal_output_moisture: list[np.ndarray] = []
    iteration_counts: list[int] = []
    previous_time = 0.0
    previous_maximum = INITIAL_MOISTURE
    continuous_threshold_time: float | None = None
    threshold_before: float | None = None
    threshold_after: float | None = None
    initial_inventory = float(np.dot(moisture, grid.volumes))
    current_inventory = initial_inventory
    cumulative_outflow = 0.0
    maximum_step_balance_residual = 0.0
    center_controls_every_step = True
    radially_nonincreasing_every_step = True
    checked_steps = 0

    for step in range(1, total_steps + 1):
        current_time = step * time_step
        radius = (
            radius_value(current_time, radius_history)
            if moving_radius
            else float(radius_history.radii[0])
        )
        room_temperature, room_moisture = problem3.boundary_value(
            current_time, boundary
        )
        temperature_iterate = temperature.copy()
        moisture_iterate = moisture.copy()

        for iteration in range(1, MAX_ITERATIONS + 1):
            temperature_candidate = advance_temperature(
                temperature,
                moisture_iterate,
                room_temperature,
                grid,
                radius,
                time_step,
                properties,
            )
            moisture_candidate = advance_moisture(
                moisture,
                temperature_candidate,
                moisture_iterate,
                room_moisture,
                grid,
                radius,
                time_step,
                properties,
            )
            new_temperature = (
                RELAXATION_FACTOR * temperature_candidate
                + (1.0 - RELAXATION_FACTOR) * temperature_iterate
            )
            new_moisture = (
                RELAXATION_FACTOR * moisture_candidate
                + (1.0 - RELAXATION_FACTOR) * moisture_iterate
            )
            temperature_error = float(
                np.max(
                    np.abs(new_temperature - temperature_iterate)
                    / (1.0 + np.abs(new_temperature))
                )
            )
            moisture_error = float(
                np.max(
                    np.abs(new_moisture - moisture_iterate)
                    / (1.0 + np.abs(new_moisture))
                )
            )
            temperature_iterate = new_temperature
            moisture_iterate = new_moisture
            if max(temperature_error, moisture_error) < CONVERGENCE_TOL:
                break
        else:
            raise RuntimeError(
                f"{current_time:.0f} s 未收敛："
                f"{temperature_error:.3e}, {moisture_error:.3e}"
            )

        temperature = temperature_iterate
        moisture = moisture_iterate
        iteration_counts.append(iteration)
        if not np.all(np.isfinite(moisture)) or np.min(moisture) <= 0.0:
            raise FloatingPointError(f"{current_time:.0f} s 含水率异常。")

        diffusivity = properties.moisture_diffusivity(moisture, temperature)
        last_diffusivity = max(float(diffusivity[-1]), 1.0e-30)
        effective_mass_transfer = 1.0 / (
            1.0 / CONVECTIVE_MASS_COEFF
            + 0.5 * radius * grid.spacing / last_diffusivity
        )
        surface_outflow = (
            effective_mass_transfer
            / radius
            * (float(moisture[-1]) - room_moisture)
        )
        next_inventory = float(np.dot(moisture, grid.volumes))
        step_balance_residual = abs(
            next_inventory - current_inventory + time_step * surface_outflow
        ) / initial_inventory
        maximum_step_balance_residual = max(
            maximum_step_balance_residual, step_balance_residual
        )
        cumulative_outflow += time_step * surface_outflow
        current_inventory = next_inventory

        reconstructed = reconstruct_output_field(
            moisture,
            diffusivity,
            CONVECTIVE_MASS_COEFF,
            room_moisture,
            output_nodes,
            grid,
            radius,
        )
        current_maximum = max(
            float(np.max(moisture)), float(np.max(reconstructed))
        )
        center_value = float(reconstructed[0])
        center_controls_every_step = center_controls_every_step and bool(
            current_maximum <= center_value + 1.0e-12
        )
        radially_nonincreasing_every_step = (
            radially_nonincreasing_every_step
            and bool(np.all(np.diff(moisture) <= 1.0e-12))
            and bool(np.all(np.diff(reconstructed) <= 1.0e-12))
        )
        checked_steps += 1

        if (
            continuous_threshold_time is None
            and previous_maximum >= critical_moisture
            and current_maximum < critical_moisture
        ):
            continuous_threshold_time = problem3.linear_threshold_crossing(
                previous_time,
                previous_maximum,
                current_time,
                current_maximum,
                critical_moisture,
            )
            threshold_before = previous_maximum
            threshold_after = current_maximum

        reached_on_report = False
        if step % report_steps == 0:
            output_times.append(current_time)
            output_moisture.append(reconstructed.copy())
            output_surface_radius.append(radius)
            if capture_internal:
                internal_output_moisture.append(moisture.copy())
            reached_on_report = current_maximum < critical_moisture

        previous_time = current_time
        previous_maximum = current_maximum
        if reached_on_report:
            break
    else:
        raise RuntimeError(f"{max_time / 3600:.1f} h 内未达到阈值。")

    moisture_field = np.asarray(output_moisture)
    times = np.asarray(output_times)
    surface_radii = np.asarray(output_surface_radius)
    final_temperature = reconstruct_output_field(
        temperature,
        properties.thermal_conductivity(moisture),
        CONVECTIVE_HEAT_COEFF,
        problem3.boundary_value(times[-1], boundary)[0],
        output_nodes,
        grid,
        surface_radii[-1],
    )
    diagnostics = {
        "continuous_threshold_time_s": float(continuous_threshold_time),
        "discrete_threshold_time_s": float(times[-1]),
        "maximum_moisture_before": float(threshold_before),
        "maximum_moisture_after": float(threshold_after),
        "radius_at_continuous_threshold_cm": 100.0
        * (
            radius_value(float(continuous_threshold_time), radius_history)
            if moving_radius
            else float(radius_history.radii[0])
        ),
        "radius_at_discrete_threshold_cm": 100.0 * surface_radii[-1],
        "moving_radius": moving_radius,
        "property_model": property_model,
        "center_controls_threshold": center_controls_every_step,
        "all_domain_checked_every_step": checked_steps == len(iteration_counts),
        "radially_nonincreasing_every_step": radially_nonincreasing_every_step,
        "moisture_balance_relative_imbalance": (
            problem3.relative_balance_imbalance(
                initial_inventory, current_inventory, cumulative_outflow
            )
        ),
        "maximum_step_balance_relative_residual": (
            maximum_step_balance_residual
        ),
        "initial_reference_moisture_integral": initial_inventory,
        "final_reference_moisture_integral": current_inventory,
        "cumulative_reference_boundary_outflow": cumulative_outflow,
        "maximum_picard_iterations": int(max(iteration_counts)),
        "mean_picard_iterations": float(np.mean(iteration_counts)),
        "internal_intervals": internal_intervals,
        "reference_spacing": grid.spacing,
        "time_step_s": time_step,
        "report_interval_s": report_interval,
        "critical_moisture": critical_moisture,
        "final_center_moisture": float(moisture_field[-1, 0]),
        "final_surface_moisture": float(moisture_field[-1, -1]),
        "final_center_temperature_c": float(final_temperature[0]),
        "final_surface_temperature_c": float(final_temperature[-1]),
        "plateau_temperature_c": boundary.plateau_temperature,
        "plateau_moisture_kgkg": boundary.plateau_moisture,
        "radius_data_final_time_s": float(radius_history.times[-1]),
        "radius_data_final_cm": 100.0 * float(radius_history.radii[-1]),
    }
    internal_field = None
    if capture_internal:
        internal_field = InternalFieldData(
            times,
            grid.centers.copy(),
            np.asarray(internal_output_moisture),
            surface_radii,
        )
    return (
        times,
        output_nodes,
        moisture_field,
        surface_radii,
        diagnostics,
        internal_field,
    )


def write_internal_field_data(
    data: InternalFieldData,
    output_path: Path = INTERNAL_FIELD_FILE,
) -> None:
    """保存内部细网格结果；该文件只供绘图与复核，不改变官方表格。"""
    np.savez_compressed(
        output_path,
        times_s=data.times,
        xi_centers=data.xi_centers,
        moisture=data.moisture,
        surface_radii_m=data.surface_radii,
    )


def write_result_workbook(
    times: np.ndarray,
    nodes: np.ndarray,
    moisture_field: np.ndarray,
    output_path: Path = OUTPUT_FILE,
) -> None:
    """按 result4 模板写入固定物理位置和动态表面含水率。"""
    workbook = load_workbook(output_path)
    worksheet = workbook.worksheets[0]
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    if worksheet.max_column > 1:
        worksheet.delete_cols(2, worksheet.max_column - 1)
    worksheet.cell(1, 1).value = "时间\\到药材中心的距离"
    worksheet.cell(1, 1).number_format = "@"
    for column, header in enumerate(output_headers(nodes), start=2):
        worksheet.cell(1, column).value = header
        worksheet.cell(1, column).number_format = (
            "0.0" if isinstance(header, float) else "@"
        )
    for row, current_time in enumerate(times, start=2):
        worksheet.cell(row, 1).value = int(round(float(current_time)))
        worksheet.cell(row, 1).number_format = "0"
        for column, value in enumerate(moisture_field[row - 2], start=2):
            worksheet.cell(row, column).value = float(value)
            worksheet.cell(row, column).number_format = "0.0000"
    workbook.save(output_path)


def _case_record(name: str, label: str, diagnostics: dict) -> dict:
    record = dict(diagnostics)
    record["name"] = name
    record["label"] = label
    record["continuous_threshold_time_h"] = (
        diagnostics["continuous_threshold_time_s"] / 3600.0
    )
    record["discrete_threshold_time_h"] = (
        diagnostics["discrete_threshold_time_s"] / 3600.0
    )
    return record


def new_diagnostics_report(baseline_diagnostics: dict) -> dict:
    """创建只含当前基准结果的报告，避免沿用旧验证工况。"""
    return {
        "scope": (
            "移动参考域离散诊断与机制对照；附录三与附录四同时改变时，"
            "不能把总时间差全部解释为收缩效应。"
        ),
        "verification_complete": False,
        "baseline": _case_record(
            "appendix4_moving_radius",
            "附录4物性 + 实测收缩半径",
            baseline_diagnostics,
        ),
        "mechanism_comparison": [],
        "numerical_refinement": [],
        "boundary_sensitivity": [],
    }


def write_diagnostics_report(report: dict) -> None:
    with DIAGNOSTICS_FILE.open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="问题四移动边界干燥求解")
    parser.add_argument(
        "--verification",
        action="store_true",
        help="额外复算固定半径对照及空间、时间加密案例",
    )
    args = parser.parse_args()

    boundary = problem3.read_drying_boundary(problem3.INPUT_FILE)
    radius_history = read_radius_history(RADIUS_FILE)
    times, nodes, moisture, surface_radii, diagnostics, internal_field = (
        solve_problem_four(
            boundary=boundary,
            radius_history=radius_history,
            capture_internal=True,
        )
    )
    write_result_workbook(times, nodes, moisture)
    if internal_field is None:
        raise RuntimeError("未生成内部参考域结果。")
    write_internal_field_data(internal_field)
    print(f"问题四重算完成：{OUTPUT_FILE}")
    print(
        f"连续阈值：{diagnostics['continuous_threshold_time_s'] / 3600:.4f} h；"
        f"首次 60 s 离散达标："
        f"{diagnostics['discrete_threshold_time_s'] / 3600:.4f} h"
    )
    print(
        f"达标时半径：{diagnostics['radius_at_discrete_threshold_cm']:.4f} cm；"
        f"输出尺寸：{moisture.shape[0]}×{moisture.shape[1]}"
    )
    print(
        "参考域 C 方程累计相对收支不平衡："
        f"{diagnostics['moisture_balance_relative_imbalance']:.3e}"
    )

    report = new_diagnostics_report(diagnostics)
    if not args.verification:
        write_diagnostics_report(report)
        print(f"基准诊断已保存（未运行完整验证）：{DIAGNOSTICS_FILE}")
        return

    print("复算机制对照：附录4物性 + 固定半径")
    _, _, _, _, fixed_diagnostics, _ = solve_problem_four(
        moving_radius=False,
        boundary=boundary,
        radius_history=radius_history,
    )
    report["mechanism_comparison"].append(
        _case_record(
            "appendix4_fixed_radius", "附录4物性 + 固定半径", fixed_diagnostics
        )
    )

    print("复算机制对照：附录3物性 + 固定半径（问题三）")
    _, _, _, problem3_diagnostics = problem3.solve_problem_three(
        boundary=boundary
    )
    report["mechanism_comparison"].append(
        _case_record(
            "appendix3_fixed_radius",
            "附录3物性 + 固定半径（问题三）",
            problem3_diagnostics,
        )
    )

    print("复算机制对照：附录3物性 + 实测收缩半径")
    _, _, _, _, appendix3_moving_diagnostics, _ = solve_problem_four(
        moving_radius=True,
        property_model="appendix3",
        boundary=boundary,
        radius_history=radius_history,
    )
    report["mechanism_comparison"].append(
        _case_record(
            "appendix3_moving_radius",
            "附录3物性 + 实测收缩半径",
            appendix3_moving_diagnostics,
        )
    )

    refinement_cases = [
        ("space_refined", "空间加密 N=640, Δt=30 s", 640, 30.0),
        ("time_refined", "时间加密 N=320, Δt=15 s", 320, 15.0),
    ]
    for name, label, intervals, refined_time_step in refinement_cases:
        print(f"复算数值加密：{label}")
        _, _, _, _, refined_diagnostics, _ = solve_problem_four(
            internal_intervals=intervals,
            time_step=refined_time_step,
            boundary=boundary,
            radius_history=radius_history,
        )
        report["numerical_refinement"].append(
            _case_record(name, label, refined_diagnostics)
        )

    for scenario in problem3.build_boundary_scenarios(problem3.INPUT_FILE):
        if scenario.name not in {
            "temperature_minus_1sigma",
            "temperature_plus_1sigma",
        }:
            continue
        print(f"复算长期边界情景：{scenario.label}")
        _, _, _, _, scenario_diagnostics, _ = solve_problem_four(
            boundary=scenario.boundary,
            radius_history=radius_history,
        )
        report["boundary_sensitivity"].append(
            _case_record(scenario.name, scenario.label, scenario_diagnostics)
        )

    moving_h = report["baseline"]["continuous_threshold_time_h"]
    fixed_h = report["mechanism_comparison"][0][
        "continuous_threshold_time_h"
    ]
    appendix3_h = report["mechanism_comparison"][1][
        "continuous_threshold_time_h"
    ]
    appendix3_moving_h = report["mechanism_comparison"][2][
        "continuous_threshold_time_h"
    ]
    report["effects"] = {
        "appendix3_shrinkage_effect_hours": appendix3_moving_h - appendix3_h,
        "appendix4_shrinkage_effect_hours": moving_h - fixed_h,
        "fixed_radius_property_effect_hours": fixed_h - appendix3_h,
        "moving_radius_property_effect_hours": moving_h - appendix3_moving_h,
        "total_problem3_to_problem4_hours": moving_h - appendix3_h,
        "interpretation": (
            "两种物性下的收缩效应不同，说明物性与几何存在交互；"
            "不作唯一的加性贡献分解。"
        ),
    }
    report["verification_complete"] = True
    write_diagnostics_report(report)
    print(f"验证诊断已保存：{DIAGNOSTICS_FILE}")


if __name__ == "__main__":
    main()
