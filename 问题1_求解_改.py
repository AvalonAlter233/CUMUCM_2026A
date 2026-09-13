from __future__ import annotations

import argparse
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook


# 读哪个文件、写到哪个文件，先都在这几行里放好。
# 后面读表和存表的时候就比较好找。
project_root = Path(__file__).resolve().parent
input_file = project_root / "附件" / "附件1.xlsx"
output_file = project_root / "附件" / "附件3" / "result1.xlsx"

pellet_radius = 0.02
pellet_density = 820.0
pellet_heat_capacity = 2600.0
thermal_conductivity = 0.36
convective_heat_coeff = 25.0
convective_mass_coeff = 8.0e-7

initial_temperature = 28.0
initial_moisture = 2.55

default_end_time = 1800.0
default_time_step = 1.0
internal_cell_count = 160
output_time_interval = 1.0
report_radii = np.linspace(0.0, pellet_radius, 21)

convergence_tol = 1.0e-10
max_iterations = 30
# 有点整齐哈哈

class DryingRoomBoundary(NamedTuple):

    times: np.ndarray
    temperatures: np.ndarray
    moisture: np.ndarray


class RadialGrid(NamedTuple):

    centers: np.ndarray
    faces: np.ndarray
    volume_factors: np.ndarray
    spacing: float


class SimulationDetails(NamedTuple):

    times: np.ndarray
    output_radii: np.ndarray
    temperature_field: np.ndarray
    moisture_field: np.ndarray
    average_temperature: np.ndarray
    average_moisture: np.ndarray
    diagnostics: dict[str, float]


def read_room_data(path: Path) -> DryingRoomBoundary:
    # 先让表里的数据进来，时间、温度和水分都从这里取。
    # 这一段读完以后先留着，后面算到哪里，再从这份数据里拿对应的值。
    book = load_workbook(path, data_only=True, read_only=True)
    sheet = book.active
    rows = [
        row for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    data = np.asarray(rows, dtype=float)

    if data.ndim != 2 or data.shape[1] < 3 or data.shape[0] < 2:
        raise ValueError("附件一必须包含时间、温度、水分浓度三列有效数据。")
    if not np.all(np.isfinite(data[:, :3])):
        raise ValueError("附件一的前三列存在空值或非有限数值。")

    times = data[:, 0]
    if not np.all(np.diff(times) > 0):
        raise ValueError("附件一中的时间必须严格递增。")

    return DryingRoomBoundary(
        times=times,
        temperatures=data[:, 1],
        moisture=data[:, 2],
    )


def linear_interp(t: float, times: np.ndarray, values: np.ndarray) -> float:
    # 表里是一串采样时刻，计算时不一定每次都正好碰上。
    # 碰在两个采样点中间的，就在这里插一下，再接着往后算。
    # 也就是说，循环算到了哪个时间，就在这串数据里找它前后挨着的两个点。
    # 中间差的那一点，在这里接上即可。
    if t < times[0] or t > times[-1]:
        raise ValueError(
            f"计算时刻 {t:g} s 超出附件一范围 "
            f"[{times[0]:g}, {times[-1]:g}] s。"
        )
    return float(np.interp(t, times, values))


def moisture_diffusivity(moisture: np.ndarray) -> np.ndarray:
    # 这里先拿含水率算出扩散系数，外面要用多少个位置，就一起算多少个。
    # 先按下面的下限处理，再接着代到式子里。
    c_safe = np.maximum(moisture, 1.0e-12)
    return 7.0e-9 * np.exp(-0.89 / c_safe)


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    # 相邻两边的系数都到了这里，先合成界面要用的那一个。
    # 这一小步单独放着，后面遇到同样的就接着调用它。
    denom = np.maximum(left + right, 1.0e-30)
    return 2.0 * left * right / denom


def build_radial_grid(cell_count: int) -> RadialGrid:
    if cell_count < 2:
        raise ValueError("径向有限体积单元数至少为 2。")

    # 半径先一段一段分好，后面的数就放在这些小段上。
    # 每段的中间在哪、两边在哪，于此，后面用哪一个就取哪一个。
    faces = np.linspace(0.0, pellet_radius, cell_count + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    vol = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return RadialGrid(
        centers=centers,
        faces=faces,
        volume_factors=vol,
        spacing=float(faces[1] - faces[0]),
    )


def solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    # 下面消元的时候会改动数组，所以先复制出一份来。
    # 接下来就在这份副本上做，算到最后再把结果拿走，原来的数组还留在外面。
    # 后面会从前往后走一遍，再从后往前走回来，跟着两个循环看就行。
    # 数组虽然长，做的事情还是在相邻的几项之间接着算。
    diagonal = diagonal.astype(float, copy=True)
    right_hand_side = right_hand_side.astype(float, copy=True)
    upper = upper.astype(float, copy=True)

    for i in range(1, len(diagonal)):
        if abs(diagonal[i - 1]) < 1.0e-30:
            raise FloatingPointError("三对角方程组出现近零主对角元。")
        factor = lower[i - 1] / diagonal[i - 1]
        diagonal[i] -= factor * upper[i - 1]
        right_hand_side[i] -= factor * right_hand_side[i - 1]

    sol = np.empty_like(right_hand_side)
    sol[-1] = right_hand_side[-1] / diagonal[-1]
    for i in range(len(diagonal) - 2, -1, -1):
        sol[i] = (
            right_hand_side[i] - upper[i] * sol[i + 1]
        ) / diagonal[i]
    return sol


def surface_transfer(
    internal_coefficient: float,
    boundary_coefficient: float,
    half_cell_width: float,
) -> float:
    # 到最外面以后，还有半个小单元的距离要经过。
    # 这一点距离也算上，再把它和表面的传递放在一起处理。
    # 走到边上时就多这一层关系，先把它单独理一下
    if internal_coefficient <= 0.0 or boundary_coefficient <= 0.0:
        raise ValueError("内部传输系数和表面传递系数必须为正。")
    return 1.0 / (
        half_cell_width / internal_coefficient + 1.0 / boundary_coefficient
    )


def assemble_system(
    storage_coefficients: np.ndarray,
    old_values: np.ndarray,
    grid: RadialGrid,
    interface_coefficients: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 先把几条对角线和右边的数组各自放好，再把内部相连的项添上去。
    # 边界还有一项要接进来，等下面都接齐了，再整套交给求解的地方。
    # 这一段先顺着看数组怎么放，后面看调用时就比较容易对上。
    cell_count = len(grid.centers)
    lower = np.zeros(cell_count - 1)
    upper = np.zeros(cell_count - 1)
    diag = storage_coefficients.copy()
    rhs = storage_coefficients * old_values

    g_face = grid.faces[1:-1] * interface_coefficients / grid.spacing
    diag[:-1] += g_face
    diag[1:] += g_face
    upper[:] = -g_face
    lower[:] = -g_face

    h_eff = surface_transfer(
        internal_coefficient=surface_internal_coefficient,
        boundary_coefficient=surface_transfer_coefficient,
        half_cell_width=0.5 * grid.spacing,
    )
    g_surf = pellet_radius * h_eff
    diag[-1] += g_surf
    rhs[-1] += g_surf * boundary_value
    return lower, diag, upper, rhs


def surface_value(
    last_cell_value: float,
    internal_coefficient: float,
    boundary_coefficient: float,
    boundary_value: float,
    grid_spacing: float,
) -> float:
    # 最后一个单元的值已经有了，表面这里再单独算一下。
    # 算完返回一个数，外面接到它以后再继续拼接输出结果。
    g_inner = 2.0 * internal_coefficient / grid_spacing
    return float(
        (g_inner * last_cell_value + boundary_coefficient * boundary_value)
        / (g_inner + boundary_coefficient)
    )


def sample_profile(
    cell_values: np.ndarray,
    grid: RadialGrid,
    output_radii: np.ndarray,
    surface_internal_coefficient: float,
    surface_transfer_coefficient: float,
    boundary_value: float,
) -> np.ndarray:
    # 计算时用的那些点，和表里最后要填的点并不是完全重合的。
    # 故而到了这里还得整理一下，把中心、内部和表面的数接到一起，再往表格要求的位置上取。
    # 最后拿出去的是表格要的那排数，so这里还得多绕这么一小步。
    # 把输出位置和计算位置对上，后面写表就可以顺着排了oh耶。
    r1_sq = grid.centers[0] ** 2
    r2_sq = grid.centers[1] ** 2
    axis_val = (
        r2_sq * cell_values[0]
        - r1_sq * cell_values[1]
    ) / (r2_sq - r1_sq)

    val_surf = surface_value(
        last_cell_value=float(cell_values[-1]),
        internal_coefficient=surface_internal_coefficient,
        boundary_coefficient=surface_transfer_coefficient,
        boundary_value=boundary_value,
        grid_spacing=grid.spacing,
    )

    r_fit = np.concatenate(([0.0], grid.centers, [pellet_radius]))
    val_fit = np.concatenate(([axis_val], cell_values, [val_surf]))
    return np.interp(output_radii, r_fit, val_fit)


def radial_average(cell_values: np.ndarray, grid: RadialGrid) -> float:
    # 需要平均值的时候就到这里来，把各个小单元的数按对应权重加起来。
    # 先把这一份算好，再交回去和当前时刻放在一起记着。
    return float(2.0 * np.sum(grid.volume_factors * cell_values) / pellet_radius**2)


def balance_residual(storage_rate: float, outward_flux: float) -> float:
    # 这一小步把两边的量放到一起看，再按下面的尺度换成一个检查数。
    # 算出来先返回，外面需要留记录的时候会把它收起来。
    scale = abs(storage_rate) + abs(outward_flux) + 1.0e-30
    return abs(storage_rate + outward_flux) / scale


def advance_temperature(
    old_temperature: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    heat_transfer_coefficient: float,
) -> np.ndarray:
    # 温度这一部分先单独往前算一步，需要的环境温度也在这里取。
    # 系数准备齐以后交给方程求解，拿到结果再回到外面的时间循环。
    storage = (
        pellet_density * pellet_heat_capacity * grid.volume_factors / time_step
    )
    k_face = np.full(len(grid.centers) - 1, thermal_conductivity)
    temp_air = linear_interp(current_time, boundary.times, boundary.temperatures)

    lin_sys = assemble_system(
        storage_coefficients=storage,
        old_values=old_temperature,
        grid=grid,
        interface_coefficients=k_face,
        surface_internal_coefficient=thermal_conductivity,
        surface_transfer_coefficient=heat_transfer_coefficient,
        boundary_value=temp_air,
    )
    return solve_tridiagonal(*lin_sys)


def advance_moisture(
    old_moisture: np.ndarray,
    current_time: float,
    time_step: float,
    grid: RadialGrid,
    boundary: DryingRoomBoundary,
    mass_transfer_coefficient: float,
) -> tuple[np.ndarray, int, float]:
    storage = grid.volume_factors / time_step
    water_iter = old_moisture.copy()
    water_air = linear_interp(current_time, boundary.times, boundary.moisture)

    # 现在的含水率算出一份扩散系数，再拿这份系数往下算。
    # 含水率更新以后又会有新的系数，就继续这样接着来，直到前后变化已经很小。
    # 这几轮都还在同一个时间步里，等这一步收拾好了，时间才继续往前走。
    for iteration in range(1, max_iterations + 1):
        d_cell = moisture_diffusivity(water_iter)
        d_face = harmonic_mean(d_cell[:-1], d_cell[1:])
        lin_sys = assemble_system(
            storage_coefficients=storage,
            old_values=old_moisture,
            grid=grid,
            interface_coefficients=d_face,
            surface_internal_coefficient=float(d_cell[-1]),
            surface_transfer_coefficient=mass_transfer_coefficient,
            boundary_value=water_air,
        )
        water_new = solve_tridiagonal(*lin_sys)
        final_error = float(
            np.max(
                np.abs(water_new - water_iter)
                / (1.0 + np.abs(water_new))
            )
        )
        water_iter = water_new
        if final_error < convergence_tol:
            return water_new, iteration, final_error

    raise RuntimeError(
        f"{current_time:g} s 的含水率 Picard 迭代未收敛，"
        f"最终相对变化量为 {final_error:.3e}。"
    )


def simulate_question1(
    end_time: float = default_end_time,
    time_step: float = default_time_step,
    cell_count: int = internal_cell_count,
    output_interval: float = output_time_interval,
    heat_transfer_coefficient: float = convective_heat_coeff,
    mass_transfer_coefficient: float = convective_mass_coeff,
) -> SimulationDetails:
    # 先看看这次给的几个时间设置能不能凑到一起。
    # 步长和输出间隔这些数先对好，后面的循环就按它们一小步一小步往后走。
    # 这些设置看起来不多，先在开头过一遍，后面就不用一边算一边再管它们了。
    if min(end_time, time_step, output_interval) <= 0.0:
        raise ValueError("终止时间、时间步长和输出间隔必须为正。")
    n_steps = int(round(end_time / time_step))
    save_every = int(round(output_interval / time_step))
    if not np.isclose(n_steps * time_step, end_time):
        raise ValueError("终止时间必须是时间步长的整数倍。")
    if save_every < 1 or not np.isclose(save_every * time_step, output_interval):
        raise ValueError("输出间隔必须是时间步长的整数倍。")

    boundary = read_room_data(input_file)
    if end_time > boundary.times[-1]:
        raise ValueError("附件一边界数据没有覆盖要求的计算时段。")
    mesh = build_radial_grid(cell_count)

    temp = np.full(cell_count, initial_temperature, dtype=float)
    water = np.full(cell_count, initial_moisture, dtype=float)

    # 这几个列表先放在这里，刚开始里面还没有东西。
    # 等走到需要保存的时刻，再把温度、水分和平均值分别往里面添。
    # 刚开始可能会有空列表，我说你先别急，下面到了保存的地方，会一行一行填进来。
    time_log: list[float] = []
    temp_log: list[np.ndarray] = []
    water_log: list[np.ndarray] = []
    temp_mean_log: list[float] = []
    water_mean_log: list[float] = []
    iter_log: list[int] = []
    max_iter_err = 0.0
    max_heat_err = 0.0
    max_water_err = 0.0

    # 上一时刻的值先另外留着，这一时刻更新时还要用，等本步算完再往下走。
    for step in range(1, n_steps + 1):
        t_now = step * time_step
        temp_old = temp.copy()
        water_old = water.copy()

        temp = advance_temperature(
            temp_old,
            t_now,
            time_step,
            mesh,
            boundary,
            heat_transfer_coefficient,
        )
        water, n_iter, iter_err = advance_moisture(
            water_old,
            t_now,
            time_step,
            mesh,
            boundary,
            mass_transfer_coefficient,
        )

        if not np.all(np.isfinite(temp)) or not np.all(np.isfinite(water)):
            raise FloatingPointError(f"{t_now:g} s 出现非有限场变量。")
        if np.min(water) <= 0.0:
            raise FloatingPointError(f"{t_now:g} s 出现非正含水率。")

        temp_air = linear_interp(t_now, boundary.times, boundary.temperatures)
        water_air = linear_interp(t_now, boundary.times, boundary.moisture)
        temp_surf = surface_value(
            float(temp[-1]),
            thermal_conductivity,
            heat_transfer_coefficient,
            temp_air,
            mesh.spacing,
        )
        d_surf = float(moisture_diffusivity(water[-1:])[0])
        water_surf = surface_value(
            float(water[-1]),
            d_surf,
            mass_transfer_coefficient,
            water_air,
            mesh.spacing,
        )

        heat_rate = float(
            np.sum(
                pellet_density * pellet_heat_capacity * mesh.volume_factors
                * (temp - temp_old) / time_step
            )
        )
        water_rate = float(
            np.sum(mesh.volume_factors * (water - water_old) / time_step)
        )
        heat_out = (
            pellet_radius * heat_transfer_coefficient
            * (temp_surf - temp_air)
        )
        water_out = (
            pellet_radius * mass_transfer_coefficient
            * (water_surf - water_air)
        )
        max_heat_err = max(
            max_heat_err,
            balance_residual(heat_rate, heat_out),
        )
        max_water_err = max(
            max_water_err,
            balance_residual(water_rate, water_out),
        )
        iter_log.append(n_iter)
        max_iter_err = max(max_iter_err, iter_err)

        # 要保存时，就把这一份结果存下来。
        # 存完以后还会继续算
        if step % save_every == 0:
            time_log.append(t_now)
            temp_log.append(
                sample_profile(
                    temp,
                    mesh,
                    report_radii,
                    thermal_conductivity,
                    heat_transfer_coefficient,
                    temp_air,
                )
            )
            water_log.append(
                sample_profile(
                    water,
                    mesh,
                    report_radii,
                    d_surf,
                    mass_transfer_coefficient,
                    water_air,
                )
            )
            temp_mean_log.append(radial_average(temp, mesh))
            water_mean_log.append(radial_average(water, mesh))

    # 这一趟算了多少轮、检查量是多少，也一起放到结果里。
    diagnostics = {
        "internal_cell_count": float(cell_count),
        "time_step": float(time_step),
        "max_picard_iterations": float(max(iter_log)),
        "mean_picard_iterations": float(np.mean(iter_log)),
        "max_final_picard_error": float(max_iter_err),
        "max_heat_balance_residual": float(max_heat_err),
        "max_moisture_balance_residual": float(max_water_err),
    }
    return SimulationDetails(
        times=np.asarray(time_log),
        output_radii=report_radii.copy(),
        temperature_field=np.asarray(temp_log),
        moisture_field=np.asarray(water_log),
        average_temperature=np.asarray(temp_mean_log),
        average_moisture=np.asarray(water_mean_log),
        diagnostics=diagnostics,
    )


def solve_question1() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    result = simulate_question1()
    return result.times, result.output_radii, result.temperature_field, result.moisture_field


def write_results(
    times: np.ndarray,
    output_radii: np.ndarray,
    temperature_field: np.ndarray,
    moisture_field: np.ndarray,
) -> None:
    # 数已经算好了，接下来就是把它们往表格里放。
    # 先把旧数据行清一下，再照着这次的时间顺序往下填，填完后一起保存。
    # 每一行先时间、再跟着对应位置的数，顺序照着表头往后接。
    book = load_workbook(output_file)
    if len(book.worksheets) < 2:
        raise ValueError("result1.xlsx 必须包含温度和水分浓度两个工作表。")

    for sheet, field in zip(
        book.worksheets[:2], (temperature_field, moisture_field)
    ):
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)

        sheet.cell(1, 1).value = "时间\\到药材中心的距离"
        sheet.cell(1, 1).number_format = "@"
        for column, radius in enumerate(output_radii, start=2):
            sheet.cell(1, column).value = round(float(radius * 100.0), 1)
            sheet.cell(1, column).number_format = "0.0"

        for row, t_now in enumerate(times, start=2):
            sheet.cell(row, 1).value = int(round(float(t_now)))
            sheet.cell(row, 1).number_format = "0"
            for column, value in enumerate(field[row - 2], start=2):
                sheet.cell(row, column).value = round(float(value), 4)
                sheet.cell(row, column).number_format = "0.0000"

    book.save(output_file)


def report_difference(
    label: str,
    times: np.ndarray,
    radii: np.ndarray,
    first_field: np.ndarray,
    second_field: np.ndarray,
) -> None:
    # 两份结果现在放在一起，先找差得最多的那个位置。
    # 找到以后，把时间、半径和对应的差值一起打印出来，回头查看时就能对上。
    diff = np.abs(first_field - second_field)
    it, ir = np.unravel_index(
        int(np.argmax(diff)), diff.shape
    )
    print(
        f"{label}：最大绝对差={diff[it, ir]:.6e}，"
        f"位置为 t={times[it]:g} s、"
        f"r={radii[ir] * 100:g} cm；"
        f"终点最大差={np.max(diff[-1]):.6e}。"
    )


def check_convergence(reference: SimulationDetails) -> None:
    # 前面主结果已经有了，这里算出的几份是接着拿来对照的。
    # 等打印结束，再一起看差值和它对应的位置
    space_fine = simulate_question1(cell_count=320)
    time_fine = simulate_question1(cell_count=160, time_step=0.5)

    print("问题一数值检验：")
    report_difference(
        "N=160 与 N=320 的温度场",
        reference.times,
        reference.output_radii,
        reference.temperature_field,
        space_fine.temperature_field,
    )
    report_difference(
        "N=160 与 N=320 的含水率场",
        reference.times,
        reference.output_radii,
        reference.moisture_field,
        space_fine.moisture_field,
    )
    report_difference(
        "时间步 1 s 与 0.5 s 的温度场",
        reference.times,
        reference.output_radii,
        reference.temperature_field,
        time_fine.temperature_field,
    )
    report_difference(
        "时间步 1 s 与 0.5 s 的含水率场",
        reference.times,
        reference.output_radii,
        reference.moisture_field,
        time_fine.moisture_field,
    )
    print(f"主计算诊断：{reference.diagnostics}")


def main() -> None:
    # 懒得写注释了这里
    parser = argparse.ArgumentParser(description="求解问题一的一维径向传热传质模型")
    parser.add_argument(
        "--validate",
        action="store_true",
        help="额外运行空间和时间步长检验，只在终端输出诊断",
    )
    args = parser.parse_args()

    result = simulate_question1()
    write_results(
        result.times,
        result.output_radii,
        result.temperature_field,
        result.moisture_field,
    )
    print(f"问题一计算完成，结果已写入：{output_file}")
    if args.validate:
        check_convergence(result)


if __name__ == "__main__":
    main()
# 噫嘘唏，写完了喵