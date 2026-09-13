from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib as mpl
import matplotlib.font_manager as font_manager
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
from openpyxl import load_workbook

if (Path(__file__).parent / "问题2_求解_改.py").is_file():
    import 问题2_求解_改 as problem2_solver
else:
    import 问题2_求解 as problem2_solver


nature_figure_scripts = Path.home() / ".codex" / "skills" / "nature-figure" / "scripts"
if str(nature_figure_scripts) not in sys.path:
    sys.path.insert(0, str(nature_figure_scripts))
from audit_panel_alignment import require_matplotlib_panel_alignment


project_root = Path(__file__).resolve().parent
result_one = project_root / "附件" / "附件3" / "result1.xlsx"
result_two = project_root / "附件" / "附件3" / "result2.xlsx"
figure_dir = project_root / "figures" / "问题二"

cjk_font_candidates = (
    "STSong", "SimSun", "Source Han Serif SC", "Noto Serif CJK SC",
)
installed_fonts = {font.name for font in font_manager.fontManager.ttflist}
available_cjk_fonts = [name for name in cjk_font_candidates if name in installed_fonts]
if not available_cjk_fonts:
    print("提示：未检测到中文字体，图中中文可能显示为方框。")
heading_font = "STZhongsong" if "STZhongsong" in installed_fonts else available_cjk_fonts[0]

mpl.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": [*available_cjk_fonts, "Times New Roman", "Times", "DejaVu Serif"],
        "mathtext.fontset": "stix",
        "axes.unicode_minus": False,
        "font.size": 8,
        "axes.titlesize": 9.2,
        "axes.titleweight": "bold",
        "axes.labelsize": 8,
        "axes.labelweight": "bold",
        "legend.fontsize": 7.3,
        "xtick.labelsize": 7.5,
        "ytick.labelsize": 7.5,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.linewidth": 0.7,
        "legend.frameon": False,
        "figure.dpi": 120,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    }
)

export_dpi = 600
pad_inches = 0.06
figure_formats = ("png",)

reference_palette = ("#44757A", "#452A3D", "#D44C3C", "#EED5B7")
color_teal, color_plum, color_coral, color_sand = reference_palette
color_center = color_teal
color_average = color_plum
color_surface = color_coral

new_figure_bases = (
    "图1_常物性与变物性模型对比",
    "图2_物性参数耦合演化",
    "图3_中心平均表面响应",
    "图4_热湿场时空演化",
    "图5_数值收敛验证",
)

# 这里沿用的是原稿记下来的全 3 h 验证数值，画图时直接从这几组数里取。
# 每一项里，前面的数是全时段最大差，后面的数是终点最大差，按这个顺序放的。
validation_metrics = {
    "grid_temperature": (1.16e-4, 9.58e-6),
    "grid_moisture": (1.67e-2, 1.71e-5),
    "time_temperature": (8.76e-4, 3.53e-5),
    "time_moisture": (7.54e-4, 5.58e-6),
}


def add_initial_state(
    times: np.ndarray, field: np.ndarray, initial_value: float
) -> tuple[np.ndarray, np.ndarray]:
    # 时间序列前面先补上初始时刻的那一行，补完再往后用。
    # 后面的几张图都可能要接上起点，所以这一小步单独放在这里处理。
    if field.ndim != 2 or field.shape[0] != len(times):
        raise ValueError("时间轴与场数组的第一维不一致。")
    initial_row = np.full((1, field.shape[1]), float(initial_value))
    return np.concatenate(([0.0], times)), np.vstack((initial_row, field))


def radial_average(radii_cm: np.ndarray, field: np.ndarray) -> np.ndarray:
    # 需要画平均曲线时，先到这里把平均值算好。
    # 算出的这一串数和时间序列对应着，后面拿过去直接画就行，不用再在画图的地方重复算。
    # 平均曲线也得先有数才能画，所以这里先多做这一步，再回去接后面的图。
    # 算完的结果等轮到平均值那条线时能用上
    radii_cm = np.asarray(radii_cm, dtype=float)
    field = np.asarray(field, dtype=float)
    if radii_cm.ndim != 1 or len(radii_cm) < 2:
        raise ValueError("径向坐标必须是一维且至少包含两个点。")
    if not np.all(np.diff(radii_cm) > 0.0):
        raise ValueError("径向坐标必须严格递增。")
    if field.ndim != 2 or field.shape[1] != len(radii_cm):
        raise ValueError("场数组列数必须与径向坐标长度一致。")
    weighted = field * radii_cm[np.newaxis, :]
    integral = np.trapezoid(weighted, radii_cm, axis=1)
    return 2.0 * integral / radii_cm[-1] ** 2


def read_fields(
    path: Path,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    # 两个工作表都先读一遍，温度和含水率分别取出来。
    # 表头里的位置、每一行的时间也一起保留，后面比较和画图都会用到这些东西
    workbook = load_workbook(path, data_only=True, read_only=True)
    if len(workbook.worksheets) < 2:
        raise ValueError(f"{path.name} 必须包含温度和含水率两个工作表。")

    def read_sheet(sheet) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
        # 这一小段还是先拿第一行的表头，再把后面的数据行收进来。
        # 几列东西分清楚以后，一起返回给外面，接着再读另一张表
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        rows = [
            row
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        ]
        times_s = np.asarray([row[0] for row in rows], dtype=float)
        radii_cm = np.asarray(list(header)[1:], dtype=float)
        field = np.asarray([row[1:] for row in rows], dtype=float)
        return times_s, radii_cm, field

    times_s, radii_cm, temperature = read_sheet(workbook.worksheets[0])
    moisture_times_s, moisture_radii_cm, moisture = read_sheet(workbook.worksheets[1])
    if not np.allclose(times_s, moisture_times_s):
        raise ValueError(f"{path.name} 两个工作表的时间轴不一致。")
    if not np.allclose(radii_cm, moisture_radii_cm):
        raise ValueError(f"{path.name} 两个工作表的径向坐标不一致。")
    return times_s, radii_cm, temperature, moisture


def load_plot_data() -> dict[str, np.ndarray]:
    # 先把第一问和第二问的结果都准备好，比较的时候两边都得用。
    # 它们共同覆盖的那一段时间也顺便检查一下，后面就照着这一段把数据接到一起
    t1, r1, temp1, water1 = read_fields(result_one)
    times_s, radii_cm, temp, water = read_fields(result_two)
    if not np.all(np.diff(times_s) > 0.0):
        raise ValueError("问题二结果时间轴必须严格递增。")
    if not np.allclose(r1, radii_cm):
        raise ValueError("问题一、问题二的径向输出网格不一致。")
    shared_times = times_s <= t1[-1]
    if not np.allclose(times_s[shared_times], t1):
        raise ValueError("问题一、问题二在共同时间窗内的输出时刻不一致。")

    return {
        "times_s": times_s,
        "radii_cm": radii_cm,
        "temperature_field": temp,
        "moisture_field": water,
        "average_temperature": radial_average(radii_cm, temp),
        "average_moisture": radial_average(radii_cm, water),
        "q1_times_s": t1,
        "q1_temperature_field": temp1,
        "q1_moisture_field": water1,
    }


def save_figure(
    figure: plt.Figure,
    base_name: str,
    exclude_axes: list[plt.Axes] | None = None,
) -> None:
    # 画布上的内容先更新好，再让这张图经过原来的对齐检查。
    # 检查通过以后再导出文件，保存完成了就把图关闭，接着处理后面那一张。
    # 一张图的事情在这里做完，再换下一张，顺着往下跑就好
    figure_dir.mkdir(parents=True, exist_ok=True)
    figure.canvas.draw()
    with TemporaryDirectory(prefix="problem2-figure-qa-") as tmp_dir:
        require_matplotlib_panel_alignment(
            figure,
            json_out=Path(tmp_dir) / f"{base_name}.alignment.json",
            exclude_axes=exclude_axes or [],
            row_groups=getattr(figure, "_alignment_row_groups", None),
            tolerance_pt=1.5,
            gutter_tolerance_pt=1.5,
            require_panel_labels=False,
            strict=True,
        )
    figure.savefig(
        figure_dir / f"{base_name}.png",
        dpi=export_dpi,
        bbox_inches="tight",
        pad_inches=pad_inches,
        facecolor="white",
    )
    plt.close(figure)


def close_figure(figure: plt.Figure) -> None:
    plt.close(figure)


def style_axis(axis: plt.Axes, *, show_grid: bool = True) -> None:
    # 坐标轴上的这些细节在这里一起处理，画到哪张图都可以接着用。
    # 先把线和刻度摆好，标题用的字体对齐
    axis.set_axisbelow(True)
    if show_grid:
        axis.grid(
            True,
            which="major",
            color="#AAA5A8",
            linestyle="--",
            linewidth=0.55,
            alpha=0.70,
        )
    else:
        axis.grid(False)
    axis.tick_params(length=3, width=0.7)
    for label in (*axis.get_xticklabels(), *axis.get_yticklabels()):
        if not any("\u4e00" <= character <= "\u9fff" for character in label.get_text()):
            label.set_fontfamily("Times New Roman")
    for text in (axis.title, axis.xaxis.label, axis.yaxis.label):
        text.set_fontfamily(heading_font)
        text.set_fontweight("bold")
    for spine in (axis.spines["left"], axis.spines["bottom"]):
        spine.set_color("#51474D")
        spine.set_linewidth(0.75)


def set_title(figure: plt.Figure, title: str, y: float = 0.97) -> None:
    text = figure.suptitle(title, x=0.5, y=y, fontsize=12, fontweight="bold")
    text.set_fontfamily(heading_font)


def plot_endpoints(
    axis: plt.Axes,
    times_h: np.ndarray,
    field: np.ndarray,
    linestyle: str,
    model_name: str,
) -> list[plt.Line2D]:
    # 这次先只拿中心和表面，按顺序往当前这块坐标轴上画。
    # 画出的线先存在列表里，外面后面需要时还能接着拿它们来整理图例
    lines = []
    for values, position, color in (
        (field[:, 0], "中心", color_center),
        (field[:, -1], "表面", color_surface),
    ):
        (line,) = axis.plot(
            times_h,
            values,
            color=color,
            linestyle=linestyle,
            linewidth=1.7 if linestyle == "-" else 1.3,
            alpha=1.0 if linestyle == "-" else 0.75,
            label=f"{position}·{model_name}",
        )
        lines.append(line)
    return lines


def plot_comparison(data: dict[str, np.ndarray]) -> plt.Figure:
    # 两种模型先对齐到共同的时间段上，再分别拿出对应的曲线。
    # 曲线放好以后，图例也一起整理一下，后面看哪一条线时就能找到它的说明。
    # 图例的名字跟着曲线走，先把线画齐再整理
    t1 = np.asarray(data["q1_times_s"], dtype=float)
    t2 = np.asarray(data["times_s"], dtype=float)
    if t1[-1] > t2[-1]:
        raise ValueError("问题二结果没有覆盖问题一的比较时间窗。")
    temp2raw = np.asarray(data["temperature_field"], dtype=float)
    water2raw = np.asarray(data["moisture_field"], dtype=float)
    temp2shared = np.column_stack(
        [np.interp(t1, t2, temp2raw[:, column])
         for column in range(temp2raw.shape[1])]
    )
    water2shared = np.column_stack(
        [np.interp(t1, t2, water2raw[:, column])
         for column in range(water2raw.shape[1])]
    )
    q1_t, temp1 = add_initial_state(
        t1,
        np.asarray(data["q1_temperature_field"], dtype=float),
        problem2_solver.initial_temperature,
    )
    _, water1 = add_initial_state(
        t1,
        np.asarray(data["q1_moisture_field"], dtype=float),
        problem2_solver.initial_moisture,
    )
    q2_t, temp2 = add_initial_state(
        t1,
        temp2shared,
        problem2_solver.initial_temperature,
    )
    _, water2 = add_initial_state(
        t1,
        water2shared,
        problem2_solver.initial_moisture,
    )
    if not np.allclose(q1_t, q2_t):
        raise ValueError("问题一、问题二的共同时间窗不一致。")

    fig, axs = plt.subplots(1, 2, figsize=(7.0, 2.9), sharex=True)
    panels = (
        (axs[0], temp1, temp2, "温度 / °C", "温度响应"),
        (axs[1], water1, water2, "干基含水率 / kg/kg", "水分响应"),
    )
    legend_lines = []
    for ax, const_field, var_field, ylabel, title in panels:
        const_lines = plot_endpoints(ax, q1_t / 3600.0, const_field, "--", "常物性")
        var_lines = plot_endpoints(ax, q2_t / 3600.0, var_field, "-", "变物性")
        if ax is axs[0]:
            legend_lines = [var_lines[0], const_lines[0], var_lines[1], const_lines[1]]
        ax.set_title(title, pad=6, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("时间 / h")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0.0, 0.5)
        style_axis(ax)
    set_title(fig, "常物性与变物性模型对比")
    fig.legend(
        handles=legend_lines,
        labels=[line.get_label() for line in legend_lines],
        ncol=4,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.855),
        handlelength=2.0,
        columnspacing=1.15,
    )
    fig.subplots_adjust(wspace=0.30, bottom=0.20, top=0.67, left=0.09, right=0.98)
    fig._alignment_row_groups = [["a", "b"]]
    return fig


def plot_properties(data: dict[str, np.ndarray]) -> plt.Figure:
    # 先用已有的温度和含水率，把要展示的物性数值算出来。
    # 数算齐了以后再去画变化曲线，这一段先把画图前需要的准备做完
    times_s, temp = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["temperature_field"], dtype=float),
        problem2_solver.initial_temperature,
    )
    _, water = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["moisture_field"], dtype=float),
        problem2_solver.initial_moisture,
    )
    rho_cp = (
        problem2_solver.density(water) * problem2_solver.heat_capacity(water)
    ) / 1.0e6
    diffusivity = problem2_solver.moisture_diffusivity(water, temp) * 1.0e9

    fig, axs = plt.subplots(1, 2, figsize=(7.0, 2.85), sharex=True)
    lines = []
    panels = (
        (axs[0], rho_cp, "体积热容 / 10^6 J/(m^3·K)", "热储存能力"),
        (axs[1], diffusivity, "扩散系数 D / 10^-9 m^2/s", "水分扩散能力"),
    )
    for ax, values, ylabel, title in panels:
        for column, label, color in ((0, "中心", color_center), (-1, "表面", color_surface)):
            (line,) = ax.plot(
                times_s / 3600.0,
                values[:, column],
                color=color,
                linewidth=1.65,
                label=label,
            )
            if ax is axs[0]:
                lines.append(line)
        ax.set_title(title, pad=6, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("时间 / h")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0.0, 3.0)
        style_axis(ax)
    set_title(fig, "物性参数耦合演化")
    fig.legend(
        handles=lines,
        labels=[line.get_label() for line in lines],
        ncol=2,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.855),
        columnspacing=1.5,
    )
    fig.subplots_adjust(wspace=0.32, bottom=0.20, top=0.67, left=0.11, right=0.98)
    fig._alignment_row_groups = [["a", "b"]]
    return fig


def plot_response(data: dict[str, np.ndarray]) -> plt.Figure:
    # 如此如此，这般这般（）
    times_s, temp = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["temperature_field"], dtype=float),
        problem2_solver.initial_temperature,
    )
    _, water = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["moisture_field"], dtype=float),
        problem2_solver.initial_moisture,
    )
    temp_mean = np.concatenate(([problem2_solver.initial_temperature], np.asarray(data["average_temperature"])))
    water_mean = np.concatenate(([problem2_solver.initial_moisture], np.asarray(data["average_moisture"])))

    fig, axs = plt.subplots(1, 2, figsize=(7.0, 2.9), sharex=True)
    lines = []
    panels = (
        (axs[0], temp, temp_mean, "温度 / °C", "温度"),
        (axs[1], water, water_mean, "干基含水率 / kg/kg", "含水率"),
    )
    for ax, field, average, ylabel, title in panels:
        for values, label, color, linestyle in (
            (field[:, 0], "中心", color_center, "-"),
            (average, "体积平均", color_average, "--"),
            (field[:, -1], "表面", color_surface, "-"),
        ):
            (line,) = ax.plot(
                times_s / 3600.0,
                values,
                color=color,
                linestyle=linestyle,
                linewidth=1.65 if linestyle == "-" else 1.3,
                label=label,
            )
            if ax is axs[0]:
                lines.append(line)
        ax.set_title(title, pad=6, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("时间 / h")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0.0, 3.0)
        style_axis(ax)
    set_title(fig, "药材内部温湿响应")
    fig.legend(
        handles=lines,
        labels=[line.get_label() for line in lines],
        ncol=3,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.855),
        columnspacing=1.4,
    )
    fig.subplots_adjust(wspace=0.30, bottom=0.20, top=0.67, left=0.09, right=0.98)
    fig._alignment_row_groups = [["a", "b"]]
    return fig


def plot_field(data: dict[str, np.ndarray]) -> plt.Figure:
    # 这一张需要的是整块场数据
    # 后面把对应的数铺到图里，能看到不同位置上的变化
    times_s, temp = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["temperature_field"], dtype=float),
        problem2_solver.initial_temperature,
    )
    _, water = add_initial_state(
        np.asarray(data["times_s"], dtype=float),
        np.asarray(data["moisture_field"], dtype=float),
        problem2_solver.initial_moisture,
    )
    radii_cm = np.asarray(data["radii_cm"], dtype=float)

    fig, axs = plt.subplots(1, 2, figsize=(7.0, 3.05), sharex=True, sharey=True)
    colorbar_axes = []
    panels = (
        (axs[0], temp, "magma", "温度场", "°C"),
        (axs[1], water, "viridis", "含水率场", "kg/kg"),
    )
    for ax, field, cmap, title, bar_label in panels:
        mesh = ax.pcolormesh(
            radii_cm,
            times_s / 3600.0,
            field,
            shading="auto",
            cmap=cmap,
            vmin=float(np.min(field)),
            vmax=float(np.max(field)),
        )
        cbar = fig.colorbar(mesh, ax=ax, fraction=0.046, pad=0.04)
        cbar.set_label(bar_label, rotation=270, labelpad=12)
        cbar.ax.yaxis.label.set_fontfamily(heading_font)
        cbar.ax.yaxis.label.set_fontweight("bold")
        for tick in cbar.ax.get_yticklabels():
            tick.set_fontfamily("Times New Roman")
        cbar.ax.grid(False)
        colorbar_axes.append(cbar.ax)
        ax.set_title(title, pad=6, fontweight="bold", fontfamily=heading_font)
        ax.set_xlabel("半径 r / cm")
        ax.set_ylabel("时间 / h")
        ax.set_xlim(radii_cm[0], radii_cm[-1])
        ax.set_ylim(0.0, 3.0)
        style_axis(ax, show_grid=False)
    set_title(fig, "热湿场时空演化")
    fig._alignment_exclude_axes = colorbar_axes
    fig._alignment_row_groups = [["a", "b"]]
    fig.subplots_adjust(wspace=0.28, bottom=0.17, top=0.78, left=0.09, right=0.92)
    return fig


def plot_validation(data: dict[str, np.ndarray]) -> plt.Figure:
    # 这张验证图直接读取前面保留下来的那几组数值。
    # 把每组的标签和对应的数放好，再照着原来的布局展示出来
    # 前面留好的数在这里排成图，先把一组放好，再接着放下一组。
    del data
    categories = np.arange(2)
    category_labels = ("全时段最大", "3 h 终点")
    figure, axes = plt.subplots(1, 2, figsize=(7.0, 2.9), sharex=True)
    panels = (
        (axes[0], validation_metrics["grid_temperature"], validation_metrics["time_temperature"], "最大绝对差 / °C", "温度场"),
        (axes[1], validation_metrics["grid_moisture"], validation_metrics["time_moisture"], "最大绝对差 / kg/kg", "含水率场"),
    )
    legend_lines = []
    for axis, grid_values, time_values, ylabel, title in panels:
        for values, label, color, marker in (
            (grid_values, "空间网格加密", color_teal, "o"),
            (time_values, "时间步长减半", color_plum, "s"),
        ):
            (line,) = axis.plot(
                categories,
                values,
                color=color,
                marker=marker,
                markersize=4.5,
                linewidth=1.5,
                label=label,
            )
            if axis is axes[0]:
                legend_lines.append(line)
        axis.set_yscale("log")
        axis.yaxis.set_major_formatter(mticker.FuncFormatter(lambda value, _: f"{value:.0e}"))
        axis.set_xticks(categories, category_labels)
        axis.set_ylabel(ylabel)
        axis.set_title(title, pad=6, fontweight="bold", fontfamily=heading_font)
        axis.set_xlim(-0.15, 1.15)
        style_axis(axis)
    set_title(figure, "空间与时间离散检验")
    figure.legend(
        handles=legend_lines,
        labels=[line.get_label() for line in legend_lines],
        ncol=2,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.855),
        columnspacing=1.5,
    )
    figure.subplots_adjust(wspace=0.31, bottom=0.19, top=0.67, left=0.11, right=0.98)
    figure._alignment_row_groups = [["a", "b"]]
    return figure


old_figure_names = (
    "图1_烘房温度与水分边界.png",
    "图2_边界采样间隔.png",
    "图3_边界变化率.png",
    "图4_温度径向剖面.png",
    "图5_含水率径向剖面.png",
    "图6_中心与表面轨迹.png",
    "图7_温度场时空分布.png",
    "图8_含水率场时空分布.png",
    "图9_空间非均匀性.png",
)


def clear_old_figures() -> None:
    # 这里按前面列好的旧图名字逐个去找，找到了才接着清理。
    # 收尾
    figure_dir.mkdir(parents=True, exist_ok=True)
    for basename in old_figure_names:
        target = figure_dir / basename
        if target.parent != figure_dir:
            raise ValueError(f"拒绝清理图集目录之外的路径：{target}")
        if target.exists():
            target.unlink()


def export_figures(data: dict[str, np.ndarray]) -> None:
    builders = (
        plot_comparison,
        plot_properties,
        plot_response,
        plot_field,
        plot_validation,
    )
    for builder, base_name in zip(builders, new_figure_bases):
        fig = builder(data)
        save_figure(
            fig,
            base_name,
            exclude_axes=list(getattr(fig, "_alignment_exclude_axes", [])),
        )

    missing = [
        figure_dir / f"{base_name}.png"
        for base_name in new_figure_bases
        if not (figure_dir / f"{base_name}.png").exists()
    ]
    if missing:
        raise RuntimeError(f"新问题二图集导出不完整：{missing}")
    clear_old_figures()


def main() -> None:
    export_figures(load_plot_data())
    print(f"问题二绘图完成，共生成 {len(new_figure_bases)} 张 PNG：{figure_dir}")

# 图还可以，就这样吧
if __name__ == "__main__":
    main()
